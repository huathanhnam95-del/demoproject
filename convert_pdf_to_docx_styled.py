import fitz
from docx import Document
from docx.shared import Pt, RGBColor, Inches
import os
import sys

# Ensure stdout can handle utf-8 to avoid charmap errors
sys.stdout.reconfigure(encoding='utf-8')

pdf_path = r'C:\Users\Admin\Downloads\Tam hon cao thuong (EdmondDeAmicis, HaMaiAnh dich)_original.pdf'
docx_path = r'C:\Users\Admin\Downloads\Tam_hon_cao_thuong_standard_styled.docx'

def convert_pdf_to_docx_styled(pdf_path, docx_path):
    doc = fitz.open(pdf_path)
    docx_doc = Document()
    
    # Set page size to standard 8.5 x 11
    section = docx_doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    for page_num in range(len(doc)):
        print(f'Processing page {page_num + 1}/{len(doc)}')
        page = doc.load_page(page_num)
        width = page.rect.width
        
        # Get dictionary of blocks
        blocks = page.get_text('dict')['blocks']
        
        # Separate into left and right columns
        left_blocks = []
        right_blocks = []
        
        for b in blocks:
            x0 = b['bbox'][0]
            if x0 < width / 2:
                left_blocks.append(b)
            else:
                right_blocks.append(b)
                
        # Sort each column by y0 (top to bottom)
        left_blocks.sort(key=lambda b: b['bbox'][1])
        right_blocks.sort(key=lambda b: b['bbox'][1])
        
        # Combine columns: read left column entirely, then right column
        ordered_blocks = left_blocks + right_blocks
        
        for b in ordered_blocks:
            if b['type'] == 0:  # Text block
                paragraph = docx_doc.add_paragraph()
                
                # Iterate through lines and spans to preserve styling
                for line_idx, line in enumerate(b['lines']):
                    for span in line['spans']:
                        text = span['text']
                        if not text.strip():
                            continue
                            
                        run = paragraph.add_run(text)
                        
                        # Apply flags
                        flags = span['flags']
                        if flags & 16: run.bold = True
                        if flags & 2: run.italic = True
                        
                        # Apply font size
                        run.font.size = Pt(span['size'])
                        
                        # Apply color
                        c = span['color']
                        # PyMuPDF color is sRGB integer: (R << 16) | (G << 8) | B
                        r = (c >> 16) & 255
                        g = (c >> 8) & 255
                        bl = c & 255
                        run.font.color.rgb = RGBColor(r, g, bl)
                    
                    # Add a space at the end of the line if it's not the last line
                    if line_idx < len(b['lines']) - 1:
                        paragraph.add_run(' ')
                        
            elif b['type'] == 1:  # Image block
                image_bytes = b['image']
                image_ext = b['ext']
                
                temp_image_path = f"temp_img_{page_num}_{b['number']}.{image_ext}"
                with open(temp_image_path, 'wb') as f:
                    f.write(image_bytes)
                
                try:
                    # Calculate reasonable width
                    img_width_px = b['width']
                    # Limit image width to roughly max page width (6.5 inches)
                    img_width_in = min(img_width_px / 72.0, 6.5)
                    
                    # Center the image
                    p = docx_doc.add_paragraph()
                    p.alignment = 1 # Center
                    r = p.add_run()
                    r.add_picture(temp_image_path, width=Inches(img_width_in))
                except Exception as e:
                    print(f"Error adding image on page {page_num}: {e}")
                finally:
                    if os.path.exists(temp_image_path):
                        os.remove(temp_image_path)
                
    docx_doc.save(docx_path)
    print(f'Successfully created {docx_path}')

if __name__ == '__main__':
    convert_pdf_to_docx_styled(pdf_path, docx_path)
