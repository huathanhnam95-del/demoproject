import os
import time
import re
import sys
import argparse
import openpyxl
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

sys.stdout.reconfigure(encoding="utf-8")

# --- API Configuration ---
OLLAMA_BASE_URL = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("LOCAL_GEMMA_MODEL", os.environ.get("OLLAMA_MODEL", "gemma4:12b"))

DEFAULT_EXCEL_PATH = r"C:\Cursor AI\public\database\LMCMA\LMCMA\LMCMA.xlsx"
EXCEL_PATH = os.environ.get("LMCMA_EXCEL_PATH", DEFAULT_EXCEL_PATH)
NEGATIVE_STEM_RE = re.compile(r"\b(false|incorrect|wrong|not true|except)\b", re.IGNORECASE)

def safe_save(workbook, path, retries=5):
    temp_path = f"{path}.tmp.{os.getpid()}.xlsx"
    workbook.save(temp_path)
    for attempt in range(1, retries + 1):
        try:
            os.replace(temp_path, path)
            return
        except PermissionError:
            if attempt == retries:
                raise
            time.sleep(0.2 * attempt)
    if os.path.exists(temp_path):
        os.remove(temp_path)

def parse_lmcma_content(text):
    if not text:
        return None
    
    # Split by hyphens (at least 3 hyphens)
    parts = [p.strip() for p in re.split(r'-{3,}', text) if p.strip()]
    if len(parts) < 2:
        return None
        
    question = parts[0]
    choices_str = parts[1]
    
    choices = []
    for line in choices_str.split('\n'):
        trimmed = line.strip()
        match = re.match(r'^\[([xX\s]*)\]\s*(.*)$', trimmed)
        if match:
            is_correct = "x" in match.group(1).lower()
            choice_text = match.group(2).strip()
            choices.append({
                'text': choice_text,
                'is_correct': is_correct
            })
            
    return {
        'question': question,
        'choices': choices
    }

def is_negative_question(question):
    return bool(NEGATIVE_STEM_RE.search(question or ""))


def clean_explanation_html(raw_text):
    text = str(raw_text or "").strip()
    text = re.sub(r"^```(?:html)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    text = re.sub(r"</?(?:html|body)[^>]*>", "", text, flags=re.IGNORECASE)
    return text.strip()


def generate_explanation(parsed_data, transcript, base_url=OLLAMA_BASE_URL, model=OLLAMA_MODEL):
    question = parsed_data['question']
    negative_question = is_negative_question(question)
    question_type = "negative stem: selected answers must be false, contradicted, or unsupported" if negative_question else "standard stem: selected answers must be true and supported"
    
    options_lines = []
    for c in parsed_data['choices']:
        if negative_question:
            label = "(Selected answer: this statement should be false, contradicted, or unsupported)" if c['is_correct'] else "(Not selected: this statement should be true or supported)"
        else:
            label = "(Selected answer: supported by the transcript)" if c['is_correct'] else "(Not selected: contradicted, too broad, or not mentioned)"
        options_lines.append(f"- {c['text']} {label}")
    options_text = "\n".join(options_lines)
    
    prompt = f"""
You are an expert PTE Academic tutor.
Analyze the following listening transcript, question, and options, and provide a clear, detailed, and easy-to-understand explanation of why each option should or should not be selected.

IMPORTANT:
- The labels identify the learner's expected selection, not whether a statement is true in ordinary language.
- For negative stems asking for false, incorrect, wrong, not true, or EXCEPT statements, selected answers are false, contradicted, or unsupported statements.
- Do not describe a supported true statement as a selected false answer.

TRANSCRIPT:
{transcript}

QUESTION:
{question}

QUESTION TYPE:
{question_type}

OPTIONS:
{options_text}

Task Instructions:
1. Explain every option in the order shown.
2. For each option, state "Select" or "Do not select" first.
3. State whether the option is supported, contradicted, or not mentioned in the transcript, and cite the specific transcript evidence.
4. For negative stems, clearly separate false statements that should be selected from true/supported statements that should not be selected.
5. Write the response in clean HTML format. Use standard HTML tags:
   - Use <p> for paragraphs.
   - Use <strong> for emphasis.
   - Use <ul> and <li> for lists.
   - Do NOT include any outer markdown backticks (such as ```html) or outer wrapping tags like <html>, <body>. Start directly with the HTML content.
   - Make the tone supportive, encouraging, and highly instructional.
"""
    
    url = f"{base_url}/api/generate"
    
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": 0.2,
            "num_ctx": 8192,
            "num_predict": 2048
        }
    }
    
    try:
        response = requests.post(url, json=payload, timeout=300)
        if response.status_code == 200:
            return clean_explanation_html(response.json().get("response", ""))
        else:
            print(f"API Error ({response.status_code}): {response.text}")
            return None
    except Exception as e:
        print(f"Exception during request: {e}")
        return None

def process_row(row_idx, q_id, title, answer_col, transcript_col, base_url, model):
    parsed = parse_lmcma_content(answer_col)
    if not parsed:
        return row_idx, q_id, None
    
    explanation = generate_explanation(parsed, transcript_col, base_url=base_url, model=model)
    return row_idx, q_id, explanation

def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Generate LMCMA explanations from workbook answers and transcripts.")
    parser.add_argument("--workbook", default=EXCEL_PATH, help="Path to the LMCMA workbook.")
    parser.add_argument("--force", action="store_true", help="Regenerate explanations even when Column E already has content.")
    parser.add_argument("--ids", default="", help="Comma-separated question IDs to regenerate regardless of existing explanation content.")
    parser.add_argument("--base-url", default=OLLAMA_BASE_URL, help="Ollama base URL.")
    parser.add_argument("--model", default=OLLAMA_MODEL, help="Ollama model name.")
    return parser.parse_args(argv)


def parse_id_filter(ids_text):
    return {item.strip() for item in str(ids_text or "").split(",") if item.strip()}


def main(argv=None):
    args = parse_args(argv)
    force_ids = parse_id_filter(args.ids)

    print(f"Loading workbook: {args.workbook}...")
    wb = openpyxl.load_workbook(args.workbook)
    sheet = wb.active
    
    # Ensure header column E is 'EXPLANATION'
    header_val = sheet.cell(row=1, column=5).value
    if header_val != "EXPLANATION":
        sheet.cell(row=1, column=5, value="EXPLANATION")
        safe_save(wb, args.workbook)
        print("Set Column E header to 'EXPLANATION'")
    
    max_row = sheet.max_row
    tasks = []
    
    # Identify rows that need processing
    for row_idx in range(2, max_row + 1):
        q_id = sheet.cell(row=row_idx, column=1).value
        title = sheet.cell(row=row_idx, column=2).value
        answer_col = sheet.cell(row=row_idx, column=3).value
        transcript_col = sheet.cell(row=row_idx, column=4).value
        existing_explanation = sheet.cell(row=row_idx, column=5).value
        q_id_text = str(q_id).strip()
        
        if not answer_col or not transcript_col:
            print(f"Row {row_idx}/{max_row} (ID {q_id}): Skipping because answer or transcript content is missing.")
            continue

        if not args.force and q_id_text not in force_ids and existing_explanation and len(str(existing_explanation).strip()) > 20:
            continue
            
        tasks.append((row_idx, q_id, title, answer_col, transcript_col))
        
    total_tasks = len(tasks)
    print(f"Pending tasks to enrich: {total_tasks}")
    if total_tasks == 0:
        print("No pending tasks. Enrichment complete.")
        return

    success_count = 0
    fail_count = 0
    
    # We use ThreadPoolExecutor with max_workers=1 to process sequentially
    with ThreadPoolExecutor(max_workers=1) as executor:
        future_to_row = {
            executor.submit(process_row, row_idx, q_id, title, answer_col, transcript_col, args.base_url, args.model): (row_idx, q_id)
            for row_idx, q_id, title, answer_col, transcript_col in tasks
        }
        
        for future in as_completed(future_to_row):
            row_idx, q_id = future_to_row[future]
            try:
                row_idx, q_id, explanation = future.result()
                if explanation:
                    # Reload workbook to avoid conflict with other saves, write, and save
                    wb_write = openpyxl.load_workbook(args.workbook)
                    sheet_write = wb_write.active
                    sheet_write.cell(row=row_idx, column=5, value=explanation)
                    safe_save(wb_write, args.workbook)
                    success_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Successfully enriched. ({success_count}/{total_tasks})")
                else:
                    fail_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Failed to generate explanation.")
            except Exception as exc:
                fail_count += 1
                print(f"Row {row_idx} (ID {q_id}) generated an exception: {exc}")
                
    print(f"\nProcessing finished! Success: {success_count}, Failed: {fail_count}")

if __name__ == "__main__":
    main()
