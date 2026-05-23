import os
import json
import random
import time
import openpyxl
import requests

KOKORO_API = "http://localhost:8880/v1/audio/speech"
EXCEL_PATH = r"C:\Cursor AI\public\database\LMCMA\LMCMA\LMCMA.xlsx"
AUDIO_DIR = r"C:\Cursor AI\public\database\LMCMA\audio"
MANIFEST_PATH = os.path.join(AUDIO_DIR, "manifest.json")

VOICES = [
    {"id": "af_alloy", "name": "Alloy", "gender": "female", "accent": "American"},
    {"id": "af_bella", "name": "Bella", "gender": "female", "accent": "American"},
    {"id": "af_heart", "name": "Heart", "gender": "female", "accent": "American"},
    {"id": "af_kore", "name": "Kore", "gender": "female", "accent": "American"},
    {"id": "af_sarah", "name": "Sarah", "gender": "female", "accent": "American"},
    {"id": "bf_emma", "name": "Emma", "gender": "female", "accent": "British"},
    {"id": "am_echo", "name": "Echo", "gender": "male", "accent": "American"},
    {"id": "am_eric", "name": "Eric", "gender": "male", "accent": "American"},
    {"id": "am_fenrir", "name": "Fenrir", "gender": "male", "accent": "American"},
    {"id": "am_liam", "name": "Liam", "gender": "male", "accent": "American"},
    {"id": "am_michael", "name": "Michael", "gender": "male", "accent": "American"},
    {"id": "am_puck", "name": "Puck", "gender": "male", "accent": "American"},
    {"id": "bm_fable", "name": "Fable", "gender": "male", "accent": "British"},
    {"id": "bm_george", "name": "George", "gender": "male", "accent": "British"},
    {"id": "bm_lewis", "name": "Lewis", "gender": "male", "accent": "British"}
]

def atomic_write_json(path, payload):
    temp_path = f"{path}.tmp.{os.getpid()}"
    with open(temp_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)
    os.replace(temp_path, path)

def generate_audio(text, voice_id, filepath, retries=3):
    payload = {
        "input": text,
        "voice": voice_id,
        "response_format": "mp3",
        "speed": 0.9
    }
    
    for attempt in range(retries):
        try:
            response = requests.post(KOKORO_API, json=payload, timeout=60)
            if response.status_code == 200:
                with open(filepath, "wb") as f:
                    f.write(response.content)
                return True
            else:
                print(f"  Attempt {attempt+1} failed with status {response.status_code}: {response.text}")
        except Exception as e:
            print(f"  Attempt {attempt+1} threw exception: {e}")
        time.sleep(2 * (attempt + 1))
    return False

def main():
    print("=== Kokoro TTS LMCMA Audio Generator ===")
    
    # Load manifest if exists
    if os.path.exists(MANIFEST_PATH):
        try:
            with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
                manifest = json.load(f)
        except Exception:
            manifest = {}
    else:
        manifest = {}
        
    os.makedirs(AUDIO_DIR, exist_ok=True)
    
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb.active
    max_row = sheet.max_row
    print(f"Loaded Excel with {max_row - 1} questions.")
    
    generated_count = 0
    skipped_count = 0
    failed_count = 0
    
    for row_idx in range(2, max_row + 1):
        q_id_val = sheet.cell(row=row_idx, column=1).value
        if q_id_val is None:
            continue
        
        q_id = str(q_id_val)
        transcript = sheet.cell(row=row_idx, column=4).value
        if not transcript or not str(transcript).strip():
            print(f"Skipping Row {row_idx} (ID {q_id}) due to empty transcript.")
            continue
            
        transcript_clean = str(transcript).strip()
        
        # Use a per-question seed for deterministic voice selection without mutating global randomness.
        selected_voices = random.Random(int(q_id_val)).sample(VOICES, 3)
        
        # Ensure question folder exists
        q_folder = os.path.join(AUDIO_DIR, q_id)
        os.makedirs(q_folder, exist_ok=True)
        
        question_voices_meta = []
        
        print(f"Question ID {q_id}: Generating 3 audios...")
        for voice in selected_voices:
            filename = f"LMCMA_{q_id}_{voice['id']}.mp3"
            filepath = os.path.join(q_folder, filename)
            
            # Check if file exists
            if os.path.exists(filepath) and os.path.getsize(filepath) > 1000:
                # Existing file is valid
                skipped_count += 1
            else:
                success = generate_audio(transcript_clean, voice['id'], filepath)
                if success:
                    generated_count += 1
                    print(f"  Generated voice: {voice['id']}")
                else:
                    failed_count += 1
                    print(f"  FAILED to generate voice: {voice['id']}")
                    
            question_voices_meta.append({
                "id": voice["id"],
                "name": voice["name"],
                "gender": voice["gender"],
                "accent": voice["accent"],
                "file": filename
            })
            
        manifest[q_id] = question_voices_meta
        
        # Save manifest incrementally
        atomic_write_json(MANIFEST_PATH, manifest)
            
    print("\n=== Audio Generation Complete ===")
    print(f"Generated: {generated_count}")
    print(f"Skipped: {skipped_count}")
    print(f"Failed: {failed_count}")
    print(f"Manifest saved to: {MANIFEST_PATH}")

if __name__ == "__main__":
    main()
