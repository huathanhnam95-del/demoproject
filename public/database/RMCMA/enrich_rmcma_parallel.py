import os
import time
import re
import openpyxl
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

# --- API Configuration ---
OLLAMA_BASE_URL = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "gemma4:latest")

EXCEL_PATH = r"C:\Cursor AI\public\database\RMCMA\RMCMA\RMCMA.xlsx"

def parse_rmcma_content(text):
    if not text:
        return None
    
    parts = [p.strip() for p in re.split(r'\n\s*-+\s*\n|\n-+', text) if p.strip()]
    
    if len(parts) < 3:
        parts = [p.strip() for p in re.split(r'-{3,}', text) if p.strip()]
        if len(parts) < 3:
            if len(parts) == 2:
                passage = parts[0]
                rest = parts[1]
                lines = rest.split('\n')
                choices = []
                question_lines = []
                for line in lines:
                    trimmed = line.strip()
                    if trimmed.startswith('[]') or trimmed.startswith('[x]'):
                        is_correct = trimmed.startswith('[x]')
                        choice_text = trimmed[3:].strip()
                        choices.append({
                            'text': choice_text,
                            'is_correct': is_correct
                        })
                    else:
                        if not choices:
                            question_lines.append(line)
                question = "\n".join(question_lines).strip()
                return {
                    'passage': passage,
                    'question': question,
                    'choices': choices
                }
            return None
            
    passage = parts[0]
    question = parts[1]
    choices_str = parts[2]
    
    choices = []
    for line in choices_str.split('\n'):
        trimmed = line.strip()
        if trimmed.startswith('[]') or trimmed.startswith('[x]'):
            is_correct = trimmed.startswith('[x]')
            choice_text = trimmed[3:].strip()
            choices.append({
                'text': choice_text,
                'is_correct': is_correct
            })
            
    return {
        'passage': passage,
        'question': question,
        'choices': choices
    }

def generate_explanation(parsed_data):
    passage = parsed_data['passage']
    question = parsed_data['question']
    
    options_lines = []
    for c in parsed_data['choices']:
        label = "(Correct)" if c['is_correct'] else "(Incorrect)"
        options_lines.append(f"- {c['text']} {label}")
    options_text = "\n".join(options_lines)
    
    prompt = f"""
You are an expert PTE Academic tutor.
Analyze the following reading passage, question, and options, and provide a clear, detailed, and easy-to-understand explanation of the correct and incorrect answers.

PASSAGE:
{passage}

QUESTION:
{question}

OPTIONS:
{options_text}

Task Instructions:
1. Explain why the correct options are correct by citing relevant information or context from the passage.
2. Explain why each incorrect option is wrong, pointing out where the passage contradicts it or why it is not mentioned/relevant.
3. Write the response in clean HTML format. Use standard HTML tags:
   - Use <p> for paragraphs.
   - Use <strong> for emphasis.
   - Use <ul> and <li> for lists.
   - Do NOT include any outer markdown backticks (such as ```html) or outer wrapping tags like <html>, <body>. Start directly with the HTML content.
   - Make the tone supportive, encouraging, and highly instructional.
"""
    
    url = f"{OLLAMA_BASE_URL}/api/generate"
    
    payload = {
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": 0.2,
            "num_predict": 1024
        }
    }
    
    try:
        response = requests.post(url, json=payload, timeout=90)
        if response.status_code == 200:
            return response.json().get("response", "").strip()
        else:
            print(f"API Error ({response.status_code}): {response.text}")
            return None
    except Exception as e:
        print(f"Exception during request: {e}")
        return None

def process_row(row_idx, q_id, title, content):
    parsed = parse_rmcma_content(content)
    if not parsed:
        return row_idx, q_id, None
    
    explanation = generate_explanation(parsed)
    return row_idx, q_id, explanation

def main():
    print(f"Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active
    
    header_val = sheet.cell(row=1, column=4).value
    if header_val != "EXPLANATION":
        sheet.cell(row=1, column=4, value="EXPLANATION")
        wb.save(EXCEL_PATH)
        print("Set Column D header to 'EXPLANATION'")
    
    max_row = sheet.max_row
    tasks = []
    
    # Identify rows that need processing
    for row_idx in range(2, max_row + 1):
        q_id = sheet.cell(row=row_idx, column=1).value
        title = sheet.cell(row=row_idx, column=2).value
        content = sheet.cell(row=row_idx, column=3).value
        existing_explanation = sheet.cell(row=row_idx, column=4).value
        
        if existing_explanation and len(str(existing_explanation).strip()) > 20:
            continue
            
        tasks.append((row_idx, q_id, title, content))
        
    total_tasks = len(tasks)
    print(f"Pending tasks to enrich: {total_tasks}")
    if total_tasks == 0:
        print("No pending tasks. Enrichment complete.")
        return

    success_count = 0
    fail_count = 0
    
    # We use ThreadPoolExecutor to run tasks concurrently
    with ThreadPoolExecutor(max_workers=8) as executor:
        future_to_row = {
            executor.submit(process_row, row_idx, q_id, title, content): (row_idx, q_id)
            for row_idx, q_id, title, content in tasks
        }
        
        for future in as_completed(future_to_row):
            row_idx, q_id = future_to_row[future]
            try:
                row_idx, q_id, explanation = future.result()
                if explanation:
                    # Reload workbook to avoid conflict with other saves, write, and save
                    wb_write = openpyxl.load_workbook(EXCEL_PATH)
                    sheet_write = wb_write.active
                    sheet_write.cell(row=row_idx, column=4, value=explanation)
                    wb_write.save(EXCEL_PATH)
                    success_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Successfully enriched. ({success_count}/{total_tasks})")
                else:
                    fail_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Failed to generate explanation.")
            except Exception as exc:
                fail_count += 1
                print(f"Row {row_idx} (ID {q_id}) generated an exception: {exc}")
                
    print(f"\nProcessing finished! Success: {success_count}, Failed: {fail_count}")

if __name__ == "__main__":
    main()
