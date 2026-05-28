import argparse
import html
import os
import re
import time

import openpyxl
import requests


OLLAMA_BASE_URL = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "gemma4:latest")
DEFAULT_EXCEL_PATH = r"C:\Cursor AI\public\database\RMCSA\RMCSA\RMCSA.xlsx"

ALLOWED_TAGS = {"p", "strong", "b", "em", "i", "ul", "ol", "li", "br", "h3", "h4"}
CHOICE_RE = re.compile(r"^\[([xX\s]*)\]\s*(.*)$")
DELIMITER_RE = re.compile(r"\r?\n\s*-{3,}\s*\r?\n")
TAG_RE = re.compile(r"</?\s*([a-zA-Z0-9-]+)(?:\s[^>]*)?>")


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Generate and audit RMCSA option explanations in the workbook."
    )
    parser.add_argument(
        "--workbook",
        default=os.environ.get("RMCSA_EXCEL_PATH", DEFAULT_EXCEL_PATH),
        help="Path to the RMCSA workbook.",
    )
    parser.add_argument(
        "--base-url",
        default=OLLAMA_BASE_URL,
        help="Ollama base URL.",
    )
    parser.add_argument(
        "--model",
        default=OLLAMA_MODEL,
        help="Ollama model name.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Maximum number of generated explanations. Omit to process all rows that need work.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Regenerate every parsed row, including rows with valid explanations.",
    )
    parser.add_argument(
        "--fallback-only",
        action="store_true",
        help="Skip Ollama and write deterministic workbook-key explanations.",
    )
    parser.add_argument(
        "--audit-only",
        action="store_true",
        help="Only audit the workbook; exit non-zero if any row fails.",
    )
    args = parser.parse_args(argv)
    if args.limit is not None and args.limit < 1:
        parser.error("--limit must be 1 or greater when provided")
    return args


def normalize_answer_text(text):
    return str(text or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_choice_lines(choices_text):
    choices = []
    for line in normalize_answer_text(choices_text).split("\n"):
        match = CHOICE_RE.match(line.strip())
        if not match:
            continue
        marker = match.group(1).strip().lower()
        choices.append(
            {
                "text": match.group(2).strip(),
                "is_correct": "x" in marker,
            }
        )
    return choices


def parse_rmcsa_content(text):
    normalized = normalize_answer_text(text)
    if not normalized:
        return None

    parts = [part.strip() for part in DELIMITER_RE.split(normalized) if part.strip()]
    if len(parts) < 3:
        parts = [part.strip() for part in re.split(r"-{3,}", normalized) if part.strip()]
    if len(parts) < 3:
        return None

    passage = parts[0]
    question = parts[1]
    choices_raw = "\n".join(parts[2:])
    choices = parse_choice_lines(choices_raw)

    return {
        "passage": passage,
        "question": question,
        "choices": choices,
    }


def strip_html(value):
    plain = re.sub(r"<[^>]+>", " ", str(value or ""))
    return re.sub(r"\s+", " ", html.unescape(plain)).strip()


def unsupported_tags(value):
    tags = set()
    for match in TAG_RE.finditer(str(value or "")):
        tag = match.group(1).lower()
        if tag not in ALLOWED_TAGS:
            tags.add(tag)
    return sorted(tags)


def sanitize_allowed_html(raw_html):
    output = str(raw_html or "").strip()
    output = re.sub(r"^```(?:html)?\s*", "", output, flags=re.IGNORECASE)
    output = re.sub(r"\s*```$", "", output)
    output = re.sub(r"<\s*/?\s*(?:html|body)\s*>", "", output, flags=re.IGNORECASE)
    output = re.sub(r"<\s*script[^>]*>.*?<\s*/\s*script\s*>", "", output, flags=re.IGNORECASE | re.DOTALL)
    output = re.sub(r"\s+on[a-zA-Z]+\s*=\s*(['\"]).*?\1", "", output)
    output = re.sub(r"\s+href\s*=\s*(['\"])\s*javascript:.*?\1", "", output, flags=re.IGNORECASE)
    output = re.sub(r"</?\s*h[1256](?:\s[^>]*)?>", lambda m: "<h3>" if not m.group(0).startswith("</") else "</h3>", output, flags=re.IGNORECASE)

    def replace_tag(match):
        text = match.group(0)
        closing = text.lstrip().startswith("</")
        tag = match.group(1).lower()
        if tag not in ALLOWED_TAGS:
            return ""
        if tag == "br":
            return "<br>"
        return f"</{tag}>" if closing else f"<{tag}>"

    output = TAG_RE.sub(replace_tag, output)
    return output.strip()


def audit_explanation(parsed_data, explanation):
    issues = []
    choices = parsed_data.get("choices") or []
    correct_choices = [choice for choice in choices if choice.get("is_correct")]
    plain = strip_html(explanation)
    lower_plain = plain.lower()

    if len(choices) < 2:
        issues.append("expected at least 2 parsed choices")
    if len(correct_choices) != 1:
        issues.append(f"expected exactly 1 correct choice, found {len(correct_choices)}")
    if len(plain) < 200:
        issues.append(f"explanation is missing or too short ({len(plain)} chars)")
    bad_tags = unsupported_tags(explanation)
    if bad_tags:
        issues.append(f"unsupported HTML tags: {', '.join(bad_tags)}")
    if re.search(r"```|<script|on\w+=|javascript:", str(explanation or ""), flags=re.IGNORECASE):
        issues.append("explanation contains unsafe or markdown artifacts")
    if correct_choices:
        correct_text = correct_choices[0]["text"].lower()
        if correct_text not in lower_plain:
            issues.append(f"explanation does not mention correct answer: {correct_choices[0]['text']}")
    if not re.search(r"\b(incorrect|wrong|not supported|not mentioned|does not|is not|fails|misses)\b", plain, flags=re.IGNORECASE):
        issues.append("explanation does not clearly address incorrect options")
    return issues


def pick_evidence_sentence(passage):
    sentences = re.split(r"(?<=[.!?])\s+", normalize_answer_text(passage))
    candidates = [sentence.strip() for sentence in sentences if len(sentence.strip()) >= 40]
    if not candidates:
        return normalize_answer_text(passage)[:260]
    return max(candidates[:5], key=len)[:260]


def build_fallback_explanation(parsed_data):
    choices = parsed_data["choices"]
    correct = next((choice for choice in choices if choice["is_correct"]), None)
    distractors = [choice for choice in choices if not choice["is_correct"]]
    if not correct:
        return ""

    question = html.escape(parsed_data["question"], quote=False)
    correct_text = html.escape(correct["text"], quote=False)
    evidence = html.escape(pick_evidence_sentence(parsed_data["passage"]), quote=False)

    items = []
    for distractor in distractors:
        option_text = html.escape(distractor["text"], quote=False)
        items.append(
            "<li><strong>{}</strong> is incorrect because the passage does not support it as the answer to this question. "
            "It changes the focus, cause, scope, or detail that the passage actually gives.</li>".format(option_text)
        )

    return (
        "<p><strong>Correct answer:</strong> {correct}</p>"
        "<p>This option is correct for the question <strong>{question}</strong> because it is the only choice that stays aligned "
        "with the passage. The key evidence is: <em>{evidence}</em> This supports the idea in the correct answer without adding "
        "a new claim or shifting the meaning.</p>"
        "<p><strong>Why the other options are incorrect:</strong></p>"
        "<ul>{items}</ul>"
    ).format(correct=correct_text, question=question, evidence=evidence, items="".join(items))


def build_prompt(parsed_data):
    options_lines = []
    for choice in parsed_data["choices"]:
        label = "CORRECT" if choice["is_correct"] else "INCORRECT"
        options_lines.append(f"- {choice['text']} ({label})")

    return f"""You are an expert PTE Academic tutor.
Write a clear option-level explanation for this Reading Multiple Choice Single Answer question.

PASSAGE:
{parsed_data['passage']}

QUESTION:
{parsed_data['question']}

OPTIONS:
{os.linesep.join(options_lines)}

Requirements:
1. Include the exact correct answer text in the explanation.
2. Explain why the correct answer is supported by the passage.
3. Explain why each incorrect option is wrong, not supported, not mentioned, too broad, too narrow, or changes the meaning.
4. Return clean HTML only. Allowed tags: <p>, <strong>, <b>, <em>, <i>, <ul>, <ol>, <li>, <br>, <h3>, <h4>.
5. Do not use <h1>, <h2>, markdown fences, JSON, <html>, or <body>.
"""


def generate_explanation(parsed_data, base_url, model):
    payload = {
        "model": model,
        "prompt": build_prompt(parsed_data),
        "stream": False,
        "options": {
            "temperature": 0.1,
            "num_ctx": 8192,
            "num_predict": 1600,
        },
    }

    try:
        response = requests.post(f"{base_url}/api/generate", json=payload, timeout=120)
        if response.status_code != 200:
            print(f"API Error ({response.status_code}): {response.text}")
            return None
        return response.json().get("response", "").strip()
    except Exception as error:
        print(f"Exception during request: {error}")
        return None


def get_valid_or_fixed_existing(parsed_data, existing_explanation):
    existing = str(existing_explanation or "").strip()
    if not existing:
        return None, ["missing explanation"]

    sanitized = sanitize_allowed_html(existing)
    issues = audit_explanation(parsed_data, sanitized)
    if not issues:
        return sanitized, []
    return sanitized, issues


def main(argv=None):
    args = parse_args(argv)
    print(f"Loading workbook: {args.workbook}...")
    workbook = openpyxl.load_workbook(args.workbook)
    sheet = workbook.active

    if sheet.cell(row=1, column=4).value != "EXPLANATION":
        sheet.cell(row=1, column=4, value="EXPLANATION")
        print("Set Column D header to EXPLANATION")

    generated_count = 0
    sanitized_count = 0
    skipped_count = 0
    failed_count = 0
    audit_failures = []

    for row_idx in range(2, sheet.max_row + 1):
        q_id = sheet.cell(row=row_idx, column=1).value
        title = sheet.cell(row=row_idx, column=2).value
        content = sheet.cell(row=row_idx, column=3).value
        existing_explanation = sheet.cell(row=row_idx, column=4).value
        parsed = parse_rmcsa_content(content)

        if not parsed:
            failed_count += 1
            audit_failures.append(f"Row {row_idx} (ID {q_id}): failed to parse answer content")
            continue

        sanitized_existing, existing_issues = get_valid_or_fixed_existing(parsed, existing_explanation)
        if args.audit_only:
            if existing_issues:
                audit_failures.append(f"Row {row_idx} (ID {q_id}): {'; '.join(existing_issues)}")
            continue

        if not args.force and not existing_issues:
            if sanitized_existing != str(existing_explanation or "").strip():
                sheet.cell(row=row_idx, column=4, value=sanitized_existing)
                sanitized_count += 1
            else:
                skipped_count += 1
            continue

        if args.limit is not None and generated_count >= args.limit:
            skipped_count += 1
            continue

        print(f"Row {row_idx}/{sheet.max_row} (ID {q_id} - {title}): generating explanation...")
        explanation = None
        if not args.fallback_only:
            explanation = generate_explanation(parsed, args.base_url, args.model)
        if not explanation:
            explanation = build_fallback_explanation(parsed)

        explanation = sanitize_allowed_html(explanation)
        issues = audit_explanation(parsed, explanation)
        if issues:
            explanation = sanitize_allowed_html(build_fallback_explanation(parsed))
            issues = audit_explanation(parsed, explanation)

        if issues:
            failed_count += 1
            audit_failures.append(f"Row {row_idx} (ID {q_id}): {'; '.join(issues)}")
            continue

        sheet.cell(row=row_idx, column=4, value=explanation)
        generated_count += 1
        workbook.save(args.workbook)

        if not args.fallback_only:
            time.sleep(0.5)

    if not args.audit_only:
        workbook.save(args.workbook)

    print(
        "Processing finished! "
        f"Generated: {generated_count}, Sanitized: {sanitized_count}, "
        f"Skipped: {skipped_count}, Failed: {failed_count}"
    )
    if audit_failures:
        print("Audit failures:")
        for failure in audit_failures[:25]:
            print(f"- {failure}")
        if len(audit_failures) > 25:
            print(f"- ... {len(audit_failures) - 25} more")
        raise SystemExit(1)
    print("Workbook saved successfully.")


if __name__ == "__main__":
    main()
