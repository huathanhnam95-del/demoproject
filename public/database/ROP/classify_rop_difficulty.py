import os
import sys
import time
import re
import json
import openpyxl
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

sys.stdout.reconfigure(encoding='utf-8')

# --- API Configuration ---
OLLAMA_BASE_URL = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("LOCAL_GEMMA_MODEL", os.environ.get("OLLAMA_MODEL", "gemma4:12b"))

EXCEL_PATH = r"C:\Cursor AI\public\database\ROP\ROP\ROP.xlsx"

def safe_save(wb, path):
    temp_path = path + ".tmp.xlsx"
    wb.save(temp_path)
    os.replace(temp_path, path)

def parse_paragraphs(text):
    if not text:
        return []
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    paragraphs = []
    for line in lines:
        match = re.match(r'^(\d+)[.)\s]\s*(.*)$', line)
        if match:
            paragraphs.append(match.group(2).strip())
        else:
            paragraphs.append(line)
    return paragraphs

def parse_json_response(raw: str):
    """Try to parse JSON from the Ollama response, handling markdown fences."""
    text = raw.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        lines = [l for l in lines if not l.strip().startswith("```")]
        text = "\n".join(lines)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", text, re.DOTALL)
        if m:
            try:
                return json.loads(m.group())
            except json.JSONDecodeError:
                pass
    return None

def classify_difficulty(paragraphs):
    paras_input = "\n".join([f"{i+1}. {p}" for i, p in enumerate(paragraphs)])
    
    prompt = f"""
You are an expert English language assessor specializing in the PTE Academic reading section, specifically "Re-order Paragraphs".
Your task is to analyze the sequential text below and classify its difficulty level as an integer: 1 (Easy), 2 (Medium), or 3 (Hard).

TEXT:
{paras_input}

CLASSIFICATION CRITERIA:
1. Topic Sentence Independence (Easy: Distinct, obvious, independent topic sentence. Hard: Subtle, multiple potential candidates, or relies on background context).
2. Cohesive Devices (Easy: Explicit connectors like "However", "Therefore", "Consequently". Hard: Implicit pronoun referents, synonym chains, or subtle thematic overlaps).
3. Logical & Chronological Flow (Easy: Sequential timeline or process, clear cause-effect. Hard: Complex, abstract academic argumentation, contrasting viewpoints, non-linear flow).
4. Sentence Structure & Vocabulary (Easy: Direct, concise sentences with high-frequency vocabulary. Hard: Long, complex compound sentences with academic or domain-specific terminology).

Please output your evaluation strictly in the following JSON format:
{{
  "reasoning": "A concise explanation (under 60 words) justifying the difficulty classification based on the criteria above.",
  "difficulty_level": 1, 2, or 3
}}

Rules:
- Do NOT wrap in markdown backticks.
- Return raw JSON content only.
- difficulty_level MUST be an integer: 1 (Easy), 2 (Medium), or 3 (Hard).
"""

    url = f"{OLLAMA_BASE_URL}/api/generate"
    payload = {
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {
            "temperature": 0.1,
            "num_ctx": 4096,
            "num_predict": 1024
        }
    }
    
    for attempt in range(1, 4):
        try:
            response = requests.post(url, json=payload, timeout=90)
            if response.status_code == 200:
                res_json = response.json()
                content_text = res_json.get("response", "")
                data = parse_json_response(content_text)
                if data is not None:
                    level = data.get("difficulty_level")
                    if level in [1, 2, 3]:
                        return data
                    # Fallback if level is string or float
                    try:
                        level_int = int(level)
                        if level_int in [1, 2, 3]:
                            data["difficulty_level"] = level_int
                            return data
                    except:
                        pass
                print(f"Attempt {attempt}/3: Invalid JSON structure or level value. Raw: {content_text[:100]}")
            else:
                print(f"Attempt {attempt}/3: API Error ({response.status_code}): {response.text[:100]}")
        except Exception as e:
            print(f"Attempt {attempt}/3: Exception during request: {e}")
        time.sleep(1)
    return None

def worker(row_idx, q_id, answer_text):
    paras = parse_paragraphs(answer_text)
    if not paras:
        return row_idx, None, None
    try:
        result = classify_difficulty(paras)
        if result:
            return row_idx, result.get("difficulty_level"), result.get("reasoning", "")
    except Exception as e:
        print(f"Row {row_idx} worker exception: {e}")
    return row_idx, None, None

def main():
    print(f"Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active
    
    # Check headers
    # Ensure there are at least 6 columns (Column F is index 6)
    c6 = sheet.cell(row=1, column=6).value
    if c6 != "LEVEL":
        sheet.cell(row=1, column=6, value="LEVEL")
        safe_save(wb, EXCEL_PATH)
        print("Added Column F header 'LEVEL'")
        
    max_row = sheet.max_row
    tasks = []
    
    # Identify rows that need difficulty classification
    for row_idx in range(2, max_row + 1):
        q_id = sheet.cell(row=row_idx, column=1).value
        answer_text = sheet.cell(row=row_idx, column=3).value
        existing_level = sheet.cell(row=row_idx, column=6).value
        
        # If level is already 1, 2, or 3, skip
        if existing_level in [1, 2, 3, "1", "2", "3"]:
            continue
            
        tasks.append((row_idx, q_id, answer_text))
        
    total_tasks = len(tasks)
    print(f"Pending tasks to classify: {total_tasks}")
    if total_tasks == 0:
        print("All tasks classified.")
        return

    success_count = 0
    fail_count = 0
    
    # Run requests concurrently using 5 workers
    # 5 workers is safe for local Ollama without heavy context thrashing
    with ThreadPoolExecutor(max_workers=5) as executor:
        futures = {
            executor.submit(worker, row_idx, q_id, answer_text): (row_idx, q_id)
            for row_idx, q_id, answer_text in tasks
        }
        
        for future in as_completed(futures):
            row_idx, q_id = futures[future]
            try:
                row_idx, level, reasoning = future.result()
                if level is not None:
                    # Write to loaded workbook in memory (safe in main thread)
                    sheet.cell(row=row_idx, column=6, value=level)
                    success_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Level {level} - Reasoning: {reasoning}")
                    
                    # Save periodically to prevent losing progress if interrupted
                    if success_count % 5 == 0:
                        safe_save(wb, EXCEL_PATH)
                        print("Saved progress to ROP.xlsx.")
                else:
                    fail_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Failed to classify.")
            except Exception as exc:
                fail_count += 1
                print(f"Row {row_idx} exception: {exc}")
                
    safe_save(wb, EXCEL_PATH)
    print(f"\nDifficulty Classification finished! Success: {success_count}, Failed: {fail_count}")

if __name__ == "__main__":
    main()
