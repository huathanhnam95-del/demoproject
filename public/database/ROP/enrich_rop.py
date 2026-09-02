import os
import sys
import time
import re
import json
import openpyxl
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

sys.stdout.reconfigure(encoding='utf-8')

# --- API Configuration ---
OLLAMA_BASE_URL = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.environ.get("LOCAL_GEMMA_MODEL", os.environ.get("OLLAMA_MODEL", "gemma4:12b"))

EXCEL_PATH = r"C:\Cursor AI\public\database\ROP\ROP\ROP.xlsx"

def safe_save(wb, path):
    temp_path = path + ".tmp.xlsx"
    wb.save(temp_path)
    os.replace(temp_path, path)

def parse_paragraphs(text):
    if not text:
        return []
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    paragraphs = []
    for line in lines:
        match = re.match(r'^(\d+)[.)\s]\s*(.*)$', line)
        if match:
            paragraphs.append(match.group(2).strip())
        else:
            paragraphs.append(line)
    return paragraphs

def parse_json_response(raw: str):
    """Try to parse JSON from the Ollama response, handling markdown fences."""
    text = raw.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        lines = [l for l in lines if not l.strip().startswith("```")]
        text = "\n".join(lines)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", text, re.DOTALL)
        if m:
            try:
                return json.loads(m.group())
            except json.JSONDecodeError:
                pass
    return None

def generate_enrichment(paragraphs):
    paras_input = "\n".join([f"{i+1}. {p}" for i, p in enumerate(paragraphs)])

    prompt = f"""
You are an expert English language tutor specializing in the PTE Academic reading section, specifically "Re-order Paragraphs".
Given the correct sequential order of paragraphs below:

{paras_input}

Task Instructions:
1. Identify 2-4 key cohesive links between the paragraphs. These include:
   - Grammatical cohesion: pronouns (he, she, it, they, these, this, that), demonstratives, substitution (one, the other, doing so), ellipsis.
   - Lexical cohesion: synonyms, keyword repetition, collocations, or transitional words/phrases (however, therefore, thus, in addition).
2. Rewrite the paragraphs by adding cohesive linking span tags. Wrap the cohesive markers (such as pronouns, repeating keywords, transitional words) and their referents in `<span class="cohesion-link" data-link="groupN" data-tooltip="Tooltip description">...</span>` tags where N is a number starting from 1 (e.g. group1, group2) representing each cohesive relationship.
   For example, in paragraph 1 you might wrap:
   "<span class=\"cohesion-link\" data-link=\"group1\" data-tooltip=\"Referenced by: This process\">endothermic reaction</span>"
   And in paragraph 2:
   "<span class=\"cohesion-link\" data-link=\"group1\" data-tooltip=\"Reference to: endothermic reaction\">This process</span>"
3. Do NOT modify the text inside the paragraphs besides wrapping elements in the `<span>` tags. Keep formatting intact.
4. Generate a concise, educational explanation in HTML format (strictly under 200 words) that details the logical flow, focusing on Grammatical Cohesion and Lexical Cohesion that link the paragraphs together (refer to the group link numbers).
5. Output the result strictly in this JSON format:
{{
  "enriched_paragraphs": [
     "Paragraph 1 text with spans...",
     "Paragraph 2 text with spans...",
     ...
  ],
  "explanation": "HTML explanation here..."
}}

Rules:
- Do NOT wrap in markdown backticks.
- Return raw JSON content only.
- Write explanation using standard HTML: <p>, <strong>, <ul>, <li>. Do NOT include <html> or <body> tags. Keep it under 200 words.
- Tone must be encouraging, supportive, and educational.
"""

    url = f"{OLLAMA_BASE_URL}/api/generate"
    payload = {
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {
            "temperature": 0.2,
            "num_ctx": 8192,
            "num_predict": 4096
        }
    }

    for attempt in range(1, 4):
        try:
            response = requests.post(url, json=payload, timeout=180)
            if response.status_code == 200:
                res_json = response.json()
                content_text = res_json.get("response", "")
                data = parse_json_response(content_text)
                if data is not None:
                    return data
                print(f"Attempt {attempt}/3: Failed to parse JSON response. Raw text length: {len(content_text)}")
            else:
                print(f"Attempt {attempt}/3: API Error ({response.status_code}): {response.text}")
        except Exception as e:
            print(f"Attempt {attempt}/3: Exception during request: {e}")
        time.sleep(2)
    return None

def worker(row_idx, q_id, answer_text):
    paras = parse_paragraphs(answer_text)
    if not paras:
        return row_idx, None, None
    try:
        result = generate_enrichment(paras)
        if result:
            enriched_paras = result.get("enriched_paragraphs", [])
            enriched_text = ""
            if enriched_paras and len(enriched_paras) == len(paras):
                cleaned_paras = []
                for p in enriched_paras:
                    p_clean = re.sub(r'^\d+[.)\s]\s*', '', p)
                    cleaned_paras.append(p_clean)
                enriched_text = "\n".join([f"{i+1}. {p}" for i, p in enumerate(cleaned_paras)])
            else:
                enriched_text = "\n".join([f"{i+1}. {p}" for i, p in enumerate(paras)])
            explanation = result.get("explanation", "")
            return row_idx, enriched_text, explanation
    except Exception as e:
        print(f"Row {row_idx} worker exception: {e}")
    return row_idx, None, None

def main():
    print(f"Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active

    # Check headers
    c4 = sheet.cell(row=1, column=4).value
    c5 = sheet.cell(row=1, column=5).value

    if c4 != "ENRICHED_ANSWER":
        sheet.cell(row=1, column=4, value="ENRICHED_ANSWER")
        safe_save(wb, EXCEL_PATH)
        print("Added Column D header 'ENRICHED_ANSWER'")
    if c5 != "EXPLANATION":
        sheet.cell(row=1, column=5, value="EXPLANATION")
        safe_save(wb, EXCEL_PATH)
        print("Added Column E header 'EXPLANATION'")

    max_row = sheet.max_row
    tasks = []

    # Identify rows that need processing
    for row_idx in range(2, max_row + 1):
        q_id = sheet.cell(row=row_idx, column=1).value
        answer_text = sheet.cell(row=row_idx, column=3).value
        existing_explanation = sheet.cell(row=row_idx, column=5).value

        if existing_explanation and len(str(existing_explanation).strip()) > 20:
            continue

        tasks.append((row_idx, q_id, answer_text))

    total_tasks = len(tasks)
    print(f"Pending tasks to enrich: {total_tasks}")
    if total_tasks == 0:
        print("No pending tasks. Enrichment complete.")
        return

    success_count = 0
    fail_count = 0

    # Using 3 workers for heavier explanation task to prevent local Ollama timeout/thrashing
    with ThreadPoolExecutor(max_workers=3) as executor:
        futures = {
            executor.submit(worker, row_idx, q_id, answer_text): (row_idx, q_id)
            for row_idx, q_id, answer_text in tasks
        }

        for future in as_completed(futures):
            row_idx, q_id = futures[future]
            try:
                row_idx, enriched_text, explanation = future.result()
                if explanation and enriched_text:
                    sheet.cell(row=row_idx, column=4, value=enriched_text)
                    sheet.cell(row=row_idx, column=5, value=explanation)
                    success_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Enriched. ({success_count}/{total_tasks})")

                    # Save periodically
                    if success_count % 3 == 0:
                        safe_save(wb, EXCEL_PATH)
                        print("Saved progress to ROP.xlsx.")
                else:
                    fail_count += 1
                    print(f"Row {row_idx}/{max_row} (ID {q_id}): Failed to enrich.")
            except Exception as exc:
                fail_count += 1
                print(f"Row {row_idx} exception: {exc}")

    safe_save(wb, EXCEL_PATH)
    print(f"\nProcessing finished! Success: {success_count}, Failed: {fail_count}")

if __name__ == "__main__":
    main()
