import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

ESSAY_QUESTIONS_DATA = [
    {
        "id": 5,
        "question": "With cities expanding, governments should look forward to creating better networks of public transport available for everyone rather than building more roads for vehicle owning population. To what extent do you agree or disagree?",
        "side1": {
            "opinion": "Opinion: Agree (Đồng ý)",
            "points": [
                {
                    "title": "Saving money for workers and students (Tiết kiệm tiền đi lại cho người lao động và học sinh)",
                    "en": "Cheap bus and train tickets help poor workers and students travel easily without worrying about high fuel costs.",
                    "vi": "Vé xe buýt và tàu điện giá rẻ giúp người lao động nghèo và học sinh, sinh viên đi lại dễ dàng mà không phải lo tiền xăng xe đắt đỏ."
                },
                {
                    "title": "Reducing traffic jams and air pollution (Giảm bớt kẹt xe và khói bụi độc hại)",
                    "en": "When more people use buses, there are fewer motorbikes and cars on the street, which makes the air cleaner and roads less crowded.",
                    "vi": "Khi nhiều người đi xe buýt hơn, đường phố sẽ bớt xe máy và ô tô, giúp không khí trong lành hơn và giảm hẳn tình trạng tắc đường."
                }
            ],
            "vocab": [
                {"en": "cheap bus fare", "vi": "vé xe buýt giá rẻ"},
                {"en": "save travel costs", "vi": "tiết kiệm chi phí đi lại"},
                {"en": "avoid heavy traffic / traffic jam", "vi": "tránh tắc đường / kẹt xe"},
                {"en": "cut down air pollution", "vi": "giảm bớt ô nhiễm không khí"},
                {"en": "eco-friendly lifestyle", "vi": "lối sống thân thiện với môi trường"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disagree (Không đồng ý)",
            "points": [
                {
                    "title": "Buses cannot enter small residential alleys (Xe buýt không thể đi vào các con ngõ nhỏ)",
                    "en": "In many cities, most families live in deep and narrow alleys where buses cannot go, so people still need motorbikes for daily life.",
                    "vi": "Ở nhiều thành phố, người dân sống trong các con ngõ nhỏ và sâu mà xe buýt không vào được, nên mọi người vẫn rất cần xe máy để đi lại hàng ngày."
                },
                {
                    "title": "Need personal vehicles for urgent family needs (Cần xe riêng cho các việc khẩn cấp của gia đình)",
                    "en": "Private motorbikes and cars are very important when someone gets sick at night or when a mother needs to buy heavy groceries for the family.",
                    "vi": "Xe máy hoặc ô tô riêng rất tiện lợi khi gia đình có người ốm đột xuất vào ban đêm, hoặc khi cha mẹ cần đi chợ chở đồ nặng cho cả nhà."
                }
            ],
            "vocab": [
                {"en": "narrow residential alley", "vi": "ngõ nhỏ trong khu dân cư"},
                {"en": "private vehicle", "vi": "phương tiện cá nhân (xe máy, ô tô)"},
                {"en": "medical emergency", "vi": "cấp cứu y tế khẩn cấp ban đêm"},
                {"en": "carry heavy groceries", "vi": "chở đồ đạc / nhu yếu phẩm nặng"},
                {"en": "deliver fresh food", "vi": "vận chuyển thực phẩm tươi sống"}
            ]
        }
    },
    {
        "id": 9,
        "question": "Climate change is a concerning global issue. Who has the main responsibility to take action to solve this problem, governments, large companies, or individuals?",
        "side1": {
            "opinion": "Opinion: Governments and Big Companies (Chính phủ & Các công ty lớn)",
            "points": [
                {
                    "title": "Making strict laws and stopping factory pollution (Ban hành luật nghiêm và phạt các nhà máy xả thải)",
                    "en": "Only governments have the power to make strict laws, ban harmful plastics, and force big factories to stop polluting the air and rivers.",
                    "vi": "Chỉ chính phủ mới có quyền ra luật nghiêm, cấm đồ nhựa độc hại và bắt các nhà máy lớn phải dừng xả khói bụi ra môi trường."
                },
                {
                    "title": "Having enough money to build green energy systems (Có đủ tiền để đầu tư điện mặt trời và năng lượng sạch)",
                    "en": "Big companies make the most smoke and waste, and governments are the only ones with enough money to build solar power and green transport.",
                    "vi": "Các công ty lớn tạo ra phần lớn khói thải độc hại, và chỉ có nhà nước mới có đủ ngân sách để làm điện mặt trời và xe buýt điện cho người dân."
                }
            ],
            "vocab": [
                {"en": "pass strict laws", "vi": "ban hành luật nghiêm khắc"},
                {"en": "ban single-use plastic", "vi": "cấm đồ nhựa dùng một lần"},
                {"en": "heavy industrial pollution", "vi": "ô nhiễm công nghiệp nặng"},
                {"en": "invest in solar power", "vi": "đầu tư vào điện mặt trời"},
                {"en": "green public transport", "vi": "giao thông công cộng xanh / thân thiện"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Individuals and Families (Từng cá nhân & Hộ gia đình)",
            "points": [
                {
                    "title": "Small daily habits save a lot of energy (Những thói quen nhỏ hàng ngày giúp tiết kiệm nhiều năng lượng)",
                    "en": "If every family turns off lights when not in use, wastes less food, and brings cloth bags to the market, it helps the environment a lot.",
                    "vi": "Nếu mỗi nhà đều tắt bớt bóng đèn, không bỏ mứa thức ăn và mang túi vải đi chợ, thì cả xã hội sẽ tiết kiệm được rất nhiều năng lượng."
                },
                {
                    "title": "Buyers choose green products to change market habits (Người mua chọn hàng sạch sẽ buộc các công ty phải thay đổi)",
                    "en": "When people refuse to buy polluting products and choose eco-friendly goods, companies must change the way they make products to stay in business.",
                    "vi": "Khi người dân không mua những món đồ gây ô nhiễm và chỉ chuộng đồ thân thiện với môi trường, các công ty buộc phải sản xuất sạch hơn."
                }
            ],
            "vocab": [
                {"en": "turn off electrical appliances", "vi": "tắt các thiết bị điện khi không dùng"},
                {"en": "bring reusable cloth bags", "vi": "mang túi vải dùng nhiều lần đi chợ"},
                {"en": "reduce food waste", "vi": "giảm lãng phí đồ ăn trong nhà"},
                {"en": "eco-friendly products", "vi": "các sản phẩm thân thiện với môi trường"},
                {"en": "change daily shopping habits", "vi": "thay đổi thói quen mua sắm hàng ngày"}
            ]
        }
    },
    {
        "id": 17,
        "question": "Many education systems assess students' learning using formal written examinations. Those kinds of exams are a valid method. To what extent do you agree or disagree? Give examples with your own experience.",
        "side1": {
            "opinion": "Opinion: Agree (Đồng ý)",
            "points": [
                {
                    "title": "Fair grading without teacher favoritism (Chấm điểm công bằng, không bị thầy cô thiên vị)",
                    "en": "All students get the same exam questions and their names are hidden on the papers, so teachers grade everyone fairly without personal feelings.",
                    "vi": "Mọi học sinh đều làm chung một đề thi và bài thi được rọc phách giấu tên, nên thầy cô sẽ chấm điểm rất công bằng, không thiên vị ai."
                },
                {
                    "title": "Pushing students to study and test their own knowledge (Thúc đẩy học sinh tự ôn bài và kiểm tra kiến thức thật)",
                    "en": "Written tests give students a clear reason to review their lessons carefully and show what each student truly understands on their own.",
                    "vi": "Kỳ thi giúp học sinh có động lực ngồi vào bàn ôn lại bài vở, đồng thời kiểm tra được thực chất mỗi em tự làm được những gì mà không nhìn bài bạn."
                }
            ],
            "vocab": [
                {"en": "fair grading standard", "vi": "thang chấm điểm công bằng"},
                {"en": "anonymous exam paper", "vi": "bài thi rọc phách ẩn tên"},
                {"en": "avoid teacher favoritism", "vi": "tránh việc thầy cô thiên vị"},
                {"en": "revise lessons carefully", "vi": "ôn tập bài vở kỹ lưỡng"},
                {"en": "test independent knowledge", "vi": "kiểm tra kiến thức tự lực của học sinh"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disagree (Không đồng ý)",
            "points": [
                {
                    "title": "Encouraging students to memorize without understanding (Học sinh dễ học vẹt đối phó rồi quên hết sau khi thi)",
                    "en": "Students often just try to remember facts and numbers for a few days to pass the exam, and then forget everything soon after the test is over.",
                    "vi": "Học sinh thường chỉ cố học thuộc lòng số liệu và câu chữ trong vài ngày để qua môn, thi xong là quên sạch kiến thức."
                },
                {
                    "title": "Creating too much stress and ignoring practical skills (Gây áp lực căng thẳng và không đo được kỹ năng thực tế)",
                    "en": "Many students feel very nervous and get sick on exam day, and a paper test cannot check important life skills like teamwork or speaking.",
                    "vi": "Nhiều bạn bị tâm lý sợ hãi, thậm chí phát ốm vào ngày thi, hơn nữa bài thi trên giấy không thể kiểm tra được kỹ năng giao tiếp hay làm việc nhóm."
                }
            ],
            "vocab": [
                {"en": "rote memorization / learn by heart", "vi": "học vẹt / học thuộc lòng đối phó"},
                {"en": "forget quickly after the test", "vi": "nhanh chóng quên hết sau khi thi xong"},
                {"en": "severe exam stress / anxiety", "vi": "áp lực / lo sợ thi cử nặng nề"},
                {"en": "practical hands-on skills", "vi": "kỹ năng thực hành ngoài đời thực"},
                {"en": "teamwork and communication", "vi": "kỹ năng làm việc nhóm và giao tiếp"}
            ]
        }
    },
    {
        "id": 24,
        "question": "The information revolution brought about by modern mass communications has both positive and negative consequences for individuals and society. To what extent do you agree with this statement? Give the reason with your own experience.",
        "side1": {
            "opinion": "Opinion: Positive Effects (Mặt tích cực / Lợi ích)",
            "points": [
                {
                    "title": "Calling family and friends for free anytime (Gọi điện và nhìn thấy người thân ở xa hoàn toàn miễn phí)",
                    "en": "People working far from home can easily make free video calls to talk with their parents and children every day, keeping families close.",
                    "vi": "Những người đi làm xa quê hay du học sinh có thể gọi video miễn phí về nhà mỗi ngày để trò chuyện với cha mẹ, con cái, giúp tình cảm gia đình luôn gắn bó."
                },
                {
                    "title": "Learning new skills and finding jobs online easily (Tự học kỹ năng mới và tìm việc làm trên mạng rất dễ)",
                    "en": "Anyone can watch free video lessons to learn cooking, English, or job skills at home, and find part-time jobs quickly on phone apps.",
                    "vi": "Bất kỳ ai cũng có thể tự học tiếng Anh, học nấu ăn hay học nghề miễn phí trên mạng, đồng thời tìm việc làm thêm nhanh chóng qua điện thoại."
                }
            ],
            "vocab": [
                {"en": "free video call", "vi": "cuộc gọi video miễn phí"},
                {"en": "stay in touch with family", "vi": "giữ liên lạc với người thân ở xa"},
                {"en": "free online learning", "vi": "học tập trực tuyến miễn phí"},
                {"en": "learn job skills at home", "vi": "tự học kỹ năng nghề nghiệp tại nhà"},
                {"en": "find part-time jobs easily", "vi": "dễ dàng tìm kiếm việc làm thêm"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Negative Effects (Mặt tiêu cực / Bất lợi)",
            "points": [
                {
                    "title": "Fake news and online scams cause panic and loss of money (Tin đồn giả và lừa đảo trên mạng gây hoang mang, mất tiền)",
                    "en": "Untrue health rumors and bad news spread very fast online, scaring elderly people and making many people lose money to online tricks.",
                    "vi": "Tin đồn nhảm về sức khỏe và các chiêu trò lừa đảo qua mạng lan truyền rất nhanh, khiến người già sợ hãi và nhiều người bị lừa mất tiền oan."
                },
                {
                    "title": "Spending too much time on screens and talking less in person (Dán mắt vào điện thoại quá nhiều và ít trò chuyện trực tiếp)",
                    "en": "People are glued to their phones all day and receive work messages at night, so they feel tired and talk much less to family members at the dinner table.",
                    "vi": "Nhiều người mải bấm điện thoại cả ngày và bị nhắn tin công việc lúc tối muộn, khiến đầu óc mệt mỏi và bữa cơm gia đình không còn ai trò chuyện với ai."
                }
            ],
            "vocab": [
                {"en": "fake news and rumors", "vi": "tin tức giả và tin đồn thất thiệt"},
                {"en": "online scams / tricks", "vi": "các chiêu trò lừa đảo qua mạng"},
                {"en": "screen addiction / glued to screens", "vi": "nghiện màn hình / dán mắt vào điện thoại"},
                {"en": "mental tiredness / burnout", "vi": "mệt mỏi / kiệt sức tinh thần"},
                {"en": "lack of face-to-face talk", "vi": "thiếu trò chuyện trực tiếp ngoài đời"}
            ]
        }
    },
    {
        "id": 30,
        "question": "In many towns and cities large shopping malls are replacing small local shops. Some people say that this is a positive development. To what extent do you agree?",
        "side1": {
            "opinion": "Opinion: Agree / Positive (Đồng ý / Tích cực)",
            "points": [
                {
                    "title": "Clean, cool place with everything under one roof (Chỗ vui chơi sạch sẽ, mát mẻ, có đủ mọi thứ ở một nơi)",
                    "en": "Malls have air-conditioning, supermarkets, restaurants, and play areas together, so families can shop and have fun comfortably on hot weekends.",
                    "vi": "Trung tâm thương mại có máy lạnh mát mẻ, có cả siêu thị, quán ăn và khu vui chơi, giúp cả nhà vừa đi sắm đồ vừa thư giãn thoải mái vào cuối tuần."
                },
                {
                    "title": "Clear prices and safe, clean food (Giá cả niêm yết rõ ràng và đồ ăn sạch sẽ, an toàn)",
                    "en": "All goods in shopping malls have printed price tags so buyers do not have to bargain, and fresh food is inspected carefully before being sold.",
                    "vi": "Hàng hóa trong siêu thị đều có tem giá rõ ràng nên không lo bị nói thách hay phải trả giá, thức ăn tươi sống cũng được kiểm tra vệ sinh kỹ càng."
                }
            ],
            "vocab": [
                {"en": "air-conditioned shopping mall", "vi": "trung tâm thương mại có máy lạnh mát mẻ"},
                {"en": "all-in-one entertainment", "vi": "khu vui chơi giải trí tất-cả-trong-một"},
                {"en": "fixed price tag", "vi": "tem giá niêm yết cố định (không phải mặc cả)"},
                {"en": "fresh and safe food", "vi": "thực phẩm tươi sạch và an toàn"},
                {"en": "family weekend outing", "vi": "chuyến đi chơi cuối tuần của cả gia đình"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disagree / Negative (Không đồng ý / Tiêu cực)",
            "points": [
                {
                    "title": "Small family shops lose their business and income (Các tiệm tạp hóa gia đình bị mất khách và khó kiếm sống)",
                    "en": "Big supermarkets take away customers from small neighborhood shops, making it hard for poor shopkeepers to earn enough money to support their families.",
                    "vi": "Các siêu thị lớn hút hết khách của các quán tạp hóa trong xóm, khiến các hộ kinh doanh nhỏ gặp khó khăn và mất đi nguồn sống nuôi gia đình."
                },
                {
                    "title": "Hard for elderly neighbors and loss of friendly connections (Người già khó đi xa mua đồ và mất đi tình làng nghĩa xóm)",
                    "en": "Old people can easily walk a few steps to a nearby corner shop to buy salt or vegetables, while big malls are too far and lack warm neighborly chats.",
                    "vi": "Các cụ già chỉ cần đi bộ vài bước ra đầu ngõ là mua được mắm muối, rau củ; nếu tiệm đóng cửa thì phải đi xa rất mệt và mất đi thói quen chào hỏi xóm giềng thân quen."
                }
            ],
            "vocab": [
                {"en": "small corner shop / grocery", "vi": "tiệm tạp hóa nhỏ ở góc phố / đầu ngõ"},
                {"en": "lose daily customers", "vi": "mất khách quen mua sắm hàng ngày"},
                {"en": "hard for elderly neighbors", "vi": "khó khăn cho các cụ già đi lại"},
                {"en": "friendly neighborhood chat", "vi": "cuộc trò chuyện xóm giềng thân tình"},
                {"en": "support family living costs", "vi": "kiếm tiền trang trải cuộc sống gia đình"}
            ]
        }
    },
    {
        "id": 35,
        "question": "The mass media, such as TV, radio and newspapers, have an influence on people, particularly on younger generations. It plays a pivotal role in shaping the opinions of people, especially teenagers and young people. To what extent do you agree with this? Please give examples.",
        "side1": {
            "opinion": "Opinion: Agree (Đồng ý)",
            "points": [
                {
                    "title": "Copying fashion, food, and shopping trends online (Bắt chước cách ăn mặc, ăn uống và mua sắm theo trào lưu trên mạng)",
                    "en": "Teenagers spend hours watching videos every day, so they quickly copy the clothes, hairstyles, and milk tea brands that online idols show.",
                    "vi": "Các bạn trẻ lướt mạng nhiều giờ mỗi ngày nên rất dễ bắt chước cách ăn mặc, kiểu tóc hay các món trà sữa theo những người nổi tiếng trên mạng."
                },
                {
                    "title": "Changing how young people talk and what jobs they want (Thay đổi cách ăn nói và ước mơ nghề nghiệp của người trẻ)",
                    "en": "The media spreads new slang words and makes online streaming look very easy and rich, so many teenagers prefer becoming video creators over normal jobs.",
                    "vi": "Phim ảnh và mạng xã hội tạo ra nhiều từ lóng mới, đồng thời tô vẽ nghề làm video trên mạng kiếm tiền dễ, khiến nhiều bạn trẻ chỉ thích làm TikToker thay vì học nghề đàng hoàng."
                }
            ],
            "vocab": [
                {"en": "social media influencer / idol", "vi": "người có sức ảnh hưởng trên mạng xã hội"},
                {"en": "copy fashion trends", "vi": "bắt chước theo các trào lưu thời trang"},
                {"en": "popular slang words", "vi": "những từ lóng phổ biến trên mạng"},
                {"en": "glamorous online jobs", "vi": "các nghề nghiệp ảo hào nhoáng trên mạng"},
                {"en": "peer pressure", "vi": "áp lực muốn giống bạn bè / theo số đông"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disagree (Không đồng ý)",
            "points": [
                {
                    "title": "Parents and teachers teach the most important life values (Cha mẹ và thầy cô mới là người dạy các nếp sống quan trọng nhất)",
                    "en": "Everyday advice from parents about being honest and polite has a much deeper and lasting effect on a child than short funny videos on the internet.",
                    "vi": "Lời dạy bảo hàng ngày của cha mẹ về tính thật thà, lễ phép có ảnh hưởng sâu đậm và lâu dài đến tính cách của con hơn là những clip ngắn xem trên mạng."
                },
                {
                    "title": "Young people are smart enough to recognize fake advertising (Giới trẻ ngày nay đủ khôn khéo để nhận ra quảng cáo thổi phồng)",
                    "en": "Many teenagers know that advertisements online are exaggerated to sell things, so they do not believe everything they see on the screen.",
                    "vi": "Rất nhiều bạn trẻ ngày nay đủ thông minh để nhận ra các chiêu trò quảng cáo nói quá sự thật để bán hàng, nên các bạn không dễ bị lừa hay tin theo mù quáng."
                }
            ],
            "vocab": [
                {"en": "moral lessons from parents", "vi": "bài học đạo đức / nếp sống từ cha mẹ"},
                {"en": "core values (honesty, kindness)", "vi": "các giá trị cốt lõi (thật thà, tốt bụng)"},
                {"en": "teacher guidance at school", "vi": "sự chỉ dẫn, uốn nắn của thầy cô"},
                {"en": "real-life experience", "vi": "trải nghiệm thực tế ngoài đời"},
                {"en": "recognize exaggerated ads", "vi": "nhận ra quảng cáo thổi phồng / nói quá"}
            ]
        }
    },
    {
        "id": 39,
        "question": "Nowadays, it is increasingly more difficult to maintain the right balance between work and the other aspects of one's life, such as leisure pursuits with family members. How important do you think this balance is and what are the reasons why some people think that this is hard to achieve？",
        "side1": {
            "opinion": "Opinion: Balance Is Very Important (Cân bằng là rất quan trọng)",
            "points": [
                {
                    "title": "Keeping the body and mind healthy without getting sick (Giữ cho cơ thể khỏe mạnh, tránh kiệt sức và đổ bệnh)",
                    "en": "Getting enough rest and having free time stops workers from feeling exhausted, helping them avoid headaches, poor sleep, and depression.",
                    "vi": "Được nghỉ ngơi đầy đủ giúp người đi làm không bị kiệt sức, tránh được chứng mất ngủ, đau đầu và bệnh trầm cảm do căng thẳng kéo dài."
                },
                {
                    "title": "Spending happy time with children and family (Dành thời gian ở bên con cái và chăm lo cho gia đình)",
                    "en": "Children grow up quickly and need parents to play and talk with them, while eating dinner together keeps the husband and wife happy.",
                    "vi": "Con cái lớn lên rất nhanh và rất cần cha mẹ ở bên trò chuyện, chơi đùa; những bữa cơm gia đình ấm cúng cũng giúp vợ chồng gắn bó và yêu thương nhau hơn."
                }
            ],
            "vocab": [
                {"en": "take regular rest", "vi": "nghỉ ngơi điều độ hàng ngày"},
                {"en": "avoid mental burnout", "vi": "tránh kiệt sức / quá tải tinh thần"},
                {"en": "spend quality time with kids", "vi": "dành thời gian trọn vẹn bên con cái"},
                {"en": "happy family dinner", "vi": "bữa cơm gia đình đầm ấm"},
                {"en": "recharge daily energy", "vi": "nạp lại năng lượng sau giờ làm"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Reasons Why It Is Hard to Achieve (Những lý do khó đạt được)",
            "points": [
                {
                    "title": "High living costs force people to work extra hours (Vật giá đắt đỏ buộc mọi người phải làm thêm giờ để kiếm tiền)",
                    "en": "House rents, food, and school fees for children are rising fast, so parents must work overtime or do a second job just to pay the monthly bills.",
                    "vi": "Tiền thuê nhà, tiền ăn và tiền học cho con cái ngày càng đắt đỏ, nên cha mẹ phải tăng ca hoặc nhận thêm việc phụ mới đủ tiền trang trải cuộc sống."
                },
                {
                    "title": "Work messages on phones follow people home at night (Tin nhắn công việc trói buộc người ta ngay cả lúc ở nhà)",
                    "en": "Because bosses and customers can send chat messages at any time, workers feel pressured to answer emails late at night and cannot truly relax.",
                    "vi": "Vì sếp và khách hàng có thể nhắn tin bất kỳ lúc nào, người lao động luôn phải canh điện thoại để trả lời công việc lúc tối muộn, không thể thảnh thơi nghỉ ngơi."
                }
            ],
            "vocab": [
                {"en": "rising cost of living", "vi": "chi phí sinh hoạt ngày càng đắt đỏ"},
                {"en": "work overtime / extra hours", "vi": "làm thêm giờ / tăng ca"},
                {"en": "pay monthly bills and rent", "vi": "trả tiền hóa đơn và tiền thuê nhà hàng tháng"},
                {"en": "late-night work messages", "vi": "tin nhắn công việc lúc tối muộn"},
                {"en": "fear of losing one's job", "vi": "nỗi sợ bị mất việc làm / sa thải"}
            ]
        }
    },
    {
        "id": 43,
        "question": "Should parents be held legally responsible for the actions of their children? Support your opinion from your study, observations or experiences.",
        "side1": {
            "opinion": "Opinion: Agree / Yes (Đồng ý - Cha mẹ nên chịu trách nhiệm)",
            "points": [
                {
                    "title": "Parents buy the phones and motorbikes used by their kids (Cha mẹ là người cho tiền và mua xe nên phải bồi thường nếu con gây hại)",
                    "en": "Underage kids have no money of their own, so when a teenager damages a neighbor's property or motorbike, the parents who give them money must pay for the repair.",
                    "vi": "Con nhỏ chưa làm ra tiền, nên khi con phá hỏng đồ đạc hay xe cộ của người khác, cha mẹ là người nuôi dưỡng và mua sắm đồ cho con thì phải đứng ra đền tiền."
                },
                {
                    "title": "Making parents care more and supervise their children (Buộc cha mẹ phải quan tâm và để mắt đến con cái nhiều hơn)",
                    "en": "If parents know they can be fined by the police, they will pay closer attention to where their teenagers go after school and who they hang out with.",
                    "vi": "Nếu biết mình sẽ bị phạt tiền, các bậc phụ huynh lơ là sẽ phải có trách nhiệm hơn, để mắt xem con mình đi đâu và chơi với những bạn bè nào sau giờ học."
                }
            ],
            "vocab": [
                {"en": "pay for damaged property", "vi": "bồi thường tài sản bị đập phá / làm hỏng"},
                {"en": "give pocket money and motorbikes", "vi": "cho tiền tiêu vặt và mua xe máy cho con"},
                {"en": "supervise teenagers closely", "vi": "giám sát con cái tuổi mới lớn chặt chẽ"},
                {"en": "parental duty and care", "vi": "trách nhiệm và sự quan tâm của cha mẹ"},
                {"en": "prevent bad juvenile behavior", "vi": "ngăn chặn các hành vi xấu của trẻ vị thành niên"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disagree / No (Không đồng ý - Trẻ tự chịu trách nhiệm)",
            "points": [
                {
                    "title": "Teenagers have their own minds and follow bad friends outside (Thanh thiếu niên đã lớn, có ý thức riêng và dễ bị bạn xấu rủ rê)",
                    "en": "Older teenagers make their own choices away from home and often do silly things because they want to show off to friends, which parents cannot foresee.",
                    "vi": "Các bạn tuổi mới lớn đã có suy nghĩ riêng và hay làm trò quậy phá vì muốn chứng tỏ với bạn bè ngoài đường, điều này cha mẹ ở nhà không thể lường trước được."
                },
                {
                    "title": "Working parents cannot watch kids 24 hours a day (Cha mẹ bận đi làm kiếm sống không thể canh chừng con cả ngày)",
                    "en": "Many poor parents work twelve hours a day just to buy food; punishing them with heavy fines is unfair and pushes poor families into debt.",
                    "vi": "Nhiều phụ huynh nghèo phải đi làm quần quật cả ngày để kiếm cơm ăn; phạt tiền nặng sẽ rất bất công và đẩy những gia đình khốn khó vào cảnh nợ nần."
                }
            ],
            "vocab": [
                {"en": "independent mind / free will", "vi": "suy nghĩ độc lập / ý muốn riêng của con"},
                {"en": "bad peer influence outside", "vi": "ảnh hưởng xấu từ bạn bè bên ngoài"},
                {"en": "show off to friends", "vi": "thích thể hiện / khoe mẽ với bạn bè"},
                {"en": "work all day to earn food", "vi": "đi làm quần quật cả ngày để kiếm sống"},
                {"en": "unfair heavy fines", "vi": "các khoản tiền phạt nặng nề, bất công"}
            ]
        }
    },
    {
        "id": 46,
        "question": "In some companies, employers involve workers in the decision-making process about products and services. What are the advantages and disadvantages of such a policy?",
        "side1": {
            "opinion": "Opinion: Advantages (Ưu điểm / Mặt tốt)",
            "points": [
                {
                    "title": "Workers talk to customers every day and know what to fix (Nhân viên tiếp xúc với khách hàng mỗi ngày nên biết rõ cần sửa đổi gì)",
                    "en": "Sales staff and waiters hear direct complaints from customers every day, so they have great practical ideas to improve products that the manager does not know.",
                    "vi": "Nhân viên bán hàng và phục vụ nghe khách phàn nàn trực tiếp mỗi ngày, nên họ biết chính xác món ăn hay sản phẩm cần sửa chỗ nào mà các sếp ngồi văn phòng không thấy."
                },
                {
                    "title": "Workers feel respected and work harder for the company (Nhân viên thấy mình được tôn trọng sẽ yêu thích công việc và gắn bó lâu hơn)",
                    "en": "When bosses listen to their ideas, workers feel proud and happy at work, so they try their best and stay with the company longer instead of quitting.",
                    "vi": "Khi được sếp lắng nghe ý kiến, người lao động cảm thấy mình được coi trọng, từ đó đi làm vui vẻ hơn, nhiệt tình hơn và gắn bó lâu dài chứ không muốn nhảy việc."
                }
            ],
            "vocab": [
                {"en": "frontline workers", "vi": "nhân viên trực tiếp bán hàng / phục vụ"},
                {"en": "listen to customer complaints", "vi": "lắng nghe lời phàn nàn của khách hàng"},
                {"en": "practical improvement ideas", "vi": "những ý tưởng cải tiến thực tế"},
                {"en": "feel respected and valued", "vi": "cảm thấy mình được tôn trọng và coi trọng"},
                {"en": "stay loyal to the company", "vi": "gắn bó trung thành, không muốn nhảy việc"}
            ]
        },
        "side2": {
            "opinion": "Opinion: Disadvantages (Nhược điểm / Mặt xấu)",
            "points": [
                {
                    "title": "Too many meetings waste time and slow down decisions (Họp hành quá nhiều làm mất thời gian và chậm trễ công việc)",
                    "en": "Asking everyone for opinions takes a lot of time and leads to long meetings, which slows down urgent store decisions and takes time away from serving customers.",
                    "vi": "Hỏi ý kiến của từng người sẽ làm các cuộc họp kéo dài lê thê, làm chậm trễ các quyết định bán hàng gấp và mất thời gian phục vụ khách."
                },
                {
                    "title": "Disagreements and hurt feelings when an idea is rejected (Dễ cãi cọ và bất hòa khi ý kiến của người này được chọn còn người kia bị từ chối)",
                    "en": "If the boss chooses one worker's suggestion and says no to others, some employees might feel jealous or hurt, which can cause bad feelings in the team.",
                    "vi": "Nếu người chủ chọn ý kiến của một bạn mà bác bỏ ý kiến của các bạn khác, nhân viên dễ sinh ra tị nạnh, giận dỗi và làm việc nhóm mất vui."
                }
            ],
            "vocab": [
                {"en": "long, time-wasting meetings", "vi": "các cuộc họp dài lê thê làm mất thời gian"},
                {"en": "slow down urgent decisions", "vi": "làm chậm trễ các quyết định khẩn cấp"},
                {"en": "unrealistic budget ideas", "vi": "ý tưởng tốn kém vượt quá ngân sách cửa hàng"},
                {"en": "interpersonal jealousy", "vi": "sự tị nạnh / ghen tị giữa các nhân viên"},
                {"en": "hurt feelings and team division", "vi": "tổn thương lòng tự ái và chia rẽ nội bộ"}
            ]
        }
    }
]

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set cell padding in twips (1 pt = 20 twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4"):
    """Set crisp table borders matching image."""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def render_opinion_block(cell, opinion_data):
    """Render an opinion block with Opinion label, Points, and Suggested Vocabulary."""
    # 1. Opinion line
    p_op = cell.paragraphs[0]
    p_op.paragraph_format.space_before = Pt(1)
    p_op.paragraph_format.space_after = Pt(3)
    p_op.paragraph_format.line_spacing = 1.15
    
    run_op = p_op.add_run(opinion_data["opinion"])
    run_op.bold = True
    run_op.font.name = "Arial"
    run_op.font.size = Pt(10)
    run_op.font.color.rgb = RGBColor(0, 0, 0)
    
    # 2. Points
    for idx, pt in enumerate(opinion_data["points"], 1):
        # Point label & title
        p_pt = cell.add_paragraph()
        p_pt.paragraph_format.space_before = Pt(2)
        p_pt.paragraph_format.space_after = Pt(1)
        p_pt.paragraph_format.line_spacing = 1.15
        
        run_lbl = p_pt.add_run(f"Point {idx}: ")
        run_lbl.bold = True
        run_lbl.font.name = "Arial"
        run_lbl.font.size = Pt(9.5)
        run_lbl.font.color.rgb = RGBColor(0, 0, 0)
        
        run_title = p_pt.add_run(pt["title"])
        run_title.font.name = "Arial"
        run_title.font.size = Pt(9.5)
        run_title.font.color.rgb = RGBColor(17, 24, 39)
        
        # EN detail
        p_en = cell.add_paragraph()
        p_en.paragraph_format.left_indent = Inches(0.15)
        p_en.paragraph_format.space_before = Pt(0)
        p_en.paragraph_format.space_after = Pt(1)
        p_en.paragraph_format.line_spacing = 1.15
        
        tag_en = p_en.add_run("• EN: ")
        tag_en.bold = True
        tag_en.font.name = "Arial"
        tag_en.font.size = Pt(9)
        tag_en.font.color.rgb = RGBColor(37, 99, 235)
        
        txt_en = p_en.add_run(pt["en"])
        txt_en.font.name = "Arial"
        txt_en.font.size = Pt(9)
        txt_en.font.color.rgb = RGBColor(31, 41, 55)
        
        # VN detail
        p_vn = cell.add_paragraph()
        p_vn.paragraph_format.left_indent = Inches(0.15)
        p_vn.paragraph_format.space_before = Pt(0)
        p_vn.paragraph_format.space_after = Pt(2)
        p_vn.paragraph_format.line_spacing = 1.15
        
        tag_vn = p_vn.add_run("• VN: ")
        tag_vn.bold = True
        tag_vn.font.name = "Arial"
        tag_vn.font.size = Pt(9)
        tag_vn.font.color.rgb = RGBColor(180, 83, 9)
        
        txt_vn = p_vn.add_run(pt["vi"])
        txt_vn.font.name = "Arial"
        txt_vn.font.size = Pt(9)
        txt_vn.font.color.rgb = RGBColor(75, 85, 99)

    # 3. Suggested Vocabulary
    if "vocab" in opinion_data and opinion_data["vocab"]:
        p_v_hdr = cell.add_paragraph()
        p_v_hdr.paragraph_format.space_before = Pt(3)
        p_v_hdr.paragraph_format.space_after = Pt(1)
        p_v_hdr.paragraph_format.line_spacing = 1.15
        
        run_v_hdr = p_v_hdr.add_run("Suggested Vocabulary (Từ vựng gợi ý):")
        run_v_hdr.bold = True
        run_v_hdr.font.name = "Arial"
        run_v_hdr.font.size = Pt(9)
        run_v_hdr.font.color.rgb = RGBColor(79, 70, 229)
        
        for v in opinion_data["vocab"]:
            p_v = cell.add_paragraph()
            p_v.paragraph_format.left_indent = Inches(0.15)
            p_v.paragraph_format.space_before = Pt(0)
            p_v.paragraph_format.space_after = Pt(1)
            p_v.paragraph_format.line_spacing = 1.12
            
            bullet = p_v.add_run("• ")
            bullet.bold = True
            bullet.font.name = "Arial"
            bullet.font.size = Pt(8.5)
            bullet.font.color.rgb = RGBColor(99, 102, 241)
            
            run_term = p_v.add_run(f"{v['en']}: ")
            run_term.bold = True
            run_term.font.name = "Arial"
            run_term.font.size = Pt(8.5)
            run_term.font.color.rgb = RGBColor(17, 24, 39)
            
            run_vi = p_v.add_run(v['vi'])
            run_vi.font.name = "Arial"
            run_vi.font.size = Pt(8.5)
            run_vi.font.color.rgb = RGBColor(75, 85, 99)

def build_docx(output_path, orientation="portrait"):
    doc = docx.Document()
    
    # Configure Section margins
    section = doc.sections[0]
    if orientation == "portrait":
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        # Total usable width = 7.5 in
        col_widths = [Inches(0.45), Inches(2.65), Inches(4.40)]
    else:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.0)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        col_widths = [Inches(0.55), Inches(3.45), Inches(6.00)]
    
    # The table starts immediately at top of page, matching Google Docs / Word table format
    num_rows = 1 + len(ESSAY_QUESTIONS_DATA) * 2
    table = doc.add_table(rows=num_rows, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="4")
    
    # Header Row: ID | Question | Ideas
    hdr_row = table.rows[0]
    hdr_titles = ["ID", "Question", "Ideas"]
    for i, title in enumerate(hdr_titles):
        cell = hdr_row.cells[i]
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if i == 0:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
    # Repeat header on every page
    header_tr = hdr_row._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    
    # Fill Question Rows
    for q_idx, item in enumerate(ESSAY_QUESTIONS_DATA):
        r1_idx = 1 + q_idx * 2
        r2_idx = r1_idx + 1
        
        row1 = table.rows[r1_idx]
        row2 = table.rows[r2_idx]
        
        # ID Cell in Row 1
        cell_id_1 = row1.cells[0]
        set_cell_margins(cell_id_1, top=120, bottom=120, left=80, right=80)
        p_id = cell_id_1.paragraphs[0]
        p_id.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_id.paragraph_format.space_before = Pt(2)
        p_id.paragraph_format.space_after = Pt(2)
        run_id = p_id.add_run(str(item["id"]))
        run_id.font.name = "Arial"
        run_id.font.size = Pt(10)
        run_id.font.color.rgb = RGBColor(0, 0, 0)
        cell_id_1.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        
        # Question Cell in Row 1
        cell_q_1 = row1.cells[1]
        set_cell_margins(cell_q_1, top=120, bottom=120, left=120, right=120)
        p_q = cell_q_1.paragraphs[0]
        p_q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_q.paragraph_format.space_before = Pt(2)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.line_spacing = 1.15
        run_q = p_q.add_run(item["question"])
        run_q.font.name = "Arial"
        run_q.font.size = Pt(9.5)
        run_q.font.color.rgb = RGBColor(0, 0, 0)
        cell_q_1.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        
        # Ideas Cell in Row 1 (Side 1)
        cell_ideas_1 = row1.cells[2]
        set_cell_margins(cell_ideas_1, top=120, bottom=120, left=140, right=140)
        cell_ideas_1.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        render_opinion_block(cell_ideas_1, item["side1"])
        
        # Ideas Cell in Row 2 (Side 2)
        cell_ideas_2 = row2.cells[2]
        set_cell_margins(cell_ideas_2, top=120, bottom=120, left=140, right=140)
        cell_ideas_2.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        render_opinion_block(cell_ideas_2, item["side2"])
        
        # Merge ID and Question across Row 1 and Row 2
        cell_id_1.merge(row2.cells[0])
        cell_q_1.merge(row2.cells[1])
        
        # Ensure row cannot split across pages mid-cell if possible
        for r in [row1, row2]:
            trPr = r._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            
    # Apply column widths across all cells
    for row in table.rows:
        for c_idx, width in enumerate(col_widths):
            row.cells[c_idx].width = width
            
    doc.save(output_path)
    print(f"Successfully generated docx file at: {output_path} ({orientation})")

if __name__ == "__main__":
    out_dir = os.path.join("public", "database", "Write Essay", "ESSAY")
    os.makedirs(out_dir, exist_ok=True)
    
    # 1. Primary files in Portrait (exact layout as shown in user's screenshot)
    target_path = os.path.join(out_dir, "PTE_Essay_Bilingual_Argument_Ideas_Table.docx")
    build_docx(target_path, orientation="portrait")
    
    root_path = "PTE_Essay_Bilingual_Argument_Ideas_Table.docx"
    build_docx(root_path, orientation="portrait")
    
    # 2. Also save an explicit portrait-named file in root
    portrait_path = "IELTS_PTE_Essay_Ideas_Table_Format.docx"
    build_docx(portrait_path, orientation="portrait")
    
    # 3. Also save a landscape copy for users who prefer wider horizontal view
    landscape_path = os.path.join(out_dir, "PTE_Essay_Bilingual_Argument_Ideas_Table_Landscape.docx")
    build_docx(landscape_path, orientation="landscape")
