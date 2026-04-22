import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook


ESSAY_JSON_WITH_VOCAB_DEFAULT = "public/database/Write Essay/essay-questions-with-vocab.json"
ESSAY_JSON_DEFAULT = "public/database/Write Essay/essay-questions.json"
ESSAY_XLSX_DEFAULT = "public/database/Write Essay/ESSAY/Essay.xlsx"
REPORT_OUT_DEFAULT = "tmp/write-essay-samples/prompt-type-classification-report.json"


PROMPT_TYPES = {
    "agree_disagree",
    "discuss_both_views",
    "problems_solutions",
    "advantages_disadvantages",
    "choose_between",
    "responsibility",
    "other",
}


_SPACE_RE = re.compile(r"\s+")


def normalize_prompt(text: str) -> str:
    # Keep this deterministic and simple.
    t = str(text or "").strip().lower()
    t = t.replace("\u2019", "'").replace("\u2018", "'")
    t = t.replace("\u201c", '"').replace("\u201d", '"')
    t = _SPACE_RE.sub(" ", t)
    return t.strip()


def classify_prompt(prompt: str, *, qid: int) -> str:
    p = normalize_prompt(prompt)

    # Manual overrides for known edge cases / product decisions.
    overrides = {
        4: "discuss_both_views",   # ads: contains both sides, best handled as discuss-both
        53: "discuss_both_views",  # explicitly says "Discuss both opinions"
    }
    if qid in overrides:
        return overrides[qid]

    # Responsibility ("who should...") is a subtype of choose-between.
    if ("responsibility" in p or "responsible" in p) and ("who" in p or "whose" in p):
        return "responsibility"
    if "main responsibility" in p or "the responsibility of" in p or "responsibility of" in p:
        return "responsibility"

    # Choose-between / preference prompts (explicit signals).
    # This must run BEFORE agree/disagree because some prompts include "agree" but still ask you to pick a side.
    if "which opinion do you agree with" in p:
        return "choose_between"
    if "which one" in p:
        return "choose_between"
    if "which would you choose" in p or "which will you choose" in p:
        return "choose_between"
    if "either" in p and " or " in p and any(k in p for k in ("choose", "which")):
        return "choose_between"
    if "is it better" in p and " or " in p:
        return "choose_between"
    if "which" in p and any(k in p for k in ("prefer", "most reliable", "deserves", "allocate funds")):
        return "choose_between"
    if "should" in p and " or " in p and "?" in prompt and not any(k in p for k in ("agree or disagree", "agree or not", "support this or not")):
        return "choose_between"
    if " or " in p and any(k in p for k in ("education or health", "health or education")):
        return "choose_between"

    # Explicit discuss-both signals.
    if any(k in p for k in ("discuss both", "discuss the two views", "discuss both sides", "discuss both opinions")):
        return "discuss_both_views"
    if ("some people" in p and "while others" in p) or ("some people" in p and "others" in p and "discuss" in p):
        return "discuss_both_views"
    if p.count("some people") >= 2 and any(k in p for k in ("but", "however")) and "discuss" in p:
        return "discuss_both_views"
    if ("some people" in p and "others" in p) and any(k in p for k in ("opinion", "what's your opinion", "what is your opinion")):
        return "discuss_both_views"

    # Problems & solutions signals.
    if any(k in p for k in ("what solutions", "give some solutions", "suggest", "suggestions", "solutions", "solution")):
        return "problems_solutions"
    if any(k in p for k in ("eradicate", "prevent", "mitigate", "address the problem", "solve the problem", "solve the problems")):
        return "problems_solutions"
    if ("how can we" in p or "what should we do" in p or "what can be done" in p) and any(k in p for k in ("problem", "problems", "issue", "issues", "challenge", "challenges")):
        return "problems_solutions"
    if "causes" in p and "solutions" in p:
        return "problems_solutions"
    if "list some causes" in p or "what are the causes" in p or "what is the cause" in p:
        return "problems_solutions"
    if "what problems" in p or "potential challenges" in p or "challenges do" in p:
        return "problems_solutions"
    if any(k in p for k in ("main reasons", "primary reasons", "what are the reasons", "what are some other reasons")):
        return "problems_solutions"

    # Advantages & disadvantages signals (incl. good/bad, positive/negative).
    if "outweigh" in p and any(k in p for k in ("disadvantage", "disadvantages", "drawback", "drawbacks")):
        return "advantages_disadvantages"
    if "advantages" in p and "disadvantages" in p:
        return "advantages_disadvantages"
    if "pros and cons" in p:
        return "advantages_disadvantages"
    if "positive" in p and "negative" in p:
        return "advantages_disadvantages"
    if ("benefit" in p or "benefits" in p) and any(k in p for k in ("drawback", "drawbacks", "problem", "problems")):
        return "advantages_disadvantages"
    if "good or bad" in p or "blessing or a curse" in p or "blessing or curse" in p or "helping or hurting" in p:
        return "advantages_disadvantages"
    if "is it good or bad" in p or "is it a positive or negative" in p:
        return "advantages_disadvantages"
    if "good or not" in p or "is it good or not" in p:
        return "advantages_disadvantages"
    if "beneficial" in p and any(k in p for k in ("detrimental", "harmful")):
        return "advantages_disadvantages"
    if "beneficial or harmful" in p or "beneficial or detrimental" in p:
        return "advantages_disadvantages"

    # Agree/disagree and general opinion prompts (fallback for most prompts).
    if any(k in p for k in ("agree", "disagree", "to what extent", "do you support", "support your opinion")):
        return "agree_disagree"
    if any(k in p for k in ("what is your opinion", "what do you think", "do you think", "express your opinion", "give your opinion")):
        return "agree_disagree"
    if any(k in p for k in ("do you believe", "what is your view", "what are your views", "your view", "point of view", "are you in favor", "in favor of", "are you against")):
        return "agree_disagree"
    if any(k in p for k in ("do you enjoy", "will you support", "what's your opinion", "for or against")):
        return "agree_disagree"
    if "as important" in p and "?" in prompt:
        return "agree_disagree"
    if p.startswith("should ") and "?" in prompt:
        return "agree_disagree"
    if "should" in p and "?" in prompt:
        return "agree_disagree"
    if "is it" in p and any(k in p for k in ("true", "right", "wrong")):
        return "agree_disagree"

    return "other"


def load_json(path: Path) -> List[Dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def update_questions_json(path: Path) -> Tuple[Counter, List[int]]:
    items = load_json(path)
    counts: Counter = Counter()
    unknown: List[int] = []

    for q in items:
        if not isinstance(q, dict):
            continue
        qid = q.get("id")
        if not isinstance(qid, int):
            continue
        prompt = str(q.get("prompt") or "")
        pt = classify_prompt(prompt, qid=qid)
        if pt not in PROMPT_TYPES:
            pt = "other"
        q["promptType"] = pt
        counts[pt] += 1
        if pt == "other":
            unknown.append(qid)

    write_json(path, items)
    return counts, unknown


def update_questions_xlsx(path: Path, prompt_types_by_id: Dict[int, str]) -> None:
    wb = load_workbook(path)
    ws = wb["Questions"]

    # Find / create PROMPT_TYPE column.
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_map = {str(v).strip(): i for i, v in enumerate(headers, start=1) if v is not None}

    col = header_map.get("PROMPT_TYPE")
    if col is None:
        col = len(headers) + 1
        ws.cell(row=1, column=col).value = "PROMPT_TYPE"

    id_col = header_map.get("ID")
    if id_col is None:
        raise ValueError("Missing ID column in Questions sheet")

    for r in range(2, ws.max_row + 1):
        qid = ws.cell(row=r, column=id_col).value
        if not isinstance(qid, int):
            continue
        ws.cell(row=r, column=col).value = prompt_types_by_id.get(qid, "other")

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--json-with-vocab", default=ESSAY_JSON_WITH_VOCAB_DEFAULT)
    ap.add_argument("--json", default=ESSAY_JSON_DEFAULT)
    ap.add_argument("--xlsx", default=ESSAY_XLSX_DEFAULT)
    ap.add_argument("--report-out", default=REPORT_OUT_DEFAULT)
    args = ap.parse_args()

    json_with_vocab_path = Path(args.json_with_vocab)
    json_path = Path(args.json)
    xlsx_path = Path(args.xlsx)
    report_out = Path(args.report_out)

    counts_vocab, unknown_vocab = update_questions_json(json_with_vocab_path)
    counts_plain, unknown_plain = update_questions_json(json_path)

    # Build id -> promptType mapping from the with-vocab JSON (source of truth).
    items = load_json(json_with_vocab_path)
    pt_by_id: Dict[int, str] = {}
    for q in items:
        if isinstance(q, dict) and isinstance(q.get("id"), int):
            pt_by_id[q["id"]] = str(q.get("promptType") or "other")

    update_questions_xlsx(xlsx_path, pt_by_id)

    report = {
        "jsonWithVocab": str(json_with_vocab_path),
        "json": str(json_path),
        "xlsx": str(xlsx_path),
        "counts": dict(counts_vocab),
        "unknownCount": len(unknown_vocab),
        "unknownIds": sorted(set(unknown_vocab + unknown_plain)),
    }
    write_json(report_out, report)

    print("[write-essay] Prompt type classification complete.")
    print("[write-essay] Counts:", dict(counts_vocab))
    if report["unknownCount"]:
        print("[write-essay] Unknown prompts:", report["unknownCount"], "(see report for IDs)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
