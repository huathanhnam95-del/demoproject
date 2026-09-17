#!/usr/bin/env python3
"""
Scrub AI Persona Boilerplate Across Database Explanations
Strips conversational greetings and generic marketing headers while preserving
100% of substantive pedagogical rationale, cohesive devices, and options analysis.
"""

import openpyxl
import os
import shutil
import datetime
import re
import sys
import argparse

sys.stdout.reconfigure(encoding='utf-8')

TARGET_FILES = [
    ('ROP', 'public/database/ROP/ROP/ROP.xlsx'),
    ('HIW', 'public/database/Highlight Incorrect Words/HIW/HIW.xlsx'),
    ('LMCMA', 'public/database/LMCMA/LMCMA/LMCMA.xlsx'),
    ('SMW', 'public/database/SMW/SMW/SMW.xlsx'),
    ('LMCSA', 'public/database/LMCSA/LMCSA/LMCSA.xlsx'),
    ('RMCMA', 'public/database/RMCMA/RMCMA/RMCMA.xlsx'),
    ('RMCSA', 'public/database/RMCSA/RMCSA/RMCSA.xlsx')
]

PERSONA_PATTERNS = [
    # Plaintext or emoji header line at start (e.g. 🔬 PTE Listening Mastery: Incorrect Words Analysis)
    r'^\s*(?:#+\s*)?(?!(?:Correct\s+Answer|Analysis\s+of\s+the\s+Correct|Option\s+Analysis|Why\s+(?:this|it|they)\s+is\s+correct|Why\s+they\s+are\s+right|Distractor\s+Analysis|Key\s+Takeaway|Correct\s+Option))[^\n<]*?(?:PTE\s+(?:Listening|Academic|Expert|Mastery|Analysis|Strategy|Deep\s+Dive|Skill|Teacher|Teaching|Tutor)|Listening\s+Skill|Highlight\s+Incorrect\s+Words|Premium\s+PTE)[^\n<]*?(?:\r?\n+|\s*(?=<))\s*',

    # Generic marketing/category headers at start (strictly bounded: CANNOT cross block elements, CANNOT match pedagogical analysis headers)
    r'^\s*(?:<p>\s*)?(?:<strong[^>]*>|<b[^>]*>|<em[^>]*>|<i[^>]*>)?\s*<h[1-6][^>]*>(?:(?!</?h[1-6]>|<p|</p|<ul|<ol|<div|<table|Correct\s+Answer|Analysis\s+of\s+the\s+Correct|Option\s+Analysis|Why\s+(?:this|it|they)\s+is\s+correct|Why\s+they\s+are\s+right|Distractor\s+Analysis|Key\s+Takeaway|Correct\s+Option).)*?(?:PTE\s+(?:Expert|Listening|Academic|Teacher|Teaching|Analysis|Strategy|Mastery|Skill|Deep\s+Dive|Detailed|Premium)|Listening\s+Skill|Premium\s+PTE|Highlight\s+Incorrect\s+Words|Understanding\s+Cohes|Understanding\s+the\s+Listening|Select\s+Missing\s+Word|Multiple\s+Choice\s+Single\s+Answer|Detailed\s+PTE\s+Listening)(?:(?!</?h[1-6]>|<p|</p|<ul|<ol|<div|<table).)*?</h[1-6]>\s*(?:</strong>|</b>|</em>|</i>)?(?:\s*</p>)?\s*',

    # Paragraph-level marketing headers (e.g. <p><b>🧠 PTE Teacher Analysis...</b></p>)
    r'^\s*<p>\s*(?:<strong[^>]*>|<b[^>]*>|<em[^>]*>|<i[^>]*>)?\s*(?:[^\w<]*\s*)?PTE\s+(?:Teacher|Teaching|Premium|Mastery|Listening|Expert|Strategy)\s+(?:Analysis|Tip|Strategy|Deep\s+Dive|Guide)[^<]*?(?:</strong>|</b>|</em>|</i>)?\s*</p>\s*',

    # Generic strategy paragraph intro
    r'^\s*<p><strong>(?:(?!</p>).)*?PTE\s+Listening\s+Strategy:(?:(?!</p>).)*?</strong></p>\s*',

    # Generic conversational greetings and introductions bounded within a single paragraph
    r'^\s*<p>(?:(?!</p>).)*?(?:Hello|Hello\s+there|Welcome|Hi|Hey|Greetings)(?:[!,\.]|\s+there[!,\.]?|\s+everyone[!,\.]?)(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>(?:(?!</p>).)*?As\s+(?:an?\s+)?(?:(?:expert|premium|your)\s*){1,3}PTE\s+(?:Academic\s+)?(?:tutor|teacher|instructor|expert)(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>(?:(?!</p>).)*?As\s+(?:an?\s+)?PTE\s+(?:Academic\s+)?(?:tutor|teacher|instructor|expert)(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>(?:(?!</p>).)*?(?:Mastering|mastering)\s+(?:the\s+)?(?:[\'\"“]?Re-order\s+Paragraphs[\'\"”]?|PTE\s+Re-order\s+Paragraphs)(?:(?!</p>).)*?</p>\s*',

    # Generic intro sentences in HIW/SMW bounded within a single paragraph
    r'^\s*<p>This\s+question\s+tests\s+your\s+ability\s+to\s+discriminate\s+between\s+phonetically\s+similar\s+words(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>In\s+this\s+question\s+type,\s+it\s+is\s+crucial\s+to\s+listen\s+not\s+only\s+for\s+keywords(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>This\s+is\s+a\s+common\s+and\s+often\s+challenging\s+question\s+type\s+in\s+the\s+PTE\s+Listening\s+section\.(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>This\s+is\s+a\s+high-level\s+vocabulary\s+and\s+contextual\s+inference\s+question\.(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>In\s+this\s+type\s+of\s+question,\s+you\s+must\s+imagine\s+the\s+missing\s+word(?:(?!</p>).)*?</p>\s*',
    r'^\s*<p>In\s+this\s+question\s+type,\s+you\s+must\s+use\s+your\s+understanding\s+of\s+context(?:(?!</p>).)*?</p>\s*',
]

COMPILED_PATTERNS = [re.compile(p, re.IGNORECASE | re.DOTALL) for p in PERSONA_PATTERNS]


def clean_explanation(text):
    if not text or not isinstance(text, str):
        return text
    
    cleaned = text
    changed = True
    iterations = 0
    while changed and iterations < 8:
        changed = False
        iterations += 1
        for cp in COMPILED_PATTERNS:
            m = cp.search(cleaned)
            if m and m.start() == 0:
                cleaned = cleaned[m.end():].lstrip()
                changed = True
                break
    return cleaned


def main():
    parser = argparse.ArgumentParser(description="Scrub AI Persona Boilerplate Across Database Explanations")
    parser.add_argument("--dry-run", action="store_true", help="Perform inspection without writing files")
    args = parser.parse_args()

    backup_dir = os.path.join('.local', 'backups')
    os.makedirs(backup_dir, exist_ok=True)
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')

    print("==================================================================")
    print(f"Scrubbing AI Persona Boilerplate (Dry run: {args.dry_run})")
    print("==================================================================")

    total_cleaned = 0
    for mode, fpath in TARGET_FILES:
        if not os.path.exists(fpath):
            print(f"Warning: File {fpath} does not exist. Skipping.")
            continue

        wb = openpyxl.load_workbook(fpath)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        if 'EXPLANATION' not in headers:
            print(f"Warning: 'EXPLANATION' column not in {fpath}. Skipping.")
            continue
        exp_col = headers.index('EXPLANATION') + 1

        file_cleaned = 0
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=exp_col)
            orig = cell.value
            if not orig or not isinstance(orig, str):
                continue
            cleaned = clean_explanation(orig)
            if cleaned != orig:
                # Safety assertion: cleaned explanation must not be wiped out
                assert len(cleaned) >= 40, f"Error: Explanation in {fpath} row {r} cleaned too aggressively (<40 chars): {repr(cleaned)}"
                stripped = orig[:len(orig) - len(cleaned)]
                assert not any(w in stripped for w in ['Correct Answer', 'Analysis of the Correct', 'Why this is correct', 'Why it is correct', 'Why they are right', 'Option Analysis', 'Key Takeaway', 'Distractor Analysis', 'Correct Option']), f"Critical Safety Error: Pedagogy stripped in {fpath} row {r}: {repr(stripped[:100])}"
                cell.value = cleaned
                file_cleaned += 1

        total_cleaned += file_cleaned
        print(f"[{mode}] {file_cleaned} rows scrubbed (out of {ws.max_row - 1} total)")

        if not args.dry_run and file_cleaned > 0:
            bak_path = os.path.join(backup_dir, f"{mode}_{os.path.basename(fpath)}.{timestamp}.bak")
            shutil.copy2(fpath, bak_path)
            print(f"  -> Backed up to {bak_path}")
            wb.save(fpath)
            print(f"  -> Saved changes to {fpath}")

    print("==================================================================")
    print(f"Total rows cleaned: {total_cleaned}")
    print("==================================================================")


if __name__ == '__main__':
    main()
