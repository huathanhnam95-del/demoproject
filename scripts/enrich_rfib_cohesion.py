"""
enrich_rfib_cohesion.py — Audit RFIB Cohesion Feature Details and generate
detailed student-facing explanations using local Gemma AI (Ollama).

Uses openpyxl for cell-level writes to preserve workbook formatting.
Follows backup/resume/locked-file patterns from run-rfib-cohesion-audit.js.

Usage:
  python scripts/enrich_rfib_cohesion.py --test 10
  python scripts/enrich_rfib_cohesion.py --ids 13,14,9,15
  python scripts/enrich_rfib_cohesion.py                        # full run, resume
  python scripts/enrich_rfib_cohesion.py --no-resume            # full rerun
  python scripts/enrich_rfib_cohesion.py --validate-only        # validate sidecar
  python scripts/enrich_rfib_cohesion.py --validate-only --print-sample
"""

import os
import re
import sys
import time
import json
import shutil
import logging
import argparse
import requests
import openpyxl
from datetime import datetime, timezone

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "gemma4:latest"

INPUT_FILE = r"C:\Cursor AI\public\database\RFIB\RFIB Final ver.xlsx"
DEFAULT_OUTPUT = INPUT_FILE
SIDECAR_FILE = r"C:\Cursor AI\public\database\RFIB\RFIB_cohesion_enrichment.jsonl"

VALID_STATUSES = {"Correct", "Partially Correct", "Incorrect", "No Cohesion"}
VALID_STATUS_PREFIXES = {f"[{s}]" for s in VALID_STATUSES}

COL_DETAILED = "Detailed Cohesion Explanation"
COL_VERIFICATION = "Cohesion Verification"

# ---------------------------------------------------------------------------
# Blank extraction (mirrors run-rfib-cohesion-audit.js extractBlanks)
# ---------------------------------------------------------------------------

BLANK_RE = re.compile(r"__([^_]+?)__")


def extract_blanks(answer_text: str) -> list[dict]:
    """Return list of {index, correct, options} from __opt1/opt2/opt3/opt4__."""
    if not answer_text:
        return []
    matches = BLANK_RE.findall(answer_text)
    blanks = []
    for i, payload in enumerate(matches, start=1):
        options = [s.strip() for s in payload.split("/")]
        blanks.append({"index": i, "correct": options[0], "options": options})
    return blanks


# ---------------------------------------------------------------------------
# Column H detection
# ---------------------------------------------------------------------------


def is_cohesion_false(value) -> bool:
    """Detect FALSE using identity check for bool, with string fallback."""
    if value is False:
        return True
    if isinstance(value, str) and value.strip().lower() == "false":
        return True
    return False


def cohesion_label(value) -> str:
    """Return 'TRUE' or 'FALSE' string from Column H raw value."""
    if is_cohesion_false(value):
        return "FALSE"
    return "TRUE"


# ---------------------------------------------------------------------------
# Prompt
# ---------------------------------------------------------------------------


def build_prompt(
    answer_with_blanks: str,
    full_text: str,
    existing_details: str,
    blanks: list[dict],
    col_h_label: str,
) -> str:
    blanks_block = "\n".join(
        f"  Blank {b['index']}: correct=\"{b['correct']}\"  options={b['options']}"
        for b in blanks
    )

    return f"""You are an expert ESL teacher and linguistic analyst specialising in textual cohesion.

TASK: Analyse a Reading Fill-in-the-Blanks (RFIB) passage. For each blank, explain
why the correct answer is the only logical choice. Separate two concerns:
  (a) selection_reason — grammar, collocation, or semantic fit that makes this word correct.
  (b) cohesion_tie (optional) — a genuine lexical or grammatical cohesion link
      to other words/clauses in the passage (reference pronoun, synonym/antonym,
      repetition, conjunction). Only claim a cohesion tie if one truly exists.
      Not every blank has a cohesion tie — many are selected purely by grammar or collocation.

We also have an existing cohesion analysis from our database (Column I).
The current Cohesion Feature label (Column H) is: {col_h_label}
The existing analysis may be correct, partially correct, incorrect, or irrelevant.

Audit it carefully:
 - Does it reference the correct blank position? (e.g. if a blank is "all" in
   "all around Tokyo", but the existing analysis says "'all' functions as a
   reference pronoun in 'all of which'" in a DIFFERENT sentence, that is INCORRECT
   because 'all of which' is not a blank — 'all around' is the blank.)
 - Are the claimed cohesion features accurate and actually helpful to a student?
 - Is it complete for cohesion-bearing blanks? (Blanks with no cohesion tie don't
   count as omissions.)

VERIFICATION RUBRIC:
 - "Correct" — every claim is factually accurate and references the right blanks.
 - "Partially Correct" — all claims are accurate but incomplete on blanks that DO
   have genuine cohesion ties (not covering non-cohesion blanks is fine).
 - "Incorrect" — one or more claims contain a materially false statement.
 - "No Cohesion" — Column H is FALSE and details correctly state no cohesion applies.

DATA:
1. Passage with blanks (__option1/option2/option3/option4__):
{answer_with_blanks}

2. Full solved text:
{full_text}

3. Detected blanks:
{blanks_block}

4. Column H label: {col_h_label}

5. Existing cohesion analysis (Column I):
{existing_details}

RESPONSE FORMAT — return ONLY valid JSON, no markdown fences:
{{
  "cohesion_verification_status": "Correct" | "Partially Correct" | "Incorrect" | "No Cohesion",
  "cohesion_verification_notes": "Brief explanation of why the original analysis is correct/partially correct/incorrect/not applicable.",
  "detailed_cohesion_explanations": [
    {{
      "blank_index": 1,
      "correct_answer": "word",
      "selection_reason": "Why this word is the only correct choice (grammar, collocation, meaning).",
      "cohesion_tie": "Lexical/grammatical cohesion link to other text, or empty string if none.",
      "detailed_student_explanation": "Full student-facing explanation combining selection_reason and cohesion_tie. Write as if explaining to a B1-B2 learner why this answer is correct and how they can figure it out from the surrounding text."
    }}
  ]
}}

RULES:
- Return exactly {len(blanks)} objects in detailed_cohesion_explanations, one per blank, in order.
- blank_index must be sequential starting from 1.
- correct_answer must match the first option in each blank exactly.
- detailed_student_explanation must be non-empty.
- cohesion_verification_status must be exactly one of: Correct, Partially Correct, Incorrect, No Cohesion.
"""


# ---------------------------------------------------------------------------
# Ollama query with retry
# ---------------------------------------------------------------------------


def query_gemma(prompt: str, max_retries: int = 3) -> str | None:
    payload = {
        "model": MODEL_NAME,
        "prompt": prompt,
        "format": "json",
        "stream": False,
        "options": {"temperature": 0.1},
    }
    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(OLLAMA_URL, json=payload, timeout=180)
            if resp.status_code == 200:
                return resp.json().get("response", "")
            logging.warning(f"Ollama HTTP {resp.status_code} (attempt {attempt})")
        except requests.exceptions.Timeout:
            logging.warning(f"Ollama timeout (attempt {attempt})")
        except Exception as e:
            logging.warning(f"Ollama error (attempt {attempt}): {e}")
        if attempt < max_retries:
            time.sleep(2 * attempt)
    return None


# ---------------------------------------------------------------------------
# Response validation
# ---------------------------------------------------------------------------


def validate_response(data: dict, blanks: list[dict]) -> list[str]:
    """Return a list of validation errors (empty = valid)."""
    errors = []

    status = data.get("cohesion_verification_status")
    if status not in VALID_STATUSES:
        errors.append(f"Invalid status '{status}'; expected one of {VALID_STATUSES}")

    explanations = data.get("detailed_cohesion_explanations")
    if not isinstance(explanations, list):
        errors.append("detailed_cohesion_explanations is not a list")
        return errors

    if len(explanations) != len(blanks):
        errors.append(
            f"Explanation count mismatch: got {len(explanations)}, expected {len(blanks)}"
        )

    seen_indexes = set()
    for i, exp in enumerate(explanations):
        idx = exp.get("blank_index")
        expected_idx = i + 1
        if idx != expected_idx:
            errors.append(f"Explanation {i}: blank_index={idx}, expected {expected_idx}")
        if idx in seen_indexes:
            errors.append(f"Duplicate blank_index={idx}")
        seen_indexes.add(idx)

        if i < len(blanks):
            expected_answer = blanks[i]["correct"]
            actual_answer = exp.get("correct_answer", "")
            if actual_answer.strip().lower() != expected_answer.strip().lower():
                errors.append(
                    f"Blank {idx}: correct_answer='{actual_answer}' != expected='{expected_answer}'"
                )

        if not exp.get("detailed_student_explanation", "").strip():
            errors.append(f"Blank {idx}: detailed_student_explanation is empty")

    return errors


# ---------------------------------------------------------------------------
# Formatting for Excel cell
# ---------------------------------------------------------------------------


def format_detailed_explanations(explanations: list[dict]) -> str:
    parts = []
    for item in explanations:
        idx = item.get("blank_index", "?")
        ans = item.get("correct_answer", "?")
        reason = item.get("selection_reason", "")
        cohesion = item.get("cohesion_tie", "")
        student_exp = item.get("detailed_student_explanation", "")

        lines = [f"Blank {idx} ('{ans}'):"]
        lines.append(f"  Why this answer: {reason}")
        if cohesion:
            lines.append(f"  Cohesion link: {cohesion}")
        lines.append(f"  Student explanation: {student_exp}")
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


# ---------------------------------------------------------------------------
# Resume check
# ---------------------------------------------------------------------------


def is_row_already_processed(ws, row_num: int, col_detailed: int, col_verification: int) -> bool:
    """A row is processed only if BOTH cells are validly filled."""
    explanation = ws.cell(row=row_num, column=col_detailed).value
    if not explanation or not str(explanation).strip():
        return False

    verification = ws.cell(row=row_num, column=col_verification).value
    if not verification:
        return False
    v_str = str(verification).strip()
    return any(v_str.startswith(prefix) for prefix in VALID_STATUS_PREFIXES)


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------


def build_header_map(ws) -> dict[str, int]:
    hmap = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None:
            hmap[str(val).strip()] = col
    return hmap


def ensure_column(ws, hmap: dict, name: str) -> int:
    if name in hmap:
        return hmap[name]
    next_col = ws.max_column + 1
    ws.cell(row=1, column=next_col, value=name)
    hmap[name] = next_col
    return next_col


def backup_suffix() -> str:
    t = time.localtime()
    return f"{t.tm_year}{t.tm_mon:02d}{t.tm_mday:02d}_{t.tm_hour:02d}{t.tm_min:02d}{t.tm_sec:02d}"


def save_workbook_safe(wb, target_path: str) -> str:
    """Save via temp file + atomic rename. Returns actual path used."""
    tmp_path = target_path + ".tmp"
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, target_path)
        return target_path
    except PermissionError:
        fallback = target_path.replace(".xlsx", f".cohesion_enriched.{backup_suffix()}.xlsx")
        logging.warning(f"File locked. Saving to fallback: {fallback}")
        try:
            os.remove(tmp_path)
        except OSError:
            pass
        wb.save(fallback)
        return fallback
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        raise


# ---------------------------------------------------------------------------
# JSONL sidecar
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# JSONL sidecar
# ---------------------------------------------------------------------------


def load_sidecar_to_dict(sidecar_path: str) -> dict[int, dict]:
    """Load sidecar JSONL into a dict of id -> record, raising ValueError on duplicate IDs."""
    records = {}
    if not os.path.exists(sidecar_path):
        return records
        
    with open(sidecar_path, "r", encoding="utf-8") as f:
        for line_num, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as e:
                logging.warning(f"Malfored JSON in sidecar line {line_num}: {e}")
                continue
                
            rid = record.get("id")
            if rid is None:
                continue
            rid = int(rid)
            if rid in records:
                raise ValueError(f"Duplicate ID {rid} detected in sidecar at line {line_num}")
            records[rid] = record
            
    return records


def save_sidecar_atomic(sidecar_path: str, records: dict[int, dict]):
    """Save in-memory sidecar dictionary atomically to file."""
    tmp_path = sidecar_path + ".tmp"
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            for rid in sorted(records.keys()):
                f.write(json.dumps(records[rid], ensure_ascii=False) + "\n")
        os.replace(tmp_path, sidecar_path)
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        raise


def checkpoint(wb, output_path: str, sidecar_path: str, records: dict[int, dict]) -> str:
    """Perform sidecar-first then workbook-second checkpointing. Returns active workbook path."""
    save_sidecar_atomic(sidecar_path, records)
    active_path = save_workbook_safe(wb, output_path)
    return active_path


def reconcile_row(ws, row_num: int, row_id: int, col_detailed: int, col_verification: int, sidecar_records: dict[int, dict]) -> str:
    """Reconcile row state. Returns 'skip', 'render', or 'reprocess'."""
    # Check Excel validity
    excel_detailed = ws.cell(row=row_num, column=col_detailed).value
    excel_ver = ws.cell(row=row_num, column=col_verification).value
    excel_valid = False
    excel_status = None
    if excel_detailed and str(excel_detailed).strip() and excel_ver:
        ver_str = str(excel_ver).strip()
        for prefix in VALID_STATUS_PREFIXES:
            if ver_str.startswith(prefix):
                excel_valid = True
                excel_status = prefix[1:-1]
                break

    # Check sidecar validity
    sidecar_valid = False
    rec = sidecar_records.get(row_id)
    if rec:
        status = rec.get("status")
        explanations = rec.get("explanations")
        if status in VALID_STATUSES and isinstance(explanations, list):
            sidecar_valid = True

    # Reconciliation logic
    if excel_valid and sidecar_valid:
        # Check alignment
        if excel_status == rec.get("status"):
            logging.info(f"Row ID {row_id}: Excel and sidecar align. Skipping.")
            return "skip"
        else:
            logging.info(f"Row ID {row_id}: Status mismatch (Excel={excel_status}, sidecar={rec.get('status')}). Reprocessing.")
            return "reprocess"
    elif sidecar_valid and not excel_valid:
        logging.info(f"Row ID {row_id}: Valid sidecar but invalid Excel cells. Recovering cells from sidecar.")
        return "render"
    elif excel_valid and not sidecar_valid:
        logging.info(f"Row ID {row_id}: Excel cells are filled but sidecar is missing/invalid. Reprocessing.")
        return "reprocess"
    else:
        return "reprocess"


def reset_sidecar_if_requested(sidecar_path: str, reset: bool):
    """Log sidecar count/path and delete if reset is True."""
    count = 0
    exists = os.path.exists(sidecar_path)
    if exists:
        try:
            records = load_sidecar_to_dict(sidecar_path)
            count = len(records)
        except Exception as e:
            logging.warning(f"Could not parse sidecar to count records before reset: {e}")
            
    logging.info(f"Sidecar path: {sidecar_path}")
    logging.info(f"Existing records count: {count} (exists={exists})")
    
    if reset and exists:
        logging.info(f"Reset requested. Deleting sidecar file: {sidecar_path}")
        os.remove(sidecar_path)


# ---------------------------------------------------------------------------
# Validate-only mode
# ---------------------------------------------------------------------------


def run_validate_only(args):
    """Read sidecar JSONL and cross-check against workbook."""
    sidecar_path = args.sidecar or SIDECAR_FILE
    logging.info(f"Validation mode: reading sidecar {sidecar_path}")

    records = load_sidecar(sidecar_path)
    if not records:
        logging.error("No records found in sidecar file.")
        sys.exit(1)

    logging.info(f"Loading workbook: {INPUT_FILE}")
    wb = openpyxl.load_workbook(INPUT_FILE, read_only=True)
    ws = wb.active
    hmap = build_header_map(ws)

    col_answer = hmap.get("ANSWER")
    col_detailed = hmap.get(COL_DETAILED)
    col_verification = hmap.get(COL_VERIFICATION)

    if not col_answer:
        logging.error("ANSWER column not found in workbook.")
        sys.exit(1)

    # Build ID → row map
    id_to_row = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=hmap["ID"]).value
        if rid is not None:
            id_to_row[int(rid)] = r

    passed = 0
    failed = 0
    failed_ids = []

    for rec in records:
        rec_id = rec.get("id")
        errors = []

        if rec_id not in id_to_row:
            errors.append(f"ID {rec_id} not found in workbook")
        else:
            row_num = id_to_row[rec_id]
            answer_text = str(ws.cell(row=row_num, column=col_answer).value or "")
            blanks = extract_blanks(answer_text)

            # Check status
            status = rec.get("status")
            if status not in VALID_STATUSES:
                errors.append(f"Invalid status '{status}'")

            # Check explanations
            explanations = rec.get("explanations", [])
            if len(explanations) != len(blanks):
                errors.append(f"Explanation count {len(explanations)} != blank count {len(blanks)}")

            seen_idx = set()
            for i, exp in enumerate(explanations):
                idx = exp.get("blank_index")
                if idx != i + 1:
                    errors.append(f"blank_index={idx}, expected {i+1}")
                if idx in seen_idx:
                    errors.append(f"Duplicate blank_index={idx}")
                seen_idx.add(idx)

                if i < len(blanks):
                    expected = blanks[i]["correct"]
                    actual = exp.get("correct_answer", "")
                    if actual.strip().lower() != expected.strip().lower():
                        errors.append(f"Blank {idx}: answer '{actual}' != '{expected}'")

                if not exp.get("detailed_student_explanation", "").strip():
                    errors.append(f"Blank {idx}: empty explanation")

            # Check Excel cells exist
            if col_detailed:
                cell_val = ws.cell(row=row_num, column=col_detailed).value
                if not cell_val or not str(cell_val).strip():
                    errors.append("Excel Detailed Cohesion Explanation cell is empty")
            if col_verification:
                cell_val = ws.cell(row=row_num, column=col_verification).value
                if not cell_val or not str(cell_val).strip():
                    errors.append("Excel Cohesion Verification cell is empty")

        if errors:
            failed += 1
            failed_ids.append(rec_id)
            logging.warning(f"FAIL ID {rec_id}: {errors}")
        else:
            passed += 1
            logging.info(f"PASS ID {rec_id}")

    # Print sample if requested
    if args.print_sample:
        _print_stratified_sample(records, ws, hmap, id_to_row)

    logging.info("=" * 60)
    logging.info("VALIDATION REPORT")
    logging.info(f"  Records checked: {len(records)}")
    logging.info(f"  Passed:          {passed}")
    logging.info(f"  Failed:          {failed}")
    if failed_ids:
        logging.info(f"  Failed IDs:      {failed_ids}")
    logging.info("=" * 60)

    wb.close()
    sys.exit(0 if failed == 0 else 1)


def _print_stratified_sample(records, ws, hmap, id_to_row):
    """Print a stratified sample for human review."""
    by_status = {}
    for rec in records:
        s = rec.get("status", "Unknown")
        by_status.setdefault(s, []).append(rec)

    col_h_idx = hmap.get("Cohesion Feature")
    col_answer_idx = hmap.get("ANSWER")

    samples = []
    for s, count in [("Incorrect", 5), ("Partially Correct", 5)]:
        for rec in by_status.get(s, [])[:count]:
            samples.append(rec)

    # FALSE rows
    false_count = 0
    for rec in records:
        rid = rec.get("id")
        if rid in id_to_row and col_h_idx:
            val = ws.cell(row=id_to_row[rid], column=col_h_idx).value
            if is_cohesion_false(val):
                samples.append(rec)
                false_count += 1
                if false_count >= 3:
                    break

    # Multi-blank rows
    multi_count = 0
    for rec in records:
        rid = rec.get("id")
        if rid in id_to_row and col_answer_idx:
            answer = str(ws.cell(row=id_to_row[rid], column=col_answer_idx).value or "")
            if len(extract_blanks(answer)) >= 5:
                if rec not in samples:
                    samples.append(rec)
                    multi_count += 1
                    if multi_count >= 3:
                        break

    logging.info("=" * 60)
    logging.info("STRATIFIED HUMAN REVIEW SAMPLE")
    logging.info("=" * 60)
    for rec in samples:
        logging.info(f"\n--- ID {rec.get('id')} [{rec.get('status')}] ---")
        logging.info(f"Notes: {rec.get('notes', '')}")
        for exp in rec.get("explanations", []):
            logging.info(
                f"  Blank {exp.get('blank_index')} ('{exp.get('correct_answer')}'): "
                f"{exp.get('detailed_student_explanation', '')[:200]}"
            )


# ---------------------------------------------------------------------------
# Main enrichment logic
# ---------------------------------------------------------------------------


def parse_arguments(args_list: list[str] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Enrich RFIB cohesion explanations using local Gemma AI."
    )
    parser.add_argument("--input", type=str, default=INPUT_FILE, help="Input workbook path.")
    parser.add_argument("--out", type=str, default="", help="Output path (default: overwrite input via temp file).")
    parser.add_argument("--sidecar", type=str, default="", help="JSONL sidecar path.")
    parser.add_argument("--test", type=int, default=0, help="Curated sample of N rows.")
    parser.add_argument("--ids", type=str, default="", help="Comma-separated IDs to process.")
    parser.add_argument("--start", type=int, default=0, help="Start from Nth data row (0-based).")
    parser.add_argument("--limit", type=int, default=0, help="Max rows to process.")
    parser.add_argument("--no-resume", action="store_true", help="Reprocess all rows.")
    parser.add_argument("--reset-sidecar", action="store_true", help="Start the selected run with an empty sidecar map.")
    parser.add_argument("--save-every", type=int, default=25, help="Save progress every N rows.")
    parser.add_argument("--validate-only", action="store_true", help="Validate sidecar against workbook; no LLM.")
    parser.add_argument("--print-sample", action="store_true", help="Print stratified sample (with --validate-only).")
    
    # Parse args (sys.argv[1:] by default if args_list is None)
    if args_list is not None:
        args = parser.parse_args(args_list)
    else:
        args = parser.parse_args()

    # Validation: save-every must be > 0
    if args.save_every <= 0:
        logging.error("Error: --save-every must be a positive integer.")
        sys.exit(2)

    # Validation: mutual exclusivity of test, ids, start
    selectors = [bool(args.test), bool(args.ids), bool(args.start != 0)]
    if sum(selectors) > 1:
        logging.error("Error: --test, --ids, and --start are mutually exclusive.")
        sys.exit(2)

    return args


def resolve_candidate_rows(ws, hmap: dict, ids_str: str = "", start: int = 0, limit: int = 0, test: int = 0) -> list[int]:
    candidate_rows = []
    
    if ids_str:
        target_ids = set(int(x.strip()) for x in ids_str.split(",") if x.strip())
        for r in range(2, ws.max_row + 1):
            cell_id = ws.cell(row=r, column=hmap["ID"]).value
            if cell_id is not None and int(cell_id) in target_ids:
                candidate_rows.append(r)
    elif test > 0:
        candidate_rows = _build_curated_sample(ws, hmap, test)
    else:
        start_row = 2 + start
        candidate_rows = list(range(start_row, ws.max_row + 1))

    if limit > 0:
        candidate_rows = candidate_rows[:limit]

    # If an explicit selector was used and resolved to empty, exit 2
    if (ids_str or test > 0 or start != 0 or limit > 0) and not candidate_rows:
        logging.error("Error: Selected candidate rows resolved to empty set.")
        sys.exit(2)

    return candidate_rows


def main():
    args = parse_arguments()

    sidecar_path = args.sidecar or SIDECAR_FILE

    if args.validate_only:
        run_validate_only(args)
        return

    output_path = args.out if args.out else args.input

    # --- Load workbook ---
    logging.info(f"Loading workbook: {args.input}")
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active
    hmap = build_header_map(ws)

    required = ["ID", "TITLE", "ANSWER", "Full Text", "Cohesion Feature", "Cohesion Feature Details"]
    missing = [c for c in required if c not in hmap]
    if missing:
        logging.error(f"Missing required columns: {missing}")
        sys.exit(1)

    col_detailed = ensure_column(ws, hmap, COL_DETAILED)
    col_verification = ensure_column(ws, hmap, COL_VERIFICATION)

    # --- Backup ---
    if output_path == args.input:
        bak_path = args.input.replace(".xlsx", f".backup.{backup_suffix()}.xlsx")
        shutil.copy2(args.input, bak_path)
        logging.info(f"Backup created: {bak_path}")

    # --- Determine rows to process ---
    total_data_rows = ws.max_row - 1
    logging.info(f"Total data rows in workbook: {total_data_rows}")

    candidate_rows = resolve_candidate_rows(ws, hmap, ids_str=args.ids, start=args.start, limit=args.limit, test=args.test)

    logging.info(f"Candidate rows: {len(candidate_rows)}")

    # --- Reset sidecar if requested ---
    reset_sidecar_if_requested(sidecar_path, args.reset_sidecar)

    # --- Load sidecar dict ---
    try:
        sidecar_records = load_sidecar_to_dict(sidecar_path)
    except ValueError as e:
        logging.error(f"Failed to load sidecar: {e}")
        sys.exit(1)

    # --- Process ---
    stats = {
        "attempted": 0,
        "skipped_resume": 0,
        "succeeded": 0,
        "failed_ids": [],
        "status_counts": {s: 0 for s in VALID_STATUSES},
    }

    for row_num in candidate_rows:
        row_id_raw = ws.cell(row=row_num, column=hmap["ID"]).value
        if row_id_raw is None:
            continue
        row_id = int(row_id_raw)

        # Reconciliation check on resume
        action = "reprocess"
        if not args.no_resume:
            action = reconcile_row(ws, row_num, row_id, col_detailed, col_verification, sidecar_records)

        if action == "skip":
            stats["skipped_resume"] += 1
            continue
        elif action == "render":
            rec = sidecar_records[row_id]
            formatted = format_detailed_explanations(rec.get("explanations", []))
            ws.cell(row=row_num, column=col_detailed, value=formatted)
            ws.cell(row=row_num, column=col_verification, value=f"[{rec.get('status')}] {rec.get('notes', '')}")
            stats["succeeded"] += 1
            if rec.get("status") in stats["status_counts"]:
                stats["status_counts"][rec.get("status")] += 1
            logging.info(f"  [+] ID {row_id} rendered from sidecar.")
            continue

        # Otherwise action is "reprocess"
        answer_text = str(ws.cell(row=row_num, column=hmap["ANSWER"]).value or "")
        full_text = str(ws.cell(row=row_num, column=hmap["Full Text"]).value or "")
        existing_details = str(ws.cell(row=row_num, column=hmap["Cohesion Feature Details"]).value or "")
        col_h_raw = ws.cell(row=row_num, column=hmap["Cohesion Feature"]).value
        col_h_label = cohesion_label(col_h_raw)

        blanks = extract_blanks(answer_text)
        if not blanks:
            logging.warning(f"Row {row_num} ID={row_id}: No blanks detected. Skipping.")
            continue

        stats["attempted"] += 1
        logging.info(
            f"Processing ID {row_id} (row {row_num}) — {len(blanks)} blanks, "
            f"ColH={col_h_label} [{stats['attempted']}/{len(candidate_rows)}]"
        )

        prompt = build_prompt(answer_text, full_text, existing_details, blanks, col_h_label)

        # Query with retries (outer retry for validation failures)
        success = False
        for outer_attempt in range(1, 4):
            response_text = query_gemma(prompt)
            if not response_text:
                logging.error(f"  ID {row_id}: No response from Ollama (attempt {outer_attempt})")
                continue

            try:
                data = json.loads(response_text.strip())
            except json.JSONDecodeError as e:
                logging.error(f"  ID {row_id}: JSON parse error (attempt {outer_attempt}): {e}")
                continue

            validation_errors = validate_response(data, blanks)
            if validation_errors:
                logging.warning(
                    f"  ID {row_id}: Validation failed (attempt {outer_attempt}): {validation_errors}"
                )
                if outer_attempt < 3:
                    continue
                break

            # Success — write to cells
            status = data.get("cohesion_verification_status", "Unknown")
            notes = data.get("cohesion_verification_notes", "")
            explanations = data.get("detailed_cohesion_explanations", [])

            formatted = format_detailed_explanations(explanations)
            verification_str = f"[{status}] {notes}"

            ws.cell(row=row_num, column=col_detailed, value=formatted)
            ws.cell(row=row_num, column=col_verification, value=verification_str)

            # Update sidecar dict
            sidecar_records[row_id] = {
                "id": row_id,
                "status": status,
                "notes": notes,
                "explanations": explanations,
                "model": MODEL_NAME,
                "timestamp": datetime.now(timezone.utc).isoformat(),
            }

            stats["succeeded"] += 1
            if status in stats["status_counts"]:
                stats["status_counts"][status] += 1
            logging.info(f"  [+] ID {row_id}: Status={status}")
            success = True
            break

        if not success:
            stats["failed_ids"].append(row_id)

        # Periodic save
        if stats["attempted"] % args.save_every == 0 and stats["attempted"] > 0:
            output_path = checkpoint(wb, output_path, sidecar_path, sidecar_records)
            logging.info(f"  💾 Saved progress -> {output_path}")

        # Pacing
        time.sleep(0.5)

    # --- Final save ---
    output_path = checkpoint(wb, output_path, sidecar_path, sidecar_records)

    # --- Workbook-wide enrichment count ---
    workbook_enriched = 0
    for r in range(2, ws.max_row + 1):
        if is_row_already_processed(ws, r, col_detailed, col_verification):
            workbook_enriched += 1

    # --- Coverage report ---
    logging.info("=" * 60)
    logging.info("COVERAGE REPORT")
    logging.info("=" * 60)
    logging.info(f"  Workbook-wide enrichment: {workbook_enriched} / {total_data_rows} rows")
    logging.info(f"  This invocation:")
    logging.info(f"    Candidate rows:    {len(candidate_rows)}")
    logging.info(f"    Skipped (resume):  {stats['skipped_resume']}")
    logging.info(f"    Attempted:         {stats['attempted']}")
    logging.info(f"    Succeeded:         {stats['succeeded']}")
    logging.info(f"    Failed:            {len(stats['failed_ids'])}")
    if stats["failed_ids"]:
        logging.info(f"    Failed IDs:        {stats['failed_ids']}")
    logging.info(f"  Status breakdown:    {stats['status_counts']}")
    logging.info(f"  Output file:         {output_path}")
    logging.info(f"  JSONL sidecar:       {sidecar_path}")
    
    if stats["failed_ids"]:
        rerun_cmd = (
            f".\\.venv\\Scripts\\python.exe scripts\\enrich_rfib_cohesion.py "
            f"--input {args.input} --out {output_path} --sidecar {sidecar_path} "
            f"--ids {','.join(map(str, stats['failed_ids']))}"
        )
        logging.info(f"  Rerun command:       {rerun_cmd}")
    logging.info("=" * 60)

    # Exit with code 1 if any failed
    if stats["failed_ids"]:
        sys.exit(1)
    else:
        sys.exit(0)


# ---------------------------------------------------------------------------
# Curated sample builder
# ---------------------------------------------------------------------------


def _build_curated_sample(ws, hmap: dict, n: int) -> list[int]:
    """Build a curated sample of row numbers prioritising user-flagged, FALSE, and multi-blank rows."""
    priority_ids = [13, 14]
    false_ids = []
    multi_blank_ids = []

    id_to_row = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=hmap["ID"]).value
        if rid is None:
            continue
        rid = int(rid)
        id_to_row[rid] = r

        col_h_val = ws.cell(row=r, column=hmap["Cohesion Feature"]).value
        if is_cohesion_false(col_h_val):
            false_ids.append(rid)

        answer = str(ws.cell(row=r, column=hmap["ANSWER"]).value or "")
        if len(extract_blanks(answer)) >= 5:
            multi_blank_ids.append(rid)

    curated = list(priority_ids)

    # Add FALSE rows (target 2)
    false_added = 0
    for fid in sorted(false_ids):
        if fid not in curated:
            curated.append(fid)
            false_added += 1
        if false_added >= 2:
            break

    # Add multi-blank rows (target 2)
    multi_added = 0
    for mid in sorted(multi_blank_ids):
        if mid not in curated:
            curated.append(mid)
            multi_added += 1
        if multi_added >= 2:
            break

    # Fill remaining
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=hmap["ID"]).value
        if rid is not None and int(rid) not in curated:
            curated.append(int(rid))
        if len(curated) >= n:
            break

    curated = curated[:n]
    logging.info(f"Curated test IDs: {curated}")
    logging.info(f"  FALSE in sample: {[x for x in curated if x in false_ids]}")
    logging.info(f"  Multi-blank in sample: {[x for x in curated if x in multi_blank_ids]}")

    return [id_to_row[rid] for rid in curated if rid in id_to_row]


if __name__ == "__main__":
    main()
