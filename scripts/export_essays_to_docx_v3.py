import json
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

def export_essays_by_level():
    json_path = 'public/database/Write Essay/essay-questions-with-vocab.json'
    output_dir = r'C:\Cursor AI\public\database\Write Essay\ESSAY\Samples'
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    levels_map = {
        "a2_b1": {"name": "A2-B1", "doc": Document(), "count": 0},
        "b2": {"name": "B2", "doc": Document(), "count": 0},
        "c1": {"name": "C1", "doc": Document(), "count": 0}
    }
    
    # Initialize Headers
    for key, info in levels_map.items():
        doc = info["doc"]
        title = f"PTE Write Essay - Full Sample Bank ({info['name']})"
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(20)
        doc.add_page_break()

    # Sort data by ID to keep it organized
    data.sort(key=lambda x: int(x.get('id', 0)))

    for q in data:
        qid = q.get('id')
        # FIX: The key in JSON is 'prompt', not 'questionText'
        prompt_text = q.get('prompt', '')
        sr = q.get('sampleResponses', {})
        levels = sr.get('levels', {})
        
        for level_key, level_data in levels.items():
            if level_key not in levels_map:
                continue
                
            doc = levels_map[level_key]["doc"]
            variants = level_data.get('variants', [])
            
            if not variants:
                continue
                
            # Question Header
            h = doc.add_heading(f"Question #{qid}", level=1)
            p_prompt = doc.add_paragraph()
            p_prompt.add_run("Prompt: ").bold = True
            p_prompt.add_run(prompt_text)
            
            for i, v in enumerate(variants):
                essay_text = v.get('essay')
                analysis = v.get('analysis', {})
                if not essay_text:
                    continue
                
                doc.add_heading(f"Variant {i+1}: {v.get('label', '')}", level=2)
                
                # Essay Section
                doc.add_heading("Essay", level=3)
                doc.add_paragraph(essay_text)
                
                # Analysis Section
                if analysis:
                    doc.add_heading("Analysis", level=3)
                    if 'point1' in analysis:
                        p_p1 = doc.add_paragraph()
                        p_p1.add_run("Point 1: ").bold = True
                        p_p1.add_run(analysis['point1'])
                    if 'point2' in analysis:
                        p_p2 = doc.add_paragraph()
                        p_p2.add_run("Point 2: ").bold = True
                        p_p2.add_run(analysis['point2'])
                    
                    # Vocabulary Table
                    vocab = analysis.get('vocabulary', [])
                    if vocab:
                        doc.add_heading("Vocabulary", level=4)
                        table = doc.add_table(rows=1, cols=3)
                        table.style = 'Table Grid'
                        hdr_cells = table.rows[0].cells
                        hdr_cells[0].text = 'Term'
                        hdr_cells[1].text = 'English Definition'
                        hdr_cells[2].text = 'Vietnamese Definition'
                        
                        for item in vocab:
                            row_cells = table.add_row().cells
                            row_cells[0].text = str(item.get('term', ''))
                            row_cells[1].text = str(item.get('enGloss', ''))
                            row_cells[2].text = str(item.get('viGloss', ''))
                
                doc.add_paragraph("\n" + "-" * 10 + "\n")
                
            levels_map[level_key]["count"] += 1
            doc.add_paragraph("\n" + "="*50 + "\n")

    # Save the 3 files
    for key, info in levels_map.items():
        filename = f"PTE_Essays_{info['name']}.docx"
        save_path = os.path.join(output_dir, filename)
        print(f"Saving {filename}...")
        info["doc"].save(save_path)
        print(f"Done: {info['count']} questions included.")

if __name__ == "__main__":
    export_essays_by_level()
