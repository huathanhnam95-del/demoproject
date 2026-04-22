import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_BREAK


ESSAY_JSON_PATH_DEFAULT = "public/database/Write Essay/essay-questions-with-vocab.json"
OUT_DOCX_DEFAULTS = {
    "a2_b1": "public/database/Write Essay/ESSAY/Write Essay Samples Pilot - A2-B1.docx",
    "b2": "public/database/Write Essay/ESSAY/Write Essay Samples Pilot - B2.docx",
    "c1": "public/database/Write Essay/ESSAY/Write Essay Samples Pilot - C1.docx",
}
LEVEL_LABELS = {"a2_b1": "A2-B1", "b2": "B2", "c1": "C1"}


PILOT_PROMPT_IDS_DEFAULT = [
    1, 2, 3, 4, 5, 6, 7, 8, 9, 10,
    11, 12, 13, 14, 15, 18, 22, 23, 53, 66,
]


def normalize_newlines(s: str) -> str:
    return s.replace("\r\n", "\n").replace("\r", "\n")


def split_paragraphs(text: str) -> List[str]:
    t = normalize_newlines(text).strip()
    if not t:
        return []
    parts = re.split(r"\n{2,}", t)
    return [p.strip() for p in parts if p.strip()]


def load_questions(path: Path) -> List[Dict[str, Any]]:
    return json.loads(path.read_text(encoding="utf-8"))


def get_approved_variants(q: Dict[str, Any], *, level_id: str) -> List[Dict[str, Any]]:
    sr = q.get("sampleResponses")
    if not isinstance(sr, dict):
        return []

    variants = None
    levels = sr.get("levels")
    if isinstance(levels, dict):
        lvl = levels.get(level_id)
        if isinstance(lvl, dict) and isinstance(lvl.get("variants"), list):
            variants = lvl.get("variants")

    # Legacy fallback: older data stores B2 under sampleResponses.variants
    if variants is None and level_id == "b2":
        legacy = sr.get("variants")
        if isinstance(legacy, list):
            variants = legacy

    if not isinstance(variants, list):
        return []
    out: List[Dict[str, Any]] = []
    for v in variants:
        if not isinstance(v, dict):
            continue
        qa = v.get("qa")
        if isinstance(qa, dict) and qa.get("status") == "approved":
            out.append(v)
    return out


def add_kv_paragraph(doc: Document, key: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{key}: ").bold = True
    p.add_run(value)


def add_pre_block(doc: Document, text: str) -> None:
    t = normalize_newlines(str(text or "")).rstrip()
    if not t:
        return
    p = doc.add_paragraph()
    lines = t.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Courier New"
        if i != len(lines) - 1:
            run.add_break()


def build_doc_for_level(*, level_id: str, questions: List[Dict[str, Any]], ids: List[int]) -> Document:
    by_id = {int(q.get("id")): q for q in questions if isinstance(q, dict) and str(q.get("id")).strip().isdigit()}

    doc = Document()
    label = LEVEL_LABELS.get(level_id, str(level_id))
    doc.add_heading(f"Write Essay Samples (Pilot) - {label}", level=0)

    for idx, prompt_id in enumerate(ids, start=1):
        q = by_id.get(prompt_id)
        if not q:
            continue

        title = str(q.get("title") or "").strip()
        prompt_text = str(q.get("prompt") or "").strip()

        doc.add_heading(title or f"Prompt {prompt_id}", level=1)
        if prompt_text:
            add_kv_paragraph(doc, "Prompt", prompt_text)

        variants = get_approved_variants(q, level_id=level_id)
        if not variants:
            add_kv_paragraph(doc, "Note", "No approved sample was generated for this prompt at this level.")
        else:
            for v in variants:
                label2 = str(v.get("label") or v.get("id") or "Sample").strip()
                doc.add_heading(label2, level=2)

                essay = str(v.get("essay") or "").strip()
                for para in split_paragraphs(essay):
                    doc.add_paragraph(para)

                analysis = v.get("analysis") if isinstance(v.get("analysis"), dict) else {}
                doc.add_heading("Analysis", level=3)
                add_kv_paragraph(doc, "Point one", str(analysis.get("point1") or "").strip())
                add_kv_paragraph(doc, "Point two", str(analysis.get("point2") or "").strip())

                vocab = analysis.get("vocabulary")
                if isinstance(vocab, list) and vocab:
                    doc.add_paragraph("Vocabulary").runs[0].bold = True
                    table = doc.add_table(rows=1, cols=3)
                    hdr = table.rows[0].cells
                    hdr[0].text = "Term"
                    hdr[1].text = "EN"
                    hdr[2].text = "VI"
                    for item in vocab:
                        if not isinstance(item, dict):
                            continue
                        row = table.add_row().cells
                        row[0].text = str(item.get("term") or "").strip()
                        row[1].text = str(item.get("enGloss") or "").strip()
                        row[2].text = str(item.get("viGloss") or "").strip()

                idea = v.get("ideaFlow") if isinstance(v, dict) else None
                if isinstance(idea, dict):
                    mindmap = str(idea.get("mindmap") or "").strip()
                    flowchart = str(idea.get("flowchart") or "").strip()
                    if mindmap or flowchart:
                        doc.add_heading("Idea Flow", level=3)
                        if mindmap:
                            doc.add_paragraph("Mindmap").runs[0].bold = True
                            add_pre_block(doc, mindmap)
                        if flowchart:
                            doc.add_paragraph("Flowchart").runs[0].bold = True
                            add_pre_block(doc, flowchart)

        if idx != len(ids):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    return doc


def save_doc(doc: Document, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out_path))
    except PermissionError:
        # Common on Windows when the target file is open in Word.
        alt = out_path.with_name(out_path.stem + " (new)" + out_path.suffix)
        doc.save(str(alt))
        print(f"[write-essay] WARNING: Could not overwrite locked file: {out_path}")
        print(f"[write-essay] Wrote new file instead: {alt}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--essay-json", default=ESSAY_JSON_PATH_DEFAULT)
    ap.add_argument("--level", default="all", choices=["a2_b1", "b2", "c1", "all"])
    ap.add_argument("--out", default="")
    ap.add_argument("--ids", default=",".join(str(i) for i in PILOT_PROMPT_IDS_DEFAULT))
    args = ap.parse_args()

    ids = [int(x.strip()) for x in str(args.ids).split(",") if x.strip()]
    questions = load_questions(Path(args.essay_json))

    if args.level == "all":
        for level_id in ("a2_b1", "b2", "c1"):
            doc = build_doc_for_level(level_id=level_id, questions=questions, ids=ids)
            out_path = Path(OUT_DOCX_DEFAULTS[level_id])
            save_doc(doc, out_path)
        return 0

    level_id = str(args.level)
    doc = build_doc_for_level(level_id=level_id, questions=questions, ids=ids)
    out_path = Path(args.out) if str(args.out or "").strip() else Path(OUT_DOCX_DEFAULTS.get(level_id, OUT_DOCX_DEFAULTS["b2"]))
    save_doc(doc, out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
