import json
import urllib.request
import urllib.error
import time
import os
import sys
import re
import openpyxl

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "gemma4:latest"
EXCEL_PATH = r"c:\Cursor AI\public\database\HCS\HCS\HCS.xlsx"
POSITION_REFERENCE_RE = r"\b((first|second|third|fourth|1st|2nd|3rd|4th)\s+(choice|option|paragraph|summary|distractor|one)|(choice|option|paragraph|summary|distractor)\s*(1|2|3|4|one|two|three|four|first|second|third|fourth))\b"

SYSTEM_PROMPT = """You are a premium PTE Academic content developer. Your task is to write a Highlight Correct Summary (HCS) question based on the provided Audio Transcript.
Output a JSON object with the following fields:
1. "question": The question prompt (always use "Click on the paragraph that best relates to the recording.").
2. "choices": An array of exactly 4 choices (paragraphs). One choice must be the correct summary, and the other three must be plausible but incorrect distractors. Each choice is an object:
   - "text": The summary paragraph text.
   - "isCorrect": true for the correct summary, false for incorrect.
3. "explanation": A detailed explanation of why the correct summary is correct (referencing/quoting the transcript) and why the other summaries are incorrect. The learner UI shuffles answer order, so NEVER refer to choices by position or label such as "first choice", "second option", "Choice 1", "Option 2", or "Distractor 3". Refer to summaries by their content instead. Use clean HTML with tags: <p>, <strong>, <b>, <em>, <i>, <ul>, <ol>, <li>, <br>, <h3>, <h4>. Do not use any markdown formatting or code blocks inside the explanation.

Output ONLY a raw, valid JSON string. Do not wrap the JSON in markdown code blocks like ```json.
"""

def generate_hcs_data(transcript):
    prompt = f"""Audio Transcript:
"{transcript}"
"""
    data = {
        "model": MODEL_NAME,
        "system": SYSTEM_PROMPT,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": 0.3
        }
    }
    
    req = urllib.request.Request(
        OLLAMA_URL, 
        data=json.dumps(data).encode("utf-8"), 
        headers={"Content-Type": "application/json"}
    )
    
    try:
        with urllib.request.urlopen(req, timeout=180) as response:
            result = json.loads(response.read().decode("utf-8"))
            output = result.get("response", "").strip()
            
            # Clean markdown code block if present
            if output.startswith("```json"):
                output = output[7:].strip()
            elif output.startswith("```"):
                output = output[3:].strip()
            if output.endswith("```"):
                output = output[:-3].strip()
                
            parsed = json.loads(output)
            return parsed
    except Exception as e:
        print(f"Error calling/parsing Ollama API: {e}")
        # If it failed to parse or run, return None
        return None

def format_answer_column(question, choices):
    # Formats into the Excel format:
    # ---
    # Click on the paragraph that best relates to the recording.
    # ---
    # [] Distractor text...
    # [x] Correct summary text...
    validate_choices(choices)
    lines = [
        "---",
        question,
        "---"
    ]
    for choice in choices:
        marker = "[x]" if choice.get("isCorrect") else "[]"
        lines.append(f"{marker} {choice.get('text')}")
    return "\n".join(lines)

def validate_choices(choices):
    if not isinstance(choices, list) or len(choices) != 4:
        raise ValueError("HCS payload must contain exactly 4 choices")

    correct_count = 0
    for idx, choice in enumerate(choices, start=1):
        if not isinstance(choice, dict):
            raise ValueError(f"HCS choice {idx} must be an object")
        if not str(choice.get("text", "")).strip():
            raise ValueError(f"HCS choice {idx} must contain non-empty text")
        if bool(choice.get("isCorrect")):
            correct_count += 1

    if correct_count != 1:
        raise ValueError("HCS payload must contain exactly 1 correct choice")

def validate_hcs_payload(payload):
    if not isinstance(payload, dict):
        raise ValueError("HCS payload must be a JSON object")
    if not str(payload.get("question", "")).strip():
        raise ValueError("HCS payload must contain a non-empty question")
    if not str(payload.get("explanation", "")).strip():
        raise ValueError("HCS payload must contain a non-empty explanation")
    if re.search(POSITION_REFERENCE_RE, str(payload.get("explanation", "")), flags=re.IGNORECASE):
        raise ValueError("HCS explanation must not refer to choices by position or number")
    validate_choices(payload.get("choices"))

def main():
    print(f"Loading Excel file: {EXCEL_PATH}...")
    if not os.path.exists(EXCEL_PATH):
        print("Excel file not found!")
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers)
    
    # Check if EXPLANATION column exists
    explanation_idx = -1
    for i, h in enumerate(headers):
        if h == "EXPLANATION":
            explanation_idx = i + 1
            break
            
    if explanation_idx == -1:
        ws.cell(row=1, column=len(headers) + 1, value="EXPLANATION")
        explanation_idx = len(headers) + 1
        print("Added EXPLANATION column at index", explanation_idx)
    
    max_row = ws.max_row
    print(f"Total rows to process: {max_row - 1}")
    
    success_count = 0
    for row in range(2, max_row + 1):
        q_id = ws.cell(row=row, column=1).value
        title = ws.cell(row=row, column=2).value
        transcript = ws.cell(row=row, column=4).value
        
        if not q_id or not transcript:
            continue
            
        # Check if explanation already exists to allow resume
        if ws.cell(row=row, column=explanation_idx).value:
            print(f"[{row-1}/{max_row-1}] Q{q_id} already has choices and explanation. Skipping.")
            continue
            
        print(f"[{row-1}/{max_row-1}] Generating HCS choices and explanation for Q{q_id} ({title})...")
        
        # Retry logic for generation
        parsed_data = None
        for attempt in range(3):
            parsed_data = generate_hcs_data(transcript)
            try:
                validate_hcs_payload(parsed_data)
                break
            except ValueError as validation_error:
                parsed_data = None
                print(f"  Attempt {attempt+1} returned invalid HCS data: {validation_error}. Retrying...")
            time.sleep(1)
            
        if parsed_data:
            question = parsed_data.get("question", "Click on the paragraph that best relates to the recording.")
            choices = parsed_data.get("choices", [])
            explanation = parsed_data.get("explanation", "")
            
            # Format choices for the ANSWER column
            answer_text = format_answer_column(question, choices)
            
            # Update columns
            ws.cell(row=row, column=3, value=answer_text)
            ws.cell(row=row, column=explanation_idx, value=explanation)
            
            success_count += 1
            print(f"  Successfully updated row {row} (Q{q_id}).")
            
            # Save progressively
            if success_count % 5 == 0:
                wb.save(EXCEL_PATH)
                print("  Saved progress to workbook.")
        else:
            print(f"  Failed to generate choices for Q{q_id} after 3 attempts.")
            
        time.sleep(0.5)
        
    wb.save(EXCEL_PATH)
    print(f"Completed! Generated {success_count} question sets. Saved workbook to {EXCEL_PATH}")

if __name__ == "__main__":
    main()
