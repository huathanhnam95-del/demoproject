import json
import urllib.request
import urllib.error
import time
import os
import openpyxl
import sys
import html
import re
from html.parser import HTMLParser

sys.stdout.reconfigure(encoding='utf-8')

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434/api/generate")
MODEL_NAME = os.getenv("LOCAL_GEMMA_MODEL", "gemma4:12b")
EXCEL_PATH = r"C:\Cursor AI\public\database\Highlight Incorrect Words\HIW\HIW.xlsx"
ALLOWED_HTML_TAGS = {"p", "strong", "b", "em", "i", "ul", "ol", "li", "br", "h3", "h4"}
SANITIZE_EXISTING_ONLY = "--sanitize-existing-only" in sys.argv

SYSTEM_PROMPT = """You are a premium PTE Academic teacher. Your task is to write a detailed, helpful explanation for a Listening Highlight Incorrect Words question.
In this question type, the user listens to an audio recording and reads a transcript. Some words in the transcript do not match what the speaker says (incorrect words). The user must click on those incorrect words to highlight them.

Analyze the passage and write a clear, helpful explanation detailing:
1. List each incorrect word shown in the transcript along with its correct equivalent spoken in the audio.
2. Provide a brief contextual, grammatical, or semantic explanation for why the correct word is correct and why the incorrect one does not fit.

Use clean HTML formatting. Only use the following allowed tags: <p>, <strong>, <b>, <em>, <i>, <ul>, <ol>, <li>, <br>, <h3>, <h4>.
Do not use markdown code blocks or code wrappers like ```html. Return only the raw HTML explanation.
"""


class ExplanationHtmlSanitizer(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ALLOWED_HTML_TAGS:
            self.parts.append(f"<{tag}>")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ALLOWED_HTML_TAGS and tag != "br":
            self.parts.append(f"</{tag}>")

    def handle_data(self, data):
        self.parts.append(html.escape(data, quote=False))


def strip_markdown_fences(raw_html):
    output = str(raw_html or "").strip()
    if output.startswith("```html"):
        output = output.replace("```html", "", 1)
        if output.endswith("```"):
            output = output[:-3]
    elif output.startswith("```"):
        output = output.replace("```", "", 1)
        if output.endswith("```"):
            output = output[:-3]
    return output.strip()


def sanitize_explanation_html(raw_html):
    parser = ExplanationHtmlSanitizer()
    parser.feed(strip_markdown_fences(raw_html))
    parser.close()
    return "".join(parser.parts).strip()


def parse_incorrect_correct_pairs(answer_text):
    pattern = r"__([^_/]+)\/([^_/]+)__"
    return re.findall(pattern, str(answer_text or ''))


def get_clean_transcript(answer_text):
    pattern = r"__([^_/]+)\/([^_/]+)__"
    return re.sub(pattern, r"\2", str(answer_text or ''))


def generate_explanation(answer_text):
    pairs = parse_incorrect_correct_pairs(answer_text)
    clean_transcript = get_clean_transcript(answer_text)

    pairs_formatted = ""
    for incorrect, correct in pairs:
        pairs_formatted += f"- Word shown in transcript: '{incorrect}' -> Spoken in audio (correct): '{correct}'\n"

    prompt = f"""Transcript displayed to user (with incorrect words highlighted as __incorrect/correct__):
"{answer_text}"

Actual spoken audio transcript (clean version):
"{clean_transcript}"

Incorrect-to-correct word mappings:
{pairs_formatted}
"""

    data = {
        "model": MODEL_NAME,
        "system": SYSTEM_PROMPT,
        "prompt": prompt,
        "stream": False
    }

    req = urllib.request.Request(
        OLLAMA_URL,
        data=json.dumps(data).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )

    try:
        with urllib.request.urlopen(req) as response:
            result = json.loads(response.read().decode("utf-8"))
            output = result.get("response", "").strip()
            return sanitize_explanation_html(output)
    except Exception as e:
        print(f"Error calling Ollama API: {e}")
        return ""


def main():
    print(f"Loading Excel file: {EXCEL_PATH}...")
    if not os.path.exists(EXCEL_PATH):
        print("Excel file not found!")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Ensure Column 4 has the transcript header
    ws.cell(row=1, column=4, value="ANSWER FOR COMPARE OR TRANSCRIPT")
    # Ensure Column 5 has the explanation header
    ws.cell(row=1, column=5, value="EXPLANATION")
    
    explanation_idx = 5
    
    # Read headers
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers)

    max_row = ws.max_row
    print(f"Total rows to process: {max_row - 1}")

    success_count = 0
    sanitized_count = 0
    for row in range(2, max_row + 1):
        q_id = ws.cell(row=row, column=1).value
        title = ws.cell(row=row, column=2).value
        answer = ws.cell(row=row, column=3).value
        existing_exp = ws.cell(row=row, column=explanation_idx).value

        if not q_id or not answer:
            continue

        if existing_exp and len(str(existing_exp).strip()) > 30:
            sanitized_existing = sanitize_explanation_html(existing_exp)
            if sanitized_existing != str(existing_exp).strip():
                ws.cell(row=row, column=explanation_idx, value=sanitized_existing)
                sanitized_count += 1
                print(f"Row {row} (Q{q_id} - {title}): Sanitized existing explanation.")
                if sanitized_count % 10 == 0:
                    wb.save(EXCEL_PATH)
                    print("  Saved sanitized progress to workbook.")
            else:
                print(f"Row {row} (Q{q_id} - {title}): Explanation already clean, skipping.")
            continue

        if SANITIZE_EXISTING_ONLY:
            print(f"Row {row} (Q{q_id} - {title}): No explanation to sanitize, skipping generation.")
            continue

        print(f"[{row-1}/{max_row-1}] Generating explanation for Q{q_id} ({title})...")

        # Call GemmaAI
        explanation = generate_explanation(answer)

        if explanation:
            ws.cell(row=row, column=explanation_idx, value=sanitize_explanation_html(explanation))
            success_count += 1
            print(f"  Successfully saved explanation (~{len(explanation)} chars).")
            # Save progressively
            if success_count % 5 == 0:
                wb.save(EXCEL_PATH)
                print("  Saved progress to workbook.")
        else:
            print("  Failed to generate explanation.")

        time.sleep(0.5)  # rate limiting

    wb.save(EXCEL_PATH)
    print(f"Completed! Generated {success_count} explanations, sanitized {sanitized_count}. Saved workbook to {EXCEL_PATH}")


if __name__ == "__main__":
    main()
