import json
import urllib.request
import urllib.error
import time
import os
import pandas as pd
import openpyxl

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434/api/generate")
MODEL_NAME = os.getenv("LOCAL_GEMMA_MODEL", "gemma4:12b")
EXCEL_PATH = r"c:\Cursor AI\public\database\LMCSA\LMCSA\LMCSA.xlsx"

SYSTEM_PROMPT = """You are a premium PTE Academic teacher. Your task is to write a detailed, helpful explanation for a Listening Multiple Choice Single Answer question.
Analyze:
1. The correct option: explain why it is correct by quoting or referencing the relevant sentence in the transcript.
2. The incorrect options: briefly explain why they are incorrect.
Use clean HTML formatting. Only use the following allowed tags: <p>, <strong>, <b>, <em>, <i>, <ul>, <ol>, <li>, <br>, <h3>, <h4>.
Do not use markdown code blocks or code wrappers like ```html. Return only the raw HTML explanation.
"""

def parse_answer_choices(answer_text):
    parts = str(answer_text or '').split('\n---')
    if len(parts) < 2:
        parts = str(answer_text or '').split('---')
    
    choices_raw = ""
    question = ""
    if len(parts) >= 2:
        question = parts[-2].replace('---', '').strip()
        choices_raw = parts[-1].strip()
    else:
        choices_raw = str(answer_text or '').strip()
        
    choices = []
    for line in choices_raw.split('\n'):
        line = line.strip()
        if not line:
            continue
        if line.startswith('[') and ']' in line:
            idx = line.find(']')
            marker = line[1:idx].strip().lower()
            text = line[idx+1:].strip()
            is_correct = 'x' in marker
            choices.append((text, is_correct))
    return question, choices

def generate_explanation(transcript, question, choices):
    choices_formatted = ""
    for text, is_correct in choices:
        status = "CORRECT" if is_correct else "INCORRECT"
        choices_formatted += f"- {text} ({status})\n"
        
    prompt = f"""Audio Transcript:
"{transcript}"

Question:
"{question}"

Choices:
{choices_formatted}
"""
    
    data = {
        "model": MODEL_NAME,
        "system": SYSTEM_PROMPT,
        "prompt": prompt,
        "stream": False
    }
    
    req = urllib.request.Request(
        OLLAMA_URL, 
        data=json.dumps(data).encode("utf-8"), 
        headers={"Content-Type": "application/json"}
    )
    
    try:
        with urllib.request.urlopen(req) as response:
            result = json.loads(response.read().decode("utf-8"))
            output = result.get("response", "").strip()
            
            # Additional cleanup of markdown code blocks
            if output.startswith("```html"):
                output = output.replace("```html", "").replace("```", "").strip()
            elif output.startswith("```"):
                output = output.replace("```", "").strip()
                
            return output
    except Exception as e:
        print(f"Error calling Ollama API: {e}")
        return ""

def main():
    print(f"Loading Excel file: {EXCEL_PATH}...")
    if not os.path.exists(EXCEL_PATH):
        print("Excel file not found!")
        return

    # Load with openpyxl to check/save
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    # Read headers
    headers = [cell.value for cell in ws[1]]
    print("Headers:", headers)
    
    # Check if EXPLANATION column already exists
    explanation_idx = -1
    for i, h in enumerate(headers):
        if h == "EXPLANATION":
            explanation_idx = i + 1
            break
            
    if explanation_idx == -1:
        # Append EXPLANATION column
        ws.cell(row=1, column=len(headers) + 1, value="EXPLANATION")
        explanation_idx = len(headers) + 1
        print("Added EXPLANATION column at index", explanation_idx)
    
    # Columns map:
    # Col 1: ID
    # Col 2: TITLE
    # Col 3: ANSWER
    # Col 4: ANSWER FOR COMPARE OR TRANSCRIPT (Transcript)
    
    max_row = ws.max_row
    print(f"Total rows to process: {max_row - 1}")
    
    success_count = 0
    for row in range(2, max_row + 1):
        q_id = ws.cell(row=row, column=1).value
        title = ws.cell(row=row, column=2).value
        answer = ws.cell(row=row, column=3).value
        transcript = ws.cell(row=row, column=4).value
        existing_exp = ws.cell(row=row, column=explanation_idx).value
        
        if not q_id or not answer or not transcript:
            continue
            
        if existing_exp and len(str(existing_exp).strip()) > 30:
            print(f"Row {row} (Q{q_id} - {title}): Explanation already exists, skipping.")
            continue
            
        print(f"[{row-1}/{max_row-1}] Generating explanation for Q{q_id} ({title})...")
        
        question, choices = parse_answer_choices(answer)
        if not choices:
            print(f"  Warning: No choices parsed for row {row}, skipping.")
            continue
            
        # Call GemmaAI
        explanation = generate_explanation(transcript, question, choices)
        
        if explanation:
            ws.cell(row=row, column=explanation_idx, value=explanation)
            success_count += 1
            print(f"  Successfully saved explanation (~{len(explanation)} chars).")
            # Save progressively
            if success_count % 5 == 0:
                wb.save(EXCEL_PATH)
                print("  Saved progress to workbook.")
        else:
            print("  Failed to generate explanation.")
            
        time.sleep(0.5) # rate limiting
        
    wb.save(EXCEL_PATH)
    print(f"Completed! Generated {success_count} explanations. Saved workbook to {EXCEL_PATH}")

if __name__ == "__main__":
    main()
