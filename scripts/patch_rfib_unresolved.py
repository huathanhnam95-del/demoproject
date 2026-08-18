"""Deterministically stage or merge the two approved RFIB unresolved fixes."""

from __future__ import annotations

import argparse
import copy
import json
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook


TARGETS = {
    (146, 3): {
        "correct_answer": "open",
        "grammar_tag": "Collocation (Verb + Predicate Adjective)",
        "final_explanation": (
            "The correct answer is 'open' because it pairs with the verb 'leaves' to form the "
            "fixed collocation 'leave open the possibility,' which means to allow for a chance, "
            "alternative interpretation, or unresolved outcome. In this sentence, the object noun "
            "clause ('the possibility that they are not...') is placed after the adjective 'open' "
            "for syntactic balance.\n\n"
            "Why the distractors are incorrect:\n"
            "- 'go': This is a base-form verb and cannot function as a predicate adjective complement "
            "after 'leaves' (one cannot say *leaves go the possibility* in standard English).\n"
            "- 'covered': While 'covered' is an adjective/participle, the phrase *'leaves covered the "
            "possibility'* is not a recognized English collocation and creates a semantic contradiction "
            "by implying concealment rather than availability.\n"
            "- 'undoubted': An adjective meaning 'certain' or 'indisputable.' Using it here would "
            "contradict the entire premise of the sentence, which emphasizes maintaining ambiguity and "
            "non-confinement rather than certainty."
        ),
        "concise_explanation": (
            "'Leaves open the possibility' is a fixed collocation meaning to allow for an alternative "
            "chance or interpretation. The distractors fail grammatically ('go' is a verb) or "
            "semantically ('covered' and 'undoubted' contradict the idea of keeping an option available)."
        ),
    },
    (520, 4): {
        "correct_answer": "tend",
        "grammar_tag": "Grammar/Usage (Catenative Verb of Tendency / Habit)",
        "final_explanation": (
            "The correct answer is 'tend' because it functions as a lexical verb expressing a general "
            "inclination, habit, or typical characteristic when paired with a to-infinitive ('tend to be'). "
            "In this context, 'tend to be higher-priced' accurately states that B2B transactions are "
            "typically or usually more expensive than consumer sales.\n\n"
            "Why the distractors are incorrect:\n"
            "- 'extend': Means to lengthen, prolong, or stretch out physically or temporally. It cannot "
            "describe a general characteristic of pricing (e.g., sales do not 'extend to be expensive').\n"
            "- 'contend': Means to assert or argue a position firmly ('contend that') or to struggle "
            "against an opponent ('contend with'). It does not express an intrinsic tendency of goods or services.\n"
            "- 'pretend': Means to behave deceptively or simulate an imaginary state. It is logically and "
            "contextually incompatible with describing factual corporate pricing structures."
        ),
        "concise_explanation": (
            "'Tend to be' is a standard construction expressing a general habit or typical state (B2B sales "
            "are usually higher-priced). The rhyming distractors mean to lengthen ('extend'), argue/compete "
            "('contend'), or fake ('pretend'), which are logically impossible here."
        ),
    },
}

MODEL_PROVENANCE = {
    "deepseek-r1:14b": "dr",
    "qwen3:14b": "qw",
    "gemma4:latest": "gm",
}
EXPECTED_TARGET_RECORD_STATUS = {146: "PASS", 520: "REVISED_WITH_CONSENSUS"}
EXPECTED_TARGET_SUMMARIES = {
    146: (5, 0, 0, 0, 5),
    520: (3, 1, 0, 0, 4),
}
ALLOWED_AUDIT_STATUSES = frozenset({"PASS", "REVISED_WITH_CONSENSUS", "UNRESOLVED"})
EXPECTED_AGGREGATE_RECORD_COUNT = 221


def _load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _write_jsonl(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records)
    _atomic_write_text(path, payload)


def _atomic_write_text(path: Path, payload: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _target_blank(record: dict, blank_index: int) -> dict:
    for blank in record.get("blanks", []):
        if int(blank.get("blank_index", -1)) == blank_index:
            return blank
    raise ValueError(f"Missing target blank {record.get('id')}:{blank_index}")


def apply_proposed_text(record: dict) -> dict:
    """Return a copy with only the two proposed revision fields patched."""
    updated = copy.deepcopy(record)
    qid = int(updated["id"])
    for (target_id, blank_index), proposed in TARGETS.items():
        if target_id == qid:
            blank = _target_blank(updated, blank_index)
            blank.update(proposed)
    return updated


def make_revision_candidate(input_path: Path, output_path: Path) -> None:
    _write_jsonl(output_path, [apply_proposed_text(record) for record in _load_jsonl(input_path)])


def _validate_fresh_audit(audited: dict[int, dict], report: dict) -> None:
    if set(audited) != {146, 520}:
        raise ValueError(f"Fresh audit must contain exactly IDs 146 and 520, got {sorted(audited)}")
    if report.get("total_target_questions") != 2 or report.get("processed_questions") != 2:
        raise ValueError("Fresh report does not cover exactly two processed target questions")
    if report.get("unresolved_count") != 0:
        raise ValueError("Fresh 3-model audit left one or more unresolved questions")
    totals = (
        report.get("round1_pass_count", 0)
        + report.get("debate_revised_count", 0)
        + report.get("unresolved_count", 0)
    )
    if totals != report.get("processed_questions"):
        raise ValueError("Fresh report totals do not reconcile")
    for qid, record in audited.items():
        expected_status = EXPECTED_TARGET_RECORD_STATUS.get(qid)
        if expected_status is None or record.get("audit_status") != expected_status:
            raise ValueError(
                f"Fresh audit Q{qid} has unexpected record status: {record.get('audit_status')}"
            )
        summary = record.get("audit_summary") or {}
        expected_summary = EXPECTED_TARGET_SUMMARIES[qid]
        actual_summary = tuple(
            summary.get(name)
            for name in (
                "round1_passed_blanks",
                "mutual_consent_fixed_blanks",
                "majority_consent_fixed_blanks",
                "unresolved_blanks",
                "total_blanks",
            )
        )
        if actual_summary != expected_summary:
            raise ValueError(f"Fresh audit Q{qid} summary is not the expected model outcome: {actual_summary}")
        if record.get("audit_timestamp") is None:
            raise ValueError(f"Fresh audit Q{qid} is missing its audit timestamp")
        if record.get("phase_models") != list(MODEL_PROVENANCE):
            raise ValueError(f"Fresh audit Q{qid} has incomplete phase model provenance")
        raw_artifacts = record.get("raw_artifacts") or {}
        if any(not isinstance(raw_artifacts.get(model_key), str) or not raw_artifacts[model_key].strip() for model_key in ("dr_raw", "qw_raw", "gm_raw")):
            raise ValueError(f"Fresh audit Q{qid} is missing one or more raw model artifacts")

        details = summary.get("details")
        if not isinstance(details, list) or len(details) != len(record.get("blanks", [])):
            raise ValueError(f"Fresh audit Q{qid} has incomplete blank detail evidence")
        details_by_index = {int(detail.get("blank_index", -1)): detail for detail in details}
        computed_summary = (
            sum(detail.get("final_verdict") == "PASS" for detail in details),
            sum(detail.get("final_verdict") == "FIXED_CONSENSUS" for detail in details),
            sum(detail.get("final_verdict") == "FIXED_MAJORITY" for detail in details),
            sum(detail.get("final_verdict") == "UNRESOLVED_DEBATE" for detail in details),
            len(details),
        )
        if computed_summary != expected_summary:
            raise ValueError(f"Fresh audit Q{qid} detail verdicts do not support its summary: {computed_summary}")
        if qid == 146 and any(detail.get("debate_history") for detail in details):
            raise ValueError("Fresh audit Q146 unexpectedly contains debate history")
        if qid == 520:
            fixed_details = [detail for detail in details if detail.get("final_verdict") == "FIXED_CONSENSUS"]
            if len(fixed_details) != 1 or fixed_details[0].get("blank_index") != 3:
                raise ValueError("Fresh audit Q520 must have exactly one fixed consensus on blank 3")
            fixed = fixed_details[0]
            if fixed.get("consensus_type") != "UNANIMOUS_MUTUAL_CONSENT":
                raise ValueError("Fresh audit Q520 blank 3 lacks unanimous consensus provenance")
            if not isinstance(fixed.get("debate_rounds"), int) or fixed["debate_rounds"] < 1:
                raise ValueError("Fresh audit Q520 blank 3 lacks a successful debate round")
            if len(fixed.get("debate_history") or []) != fixed["debate_rounds"]:
                raise ValueError("Fresh audit Q520 blank 3 debate history is incomplete")
        for blank in record.get("blanks", []):
            blank_index = int(blank.get("blank_index", -1))
            detail = details_by_index.get(blank_index)
            if detail is None:
                raise ValueError(f"Fresh audit Q{qid} is missing detail evidence for blank {blank_index}")
            if detail.get("final_verdict") == "UNRESOLVED_DEBATE":
                raise ValueError(f"Fresh audit Q{qid} blank {blank_index} is unresolved")
            if detail.get("final_verdict") != blank.get("audit_verdict"):
                raise ValueError(f"Fresh audit Q{qid} blank {blank_index} verdict does not match its detail")
            if detail.get("final_grammar_tag") != blank.get("grammar_tag"):
                raise ValueError(f"Fresh audit Q{qid} blank {blank_index} grammar tag lacks provenance")
            if detail.get("final_explanation") != blank.get("final_explanation"):
                raise ValueError(f"Fresh audit Q{qid} blank {blank_index} explanation lacks provenance")
            if detail.get("final_concise") != blank.get("concise_explanation"):
                raise ValueError(f"Fresh audit Q{qid} blank {blank_index} concise text lacks provenance")

        for target_index in (3, 4):
            if (qid, target_index) not in TARGETS:
                continue
            blank = _target_blank(record, target_index)
            detail = details_by_index[target_index]
            if blank.get("grammar_tag") != TARGETS[(qid, target_index)]["grammar_tag"]:
                raise ValueError(f"Fresh audit Q{qid} blank {target_index} has the wrong grammar tag")
            if blank.get("final_explanation") != TARGETS[(qid, target_index)]["final_explanation"]:
                raise ValueError(f"Fresh audit Q{qid} blank {target_index} has the wrong explanation")
            if blank.get("concise_explanation") != TARGETS[(qid, target_index)]["concise_explanation"]:
                raise ValueError(f"Fresh audit Q{qid} blank {target_index} has the wrong concise explanation")
            _validate_round1_vote_detail(detail, qid, target_index)

    _validate_fresh_transcripts(audited, report)


def _validate_round1_vote_detail(detail: dict, qid: int, blank_index: int) -> None:
    if detail.get("final_verdict") != "PASS" or detail.get("consensus_type") != "ROUND1_APPROVAL":
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} is not a Round-1 approval")
    if detail.get("debate_rounds") != 0 or detail.get("debate_history") != []:
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} has unexpected debate provenance")
    votes = detail.get("round_1_votes") or {}
    if votes.get("pass") != 3 or votes.get("fail") != 0:
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} does not have 3/3 Round-1 votes")
    vote_details = votes.get("details")
    if not isinstance(vote_details, list) or len(vote_details) != 3:
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} is missing one or more juror votes")
    expected_models = set(MODEL_PROVENANCE.items())
    actual_models = {(vote.get("model"), vote.get("model_key")) for vote in vote_details}
    if actual_models != expected_models:
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} has incomplete model provenance")
    if any(vote.get("verdict") != "PASS" for vote in vote_details):
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} has a non-PASS juror vote")
    if any(not isinstance(vote.get("overall_notes"), str) or not vote["overall_notes"].strip() for vote in vote_details):
        raise ValueError(f"Fresh target Q{qid} blank {blank_index} has an empty juror rationale")


def _validate_fresh_transcripts(audited: dict[int, dict], report: dict) -> None:
    transcripts = report.get("debate_transcripts")
    if not isinstance(transcripts, list):
        raise ValueError("Fresh report is missing debate transcripts")
    for transcript in transcripts:
        qid = int(transcript.get("question_id", -1))
        blank_index = int(transcript.get("blank_index", -1))
        record = audited.get(qid)
        if record is None:
            raise ValueError(f"Fresh transcript references unknown Q{qid}")
        detail = next(
            (item for item in (record.get("audit_summary") or {}).get("details", []) if int(item.get("blank_index", -1)) == blank_index),
            None,
        )
        if detail is None:
            raise ValueError(f"Fresh transcript references missing Q{qid} blank {blank_index} detail")
        for field in ("round_1_votes", "debate_rounds", "consensus_type", "final_explanation", "debate_history"):
            if transcript.get(field) != detail.get(field):
                raise ValueError(f"Fresh transcript Q{qid} blank {blank_index} diverges on {field}")


def _transcripts_from_records(records: dict[int, dict]) -> list[dict]:
    transcripts = []
    for qid in sorted(records):
        summary = records[qid].get("audit_summary") or {}
        for detail in summary.get("details", []):
            if detail.get("debate_history"):
                transcripts.append(
                    {
                        "question_id": qid,
                        "blank_index": detail.get("blank_index"),
                        "round_1_votes": detail.get("round_1_votes"),
                        "debate_rounds": detail.get("debate_rounds"),
                        "consensus_type": detail.get("consensus_type"),
                        "final_explanation": detail.get("final_explanation"),
                        "debate_history": detail.get("debate_history"),
                    }
                )
    return transcripts


def _aggregate_status_counts(audited: dict[int, dict]) -> dict[str, int]:
    statuses = [record.get("audit_status") for record in audited.values()]
    unknown = sorted(set(statuses) - ALLOWED_AUDIT_STATUSES)
    if unknown:
        raise ValueError(f"Audited sidecar contains unsupported audit statuses: {unknown}")
    return {
        "PASS": statuses.count("PASS"),
        "REVISED_WITH_CONSENSUS": statuses.count("REVISED_WITH_CONSENSUS"),
        "UNRESOLVED": statuses.count("UNRESOLVED"),
    }


def _validate_aggregate_report(report: dict, audited: dict[int, dict]) -> None:
    if len(audited) != EXPECTED_AGGREGATE_RECORD_COUNT:
        raise ValueError("Aggregate report cannot be built from a partial audited sidecar")
    status_counts = _aggregate_status_counts(audited)
    expected_report_counts = {
        "total_target_questions": len(audited),
        "processed_questions": len(audited),
        "round1_pass_count": status_counts["PASS"],
        "debate_revised_count": status_counts["REVISED_WITH_CONSENSUS"],
        "unresolved_count": status_counts["UNRESOLVED"],
    }
    if any(report.get(field) != expected for field, expected in expected_report_counts.items()):
        raise ValueError("Aggregate report counts do not match the audited record statuses")
    if report.get("processed_questions") != report.get("round1_pass_count", 0) + report.get("debate_revised_count", 0) + report.get("unresolved_count", 0):
        raise ValueError("Aggregate report totals do not reconcile")
    if report.get("unresolved_count") == 0 and any(
        (record.get("audit_summary") or {}).get("unresolved_blanks", 0) for record in audited.values()
    ):
        raise ValueError("Aggregate report cannot claim zero unresolved records while sidecar has unresolved blanks")
    expected_transcripts = _transcripts_from_records(audited)
    actual_transcripts = report.get("debate_transcripts")
    if actual_transcripts != expected_transcripts:
        raise ValueError("Aggregate report transcripts do not match audited sidecar evidence")


def _validate_aggregate_shape(report: dict, audited: dict[int, dict]) -> None:
    if report.get("total_target_questions") != EXPECTED_AGGREGATE_RECORD_COUNT or report.get("processed_questions") != EXPECTED_AGGREGATE_RECORD_COUNT:
        raise ValueError("Existing report is not the full 221-question aggregate; refusing target-only replacement")
    if len(audited) != EXPECTED_AGGREGATE_RECORD_COUNT:
        raise ValueError("Existing audited sidecar is not the full 221-question sample")
    _validate_aggregate_report(report, audited)


def build_aggregate_report(audited_records: list[dict], base_report: dict | None = None) -> dict:
    audited = {int(record["id"]): record for record in audited_records}
    report = copy.deepcopy(base_report) if base_report else {"sample_pct": 0.2, "sample_seed": 42}
    status_counts = _aggregate_status_counts(audited)
    report.update(
        {
            "total_target_questions": len(audited),
            "processed_questions": len(audited),
            "round1_pass_count": status_counts["PASS"],
            "debate_revised_count": status_counts["REVISED_WITH_CONSENSUS"],
            "unresolved_count": status_counts["UNRESOLVED"],
        }
    )
    report["debate_transcripts"] = _transcripts_from_records(audited)
    _validate_aggregate_report(report, audited)
    return report


def _replace_workbook_target(worksheet, row: int, blank_index: int, answer: str, proposed: dict) -> None:
    analysis = json.loads(worksheet.cell(row, 11).value or "[]")
    entry = next(item for item in analysis if int(item.get("blank_index", -1)) == blank_index)
    entry["grammar_tag"] = proposed["grammar_tag"]
    entry["concise_explanation"] = proposed["concise_explanation"]
    worksheet.cell(row, 11).value = json.dumps(analysis, ensure_ascii=False)

    detailed = worksheet.cell(row, 13).value or ""
    marker = f"Blank {blank_index} ('{answer}'):\n"
    start = detailed.find(marker)
    if start < 0:
        raise ValueError(f"Workbook row {row} is missing {marker}")
    next_start = detailed.find("\nBlank ", start + len(marker))
    if next_start < 0:
        next_start = len(detailed)
    replacement = (
        f"{marker}\n"
        f"  Gold-standard explanation: {proposed['final_explanation']}\n"
        f"  Concise UI explanation: {proposed['concise_explanation']}\n"
    )
    worksheet.cell(row, 13).value = detailed[:start] + replacement + detailed[next_start:]


def merge_production(
    revision_path: Path,
    audited_existing_path: Path,
    audited_fresh_path: Path,
    report_path: Path,
    fresh_report_path: Path,
    workbook_path: Path,
) -> None:
    fresh_records = {int(record["id"]): record for record in _load_jsonl(audited_fresh_path)}
    fresh_report = json.loads(fresh_report_path.read_text(encoding="utf-8"))
    _validate_fresh_audit(fresh_records, fresh_report)

    revision_records = [apply_proposed_text(record) for record in _load_jsonl(revision_path)]
    audited_records = _load_jsonl(audited_existing_path)
    audited_by_id = {int(record["id"]): record for record in audited_records}
    existing_report = json.loads(report_path.read_text(encoding="utf-8"))
    _validate_aggregate_shape(existing_report, audited_by_id)
    for qid in (146, 520):
        audited_by_id[qid] = fresh_records[qid]
    merged_audited = [audited_by_id[int(record["id"])] for record in audited_records]
    aggregate_report = build_aggregate_report(merged_audited, existing_report)

    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    worksheet = workbook["Sheet1"]
    _replace_workbook_target(worksheet, 147, 3, "open", TARGETS[(146, 3)])
    _replace_workbook_target(worksheet, 480, 4, "tend", TARGETS[(520, 4)])

    # Validate and stage every artifact before replacing any production path.
    staged = []
    try:
        for path, writer in (
            (revision_path, lambda p: _write_jsonl(p, revision_records)),
            (audited_existing_path, lambda p: _write_jsonl(p, merged_audited)),
            (report_path, lambda p: _atomic_write_text(p, json.dumps(aggregate_report, ensure_ascii=False, indent=2) + "\n")),
        ):
            fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
            os.close(fd)
            temporary_path = Path(temporary)
            staged.append((temporary_path, path))
            writer(temporary_path)

        workbook_fd, workbook_temporary = tempfile.mkstemp(
            prefix=f".{workbook_path.name}.", suffix=".tmp", dir=workbook_path.parent
        )
        os.close(workbook_fd)
        workbook_temporary_path = Path(workbook_temporary)
        staged.append((workbook_temporary_path, workbook_path))
        workbook.save(workbook_temporary_path)

        for temporary_path, destination in staged:
            os.replace(temporary_path, destination)
    finally:
        for temporary_path, _ in staged:
            if temporary_path.exists():
                temporary_path.unlink()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    candidate = subparsers.add_parser("candidate")
    candidate.add_argument("--input", type=Path, required=True)
    candidate.add_argument("--output", type=Path, required=True)

    merge = subparsers.add_parser("merge")
    merge.add_argument("--revision", type=Path, required=True)
    merge.add_argument("--audited-existing", type=Path, required=True)
    merge.add_argument("--audited-fresh", type=Path, required=True)
    merge.add_argument("--report", type=Path, required=True)
    merge.add_argument("--fresh-report", type=Path, required=True)
    merge.add_argument("--workbook", type=Path, required=True)

    aggregate = subparsers.add_parser("aggregate")
    aggregate.add_argument("--audited", type=Path, required=True)
    aggregate.add_argument("--output", type=Path, required=True)

    args = parser.parse_args()
    if args.command == "candidate":
        make_revision_candidate(args.input, args.output)
    elif args.command == "merge":
        merge_production(
            args.revision,
            args.audited_existing,
            args.audited_fresh,
            args.report,
            args.fresh_report,
            args.workbook,
        )
    else:
        records = _load_jsonl(args.audited)
        report = build_aggregate_report(records)
        _atomic_write_text(args.output, json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
