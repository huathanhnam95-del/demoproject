import os
import re
import sys
import json
import logging
import argparse
import openpyxl
import requests

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
)

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434/api/generate")
MODEL_NAME = os.getenv("LOCAL_GEMMA_MODEL", "gemma4:12b")

INPUT_FILE = r"public\database\RFIB\RFIB Final ver.xlsx"
SIDECAR_FILE = r"public\database\RFIB\RFIB_cohesion_enrichment.jsonl"
REPORT_FILE = r"public\database\RFIB\RFIB_cohesion_review_report.jsonl"

BLANK_RE = re.compile(r"__([^_]+?)__")


def extract_blanks(answer_text: str) -> list[dict]:
    if not answer_text:
        return []
    matches = BLANK_RE.findall(answer_text)
    blanks = []
    for i, payload in enumerate(matches, start=1):
        options = [s.strip() for s in payload.split("/")]
        blanks.append({"index": i, "correct": options[0], "options": options})
    return blanks


def build_header_map(ws) -> dict:
    hmap = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            hmap[val.strip()] = col
    return hmap


def check_distractor_heuristic(distractor: str, text: str) -> bool:
    """Return True if the distractor is mentioned in the explanation text."""
    # Escape special characters for regex
    escaped = re.escape(distractor)
    
    # Try match with word boundaries
    pattern = re.compile(rf"\b{escaped}\b", re.IGNORECASE)
    if pattern.search(text):
        return True
        
    # Fallback to direct substring search if distractor contains non-alphanumeric chars
    if not distractor.isalnum():
        if distractor.lower() in text.lower():
            return True
            
    return False


def query_gemma_review(passage: str, blank_index: int, correct_answer: str, options: list[str], explanation: str) -> dict:
    """Ask local Gemma AI if the explanation covers all distractors."""
    distractors = options[1:]
    prompt = f"""You are a quality assurance auditor. Verify if the provided student-facing explanation for the blank in this context successfully explains:
1. Why the correct answer is correct.
2. Why each of the following incorrect options (distractors) is wrong: {distractors}

Context: {passage}
Blank Index: {blank_index}
Correct Answer: {correct_answer}
Options: {options}
Explanation: {explanation}

Format your response strictly as JSON with no markdown fences, matching:
{{
  "compliant": true | false,
  "feedback": "Details if compliant is false (e.g. which distractors are not explained or why the explanation is insufficient), otherwise empty string"
}}
"""
    payload = {
        "model": MODEL_NAME,
        "prompt": prompt,
        "format": "json",
        "stream": False,
        "options": {"temperature": 0.1},
    }
    try:
        resp = requests.post(OLLAMA_URL, json=payload, timeout=60)
        if resp.status_code == 200:
            result_str = resp.json().get("response", "")
            return json.loads(result_str)
    except Exception as e:
        logging.warning(f"Ollama review call failed for blank {blank_index}: {e}")
        
    return {"compliant": False, "feedback": "Auditor query failed or returned invalid JSON"}


def main():
    parser = argparse.ArgumentParser(description="Review generated RFIB explanations.")
    parser.add_argument("--input", type=str, default=INPUT_FILE, help="Workbook path.")
    parser.add_argument("--sidecar", type=str, default=SIDECAR_FILE, help="Sidecar JSONL path.")
    parser.add_argument("--report", type=str, default=REPORT_FILE, help="Output review report path.")
    parser.add_argument("--mode", type=str, default="heuristic", choices=["heuristic", "llm"], help="Review mode.")
    parser.add_argument("--ids", type=str, default="", help="Comma-separated list of IDs to review.")
    parser.add_argument("--limit", type=int, default=0, help="Max records to review.")
    args = parser.parse_args()

    # Load sidecar records
    if not os.path.exists(args.sidecar):
        logging.error(f"Sidecar file not found: {args.sidecar}")
        sys.exit(1)

    sidecar_records = {}
    with open(args.sidecar, "r", encoding="utf-8") as f:
        for line in f:
            rec = json.loads(line.strip())
            sidecar_records[rec["id"]] = rec

    logging.info(f"Loaded {len(sidecar_records)} records from sidecar.")

    # Load workbook
    if not os.path.exists(args.input):
        logging.error(f"Workbook not found: {args.input}")
        sys.exit(1)

    wb = openpyxl.load_workbook(args.input, data_only=True)
    ws = wb.active
    hmap = build_header_map(ws)
    
    col_id = hmap.get("ID")
    col_answer = hmap.get("ANSWER")
    col_passage = hmap.get("Full Text")

    if not col_id or not col_answer or not col_passage:
        logging.error("Required columns (ID, ANSWER, Full Text) missing from workbook.")
        sys.exit(1)

    # Build ID -> row mapping
    id_to_row = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=col_id).value
        if rid is not None:
            id_to_row[int(rid)] = r

    # Open report file
    out_f = open(args.report, "w", encoding="utf-8")

    stats = {
        "total_records": 0,
        "total_blanks": 0,
        "compliant_blanks": 0,
        "non_compliant_blanks": 0,
        "non_compliant_records": [],
    }

    count = 0
    target_ids = set(int(x.strip()) for x in args.ids.split(",") if x.strip()) if args.ids else None

    for record_id, record in sidecar_records.items():
        if target_ids and record_id not in target_ids:
            continue
            
        if args.limit > 0 and count >= args.limit:
            break
            
        row = id_to_row.get(record_id)
        if not row:
            logging.warning(f"Record ID {record_id} not found in workbook. Skipping.")
            continue

        stats["total_records"] += 1
        passage = ws.cell(row=row, column=col_passage).value or ""
        answer_val = ws.cell(row=row, column=col_answer).value or ""
        blanks = extract_blanks(answer_val)
        
        explanations = record.get("explanations", [])
        
        record_compliant = True
        record_feedback = []
        blank_verdicts = []

        for idx, blank in enumerate(blanks):
            stats["total_blanks"] += 1
            exp_item = next((e for e in explanations if e.get("blank_index") == blank["index"]), None)
            
            if not exp_item:
                stats["non_compliant_blanks"] += 1
                record_compliant = False
                record_feedback.append(f"Blank {blank['index']} is missing explanation in sidecar")
                blank_verdicts.append({
                    "blank_index": blank["index"],
                    "compliant": False,
                    "feedback": "Missing explanation"
                })
                continue

            explanation_text = exp_item.get("detailed_student_explanation", "")
            distractors = blank["options"][1:]
            
            # Check Heuristic first
            missing_distractors = []
            for dist in distractors:
                if not check_distractor_heuristic(dist, explanation_text):
                    missing_distractors.append(dist)

            if missing_distractors:
                compliant = False
                feedback = f"Heuristic Fail: Distractors not mentioned in text: {missing_distractors}"
            elif args.mode == "llm":
                # Call Ollama Gemma reviewer
                review_verdict = query_gemma_review(
                    passage=passage,
                    blank_index=blank["index"],
                    correct_answer=blank["correct"],
                    options=blank["options"],
                    explanation=explanation_text
                )
                compliant = review_verdict.get("compliant", False)
                feedback = review_verdict.get("feedback", "")
            else:
                compliant = True
                feedback = ""

            if compliant:
                stats["compliant_blanks"] += 1
            else:
                stats["non_compliant_blanks"] += 1
                record_compliant = False
                record_feedback.append(f"Blank {blank['index']}: {feedback}")

            blank_verdicts.append({
                "blank_index": blank["index"],
                "compliant": compliant,
                "feedback": feedback
            })

        # Save record review verdict
        verdict_rec = {
            "id": record_id,
            "compliant": record_compliant,
            "feedback": "; ".join(record_feedback) if record_feedback else "All blanks compliant",
            "blanks": blank_verdicts
        }
        out_f.write(json.dumps(verdict_rec, ensure_ascii=False) + "\n")
        out_f.flush()

        if not record_compliant:
            stats["non_compliant_records"].append(record_id)
            logging.info(f"[-] ID {record_id}: NON-COMPLIANT - {verdict_rec['feedback']}")
        else:
            logging.info(f"[+] ID {record_id}: COMPLIANT")

        count += 1

    out_f.close()
    wb.close()

    logging.info("=" * 60)
    logging.info("REVIEW SUMMARY")
    logging.info("=" * 60)
    logging.info(f"Total records reviewed: {stats['total_records']}")
    logging.info(f"Total blanks checked:   {stats['total_blanks']}")
    logging.info(f"Compliant blanks:       {stats['compliant_blanks']}")
    logging.info(f"Non-compliant blanks:   {stats['non_compliant_blanks']}")
    logging.info(f"Non-compliant records:  {len(stats['non_compliant_records'])}")
    if stats["non_compliant_records"]:
        logging.info(f"Non-compliant IDs:      {stats['non_compliant_records']}")
    logging.info("=" * 60)


if __name__ == "__main__":
    main()
