# revise_rfib_explanations.py — 3-Model Local LLM Pipeline for RFIB Explanation Revision

import os
import re
import sys
import time
import json
import logging
import argparse
import requests
from datetime import datetime, timezone

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler(sys.stdout)],
)

OLLAMA_URL = "http://localhost:11434/api/generate"

MODELS = {
    "dr": "deepseek-r1:14b",
    "qw": "qwen3:14b",
    "gm": "gemma4:latest",
}

INPUT_WORKBOOK = r"C:\Cursor AI\public\database\RFIB\RFIB Final ver.xlsx"
INPUT_JSONL = r"C:\Cursor AI\public\database\RFIB\RFIB_cohesion_enrichment.jsonl"
OUTPUT_JSONL = r"C:\Cursor AI\public\database\RFIB\RFIB_3model_revision.jsonl"

NON_COMPLIANT_IDS = [375, 398, 597, 677, 944, 1315]
CURATED_QUALITY_GATE_IDS = [375, 398, 597, 677, 944, 1315, 1, 2, 9, 13]

BLANK_RE = re.compile(r"__([^_]+?)__")


def extract_blanks(answer_text: str) -> list[dict]:
    if not answer_text:
        return []
    matches = BLANK_RE.findall(answer_text)
    blanks = []
    for i, payload in enumerate(matches, start=1):
        options = [s.strip() for s in payload.split("/")]
        blanks.append({"index": i, "correct": options[0], "options": options})
    return blanks


def clean_model_response(raw: str) -> str:
    if not raw:
        return raw
    cleaned = re.sub(r"<think>.*?</think>", "", raw, flags=re.DOTALL).strip()
    cleaned = re.sub(r"</?(no_)?think>", "", cleaned).strip()
    if cleaned.startswith("```"):
        match = re.search(r"```(?:json)?\s*\n?(.*?)\n?```", cleaned, re.DOTALL)
        if match:
            cleaned = match.group(1).strip()
    if not cleaned.startswith("{") and not cleaned.startswith("["):
        match = re.search(r"(\[.*\]|\{.*\})", cleaned, re.DOTALL)
        if match:
            cleaned = match.group(1).strip()
    return cleaned


def query_ollama(model: str, prompt: str, temperature: float = 0.1,
                 max_retries: int = 5, timeout: int = 450) -> str | None:
    effective_prompt = prompt
    if "qwen3" in model.lower():
        effective_prompt = "/no_think\n\n" + prompt

    # Adjust context size per model architecture to prevent CUDA VRAM memory bounds
    num_ctx = 4096 if "gemma" in model.lower() else 8192
    num_predict = 2048 if "gemma" in model.lower() else 4096

    for attempt in range(1, max_retries + 1):
        temp = temperature if attempt == 1 else 0.0
        payload = {
            "model": model,
            "prompt": effective_prompt,
            "format": "json",
            "stream": False,
            "options": {
                "temperature": temp,
                "num_predict": num_predict,
                "num_ctx": num_ctx,
            },
        }
        try:
            resp = requests.post(OLLAMA_URL, json=payload, timeout=timeout)
            if resp.status_code == 200:
                raw = resp.json().get("response", "")
                cleaned = clean_model_response(raw)
                if not cleaned or len(cleaned.strip()) < 10:
                    if attempt < max_retries:
                        time.sleep(3 * attempt)
                        continue
                return cleaned
            logging.warning(f"Ollama HTTP {resp.status_code} (attempt {attempt}, model={model})")
        except requests.exceptions.Timeout:
            logging.warning(f"Ollama timeout (attempt {attempt}, model={model})")
        except Exception as e:
            logging.warning(f"Ollama error (attempt {attempt}, model={model}): {e}")
        if attempt < max_retries:
            time.sleep(3 * attempt)
    return None


def query_ollama_text(model: str, prompt: str, temperature: float = 0.3,
                      max_retries: int = 3, timeout: int = 120) -> str | None:
    """Query Ollama for free-text output (no JSON format constraint)."""
    effective_prompt = prompt
    if "qwen3" in model.lower():
        effective_prompt = "/no_think\n\n" + prompt

    for attempt in range(1, max_retries + 1):
        temp = temperature if attempt == 1 else 0.1
        payload = {
            "model": model,
            "prompt": effective_prompt,
            "stream": False,
            "options": {
                "temperature": temp,
                "num_predict": 2048,
                "num_ctx": 4096,
            },
        }
        try:
            resp = requests.post(OLLAMA_URL, json=payload, timeout=timeout)
            if resp.status_code == 200:
                raw = resp.json().get("response", "")
                cleaned = clean_model_response(raw)
                if not cleaned or len(cleaned.strip()) < 20:
                    if attempt < max_retries:
                        time.sleep(3 * attempt)
                        continue
                return cleaned
            logging.warning(f"Ollama HTTP {resp.status_code} (attempt {attempt}, model={model})")
        except requests.exceptions.Timeout:
            logging.warning(f"Ollama timeout (attempt {attempt}, model={model})")
        except Exception as e:
            logging.warning(f"Ollama error (attempt {attempt}, model={model}): {e}")
        if attempt < max_retries:
            time.sleep(3 * attempt)
    return None


def build_translation_prompt(final_explanation: str) -> str:
    """Build the Qwen3 interactive teacher persona prompt for Vietnamese translation."""
    return (
        'You are a warm, engaging, and expert English teacher having a 1-on-1 '
        'interactive conversation with a Vietnamese student studying PTE '
        'Reading Fill-in-the-Blanks.\n\n'
        'GOAL: Explain the answer in friendly, natural, human Vietnamese, '
        'making it feel like a real conversation.\n\n'
        'RULES:\n'
        '1. Use an interactive, human teacher persona (calling the student '
        '"em" and asking friendly reflective questions like '
        '"Em có chú ý... không?").\n'
        '2. DO NOT repeat rigid repetitive headings like "Why \'X\' is '
        'incorrect:". PARAPHRASE and translate distractor questions naturally '
        'into conversational Vietnamese (e.g., "Thế còn phương án \'was '
        'receiving\' thì sao nhỉ?", "Tại sao \'had received\' lại chưa '
        'chuẩn trong câu này?").\n'
        '3. KEEP exact option choices (e.g., \'received\', \'was receiving\'), '
        'grammar terms (e.g., \'Past Simple Tense\', \'Past Continuous\'), '
        'and English quotes in ENGLISH.\n'
        '4. End with a "Tóm lại..." conclusion that includes a merged '
        'reflective takeaway or thought-provoking question for the student. '
        'Do NOT write a separate "Reflective Check" section.\n\n'
        'English Explanation:\n'
        f'{final_explanation}\n\n'
        'Interactive Vietnamese Teacher Explanation:'
    )


def translate_blanks(record: dict, no_resume: bool = False) -> tuple[dict, list[int], list[int]]:
    """Add vi_explanation, vi_raw, vi_model, vi_timestamp to each blank in the record using Qwen3."""
    qid = record.get("id")
    blanks = record.get("blanks", [])
    translate_model = MODELS["qw"]
    ok_indexes = []
    failed_indexes = []

    for i, blank in enumerate(blanks):
        final_exp = blank.get("final_explanation", "")
        b_idx = blank.get("blank_index", i + 1)
        if not final_exp:
            continue

        if not no_resume:
            has_exp = bool(blank.get("vi_explanation"))
            has_raw = bool(blank.get("vi_raw"))
            has_model = bool(blank.get("vi_model"))
            has_ts = bool(blank.get("vi_timestamp"))
            if has_exp and has_raw and has_model and has_ts:
                continue

        logging.info(f"  Translating blank {b_idx} for Question {qid} ...")

        prompt = build_translation_prompt(final_exp)
        raw_text = query_ollama_text(translate_model, prompt, temperature=0.3)

        if raw_text:
            cleaned_text = clean_model_response(raw_text)
            blank["vi_explanation"] = cleaned_text
            blank["vi_raw"] = raw_text
            blank["vi_model"] = translate_model
            blank["vi_timestamp"] = datetime.now(timezone.utc).isoformat()
            ok_indexes.append(b_idx)
            logging.info(f"  Translated blank {b_idx} OK ({len(cleaned_text)} chars)")
        else:
            failed_indexes.append(b_idx)
            logging.warning(f"  Translation FAILED for blank {b_idx} of Question {qid}")

    return record, ok_indexes, failed_indexes


def compute_objective_confidence(
    blank_index: int,
    expected_correct: str,
    all_options: list[str],
    revised_blank: dict,
) -> tuple[str, list[str]]:
    flags = []
    correct_ans = revised_blank.get("correct_answer", "").strip().lower()
    if correct_ans != expected_correct.strip().lower():
        flags.append(f"Correct answer mismatch: '{correct_ans}' != '{expected_correct}'")

    explanation = revised_blank.get("final_explanation", "").strip()
    if not explanation:
        flags.append("Final explanation is empty")

    if len(explanation) < 100:
        flags.append(f"Explanation too short ({len(explanation)} chars, min 100)")
    elif len(explanation) > 1000:
        flags.append(f"Explanation too long ({len(explanation)} chars, max 1000)")

    distractors = [opt.strip().lower() for opt in all_options[1:]]
    for d in distractors:
        if d not in explanation.lower():
            flags.append(f"Distractor '{d}' not explicitly mentioned in explanation")

    tag = revised_blank.get("grammar_tag", "").strip()
    if not tag or tag.lower() in ("grammar", "general", "other", "n/a"):
        flags.append(f"Generic or empty grammar tag: '{tag}'")

    if not flags:
        return "high", []
    elif len(flags) <= 2 and "Final explanation is empty" not in flags:
        return "medium", flags
    else:
        return "low", flags


def build_phase1_prompt(
    answer_with_blanks: str,
    full_text: str,
    blanks: list[dict],
    existing_explanation: dict | None,
) -> str:
    blanks_block = "\n".join(
        f"  Blank {b['index']}: correct=\"{b['correct']}\"  options={b['options']}"
        for b in blanks
    )

    existing_block = ""
    if existing_explanation:
        exps = existing_explanation.get("explanations", [])
        for exp in exps:
            idx = exp.get("blank_index", "?")
            sel = exp.get("selection_reason", "N/A")
            if isinstance(sel, dict):
                sel = sel.get("reason") or sel.get("explanation") or str(sel)
            std = exp.get("detailed_student_explanation", "N/A")
            if isinstance(std, dict):
                std = std.get("explanation") or std.get("student_explanation") or str(std)
            existing_block += f"\n  Blank {idx}:\n"
            existing_block += f"    selection_reason: {str(sel)[:200]}\n"
            existing_block += f"    student_explanation: {str(std)[:200]}\n"

    return f"""You are an expert linguist and ESL assessment specialist. Your role is DIAGNOSTIC REASONING - explain WHY things are the way they are.

TASK: Deeply analyze each blank in this Reading Fill-in-the-Blanks (RFIB) passage. For each blank:

1. Explain precisely WHY the correct answer is the only correct choice. Be specific about the grammar rule, collocation pattern, or semantic constraint.
2. For EACH distractor (wrong option), explain specifically WHY it is wrong. Cite the exact grammar rule violation, semantic mismatch, or collocational impossibility.
3. Evaluate the existing AI-generated explanation (provided below). Identify any errors, imprecisions, missing reasoning, or misleading claims in it.

DATA:
1. Passage with blanks (__option1/option2/option3/option4__):
{answer_with_blanks}

2. Full solved text:
{full_text}

3. Detected blanks:
{blanks_block}

4. Existing AI explanation to evaluate:
{existing_block if existing_block else "(No existing explanation available)"}

RESPONSE FORMAT - return ONLY valid JSON, no markdown fences:
{{
  "diagnostic_blanks": [
    {{
      "blank_index": 1,
      "correct_answer": "word",
      "correct_reason": "Detailed grammar/collocation/semantic reasoning for why this is correct.",
      "distractor_analysis": [
        {{
          "option": "wrong_word",
          "why_wrong": "Specific reason this option fails - cite grammar rule, semantic clash, etc."
        }}
      ],
      "existing_errors": "What the existing AI explanation got wrong or missed. Empty string if the existing explanation is accurate.",
      "cohesion_tie": "Genuine lexical/grammatical cohesion link to other text, or empty string if none."
    }}
  ]
}}

RULES:
- Return exactly {len(blanks)} objects in diagnostic_blanks, one per blank, in order.
- blank_index must be sequential starting from 1.
- correct_answer must match the first option in each blank exactly.
- Each distractor_analysis must cover ALL distractors (every option except the correct one).
- Be rigorous and specific. Avoid vague claims like "it doesn't fit" - explain exactly WHY.
- MULTI-DIMENSION RULE: For each distractor, evaluate ALL applicable dimensions independently:
  (a) TENSE — Does the tense match the time frame of the passage?
  (b) VOICE — Is Active/Passive voice appropriate for the subject-verb relationship? If a distractor uses passive voice, does it make semantic sense for the subject to be the receiver of the action?
  (c) ASPECT — Is the aspect (simple/continuous/perfect) appropriate for the event type?
  (d) COLLOCATION — Does the word collocate naturally with its surrounding words?
  (e) MEANING — Does the word fit the semantic context?
  List ALL dimensions that apply in why_wrong, not just the first one you find. Separate multiple reasons with numbered points if more than one dimension fails.
"""


def build_phase2_prompt(
    answer_with_blanks: str,
    full_text: str,
    blanks: list[dict],
    phase1_output: dict,
) -> str:
    dr_block = json.dumps(phase1_output, indent=2, ensure_ascii=False)

    return f"""You are an expert ESL content editor. Your role is STRUCTURED SYNTHESIS - take raw diagnostic analysis and transform it into clear, well-organized, student-facing explanations.

TASK: Transform the diagnostic reasoning below into polished student-facing explanations for B1-B2 English learners. For each blank:

1. Write a clear, structured explanation with this flow:
   a. Start with WHY the correct answer fits (1-2 sentences, clear grammar/meaning reason)
   b. Then explain each wrong option with a bullet point, using this format:
      * Why '[option]' is incorrect: [specific reason]
   c. If there's a cohesion tie, add it as a final note

2. Use simple, direct English appropriate for B1-B2 learners
3. Assign a specific grammar_tag (e.g., "Past Simple Tense", "Relative Pronoun", "Collocation with Preposition")
4. Flag any gaps in the diagnostic reasoning that seem incomplete

ORIGINAL PASSAGE:
{answer_with_blanks}

FULL TEXT:
{full_text}

DIAGNOSTIC REASONING (from Phase 1):
{dr_block}

RESPONSE FORMAT - return ONLY valid JSON, no markdown fences:
{{
  "synthesized_blanks": [
    {{
      "blank_index": 1,
      "correct_answer": "word",
      "grammar_tag": "Grammar Category",
      "student_explanation": "Full student-facing explanation combining correct-answer reasoning and distractor analysis. Written for B1-B2 learners.",
      "concise_explanation": "One-sentence summary of why this answer is correct (max 50 words).",
      "gaps_flagged": "Any gaps or issues in the diagnostic reasoning. Empty string if none."
    }}
  ]
}}

RULES:
- Return exactly {len(blanks)} objects in synthesized_blanks, one per blank, in order.
- blank_index must be sequential starting from 1.
- correct_answer must match the first option in each blank exactly.
- student_explanation MUST address the correct answer AND every distractor using * bullet points.
- When a distractor fails on MULTIPLE dimensions (e.g., both wrong tense AND wrong voice), list each reason as a separate numbered point within that distractor's bullet. Do NOT collapse multiple failures into a single vague sentence.
- When labelling a grammar form (e.g., "Present Passive", "Past Continuous"), ALWAYS explain what that label means in plain English so B1-B2 learners understand it.
"""


def build_phase3_prompt(
    answer_with_blanks: str,
    full_text: str,
    blanks: list[dict],
    phase2_output: dict,
    existing_explanation: dict | None,
) -> str:
    qw_block = json.dumps(phase2_output, indent=2, ensure_ascii=False)

    existing_block = ""
    if existing_explanation:
        exps = existing_explanation.get("explanations", [])
        for exp in exps:
            idx = exp.get("blank_index", "?")
            existing_block += f"\n  Blank {idx}: {str(exp.get('detailed_student_explanation', 'N/A'))[:300]}\n"

    return f"""You are an expert English language reviewer and quality auditor. Your role is CROSS-REVIEW - verify accuracy, naturalness, and correctness of explanations written by other AI models.

TASK: Review the synthesized student-facing explanations below. For each blank:

1. VERIFY: Are all grammar rules cited factually correct? Is the reasoning sound?
2. CHECK ENGLISH: Is the language natural and appropriate for B1-B2 learners? Fix awkward phrasing.
3. COMPARE: Compare with the original explanation (if provided). Flag any important disagreements.
4. COMPLETENESS CHECK: For each distractor, verify that ALL relevant failure dimensions are covered. If a distractor fails on Tense but ALSO fails on Voice/Transitivity (or vice versa), add the missing dimension. Common missed dimensions:
   - Passive voice used where Active is semantically required (subject performs the action)
   - Active voice used where Passive is semantically required (subject receives the action)
   - Aspect mismatch (simple vs continuous vs perfect) beyond just tense
5. FINALIZE: Produce the definitive final explanation, incorporating any corrections and completeness additions.

ORIGINAL PASSAGE:
{answer_with_blanks}

FULL TEXT:
{full_text}

SYNTHESIZED EXPLANATIONS (from Phase 2):
{qw_block}

ORIGINAL AI EXPLANATION (for comparison):
{existing_block if existing_block else "(No original explanation available)"}

RESPONSE FORMAT - return ONLY valid JSON, no markdown fences:
{{
  "reviewed_blanks": [
    {{
      "blank_index": 1,
      "correct_answer": "word",
      "grammar_tag": "Grammar Category (verified or corrected)",
      "final_explanation": "The definitive student-facing explanation after review and corrections.",
      "concise_explanation": "One-sentence summary (max 50 words), verified for accuracy.",
      "review_notes": "What was verified, corrected, or flagged. Empty string if no changes needed.",
      "disagreements": "Any disagreements with original AI explanation. Empty string if aligned."
    }}
  ]
}}

RULES:
- Return exactly {len(blanks)} objects in reviewed_blanks, one per blank, in order.
- blank_index must be sequential starting from 1.
- correct_answer must match the first option in each blank exactly.
- final_explanation must be non-empty and address correct answer AND every distractor.
"""


def _find_list_key(data: dict, preferred: str, fallbacks: list[str]) -> list | None:
    if not isinstance(data, dict):
        return None
    result = data.get(preferred)
    if isinstance(result, list) and len(result) > 0:
        return result
    for key in fallbacks:
        result = data.get(key)
        if isinstance(result, list) and len(result) > 0:
            return result
    for key, val in data.items():
        if isinstance(val, list) and len(val) > 0 and isinstance(val[0], dict):
            return val
        if isinstance(val, dict):
            sub_res = _find_list_key(val, preferred, fallbacks)
            if sub_res:
                return sub_res
    return None


def parse_phase1_response(raw: str, blanks: list[dict]) -> dict | None:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return None

    if isinstance(data, list):
        diagnostics = data
        data = {"diagnostic_blanks": diagnostics}
    elif isinstance(data, dict):
        diagnostics = _find_list_key(data, "diagnostic_blanks", ["diagnostics", "blanks", "analysis", "items", "questions"])
    else:
        return None

    if not isinstance(diagnostics, list) or len(diagnostics) < len(blanks):
        return None

    if len(diagnostics) > len(blanks):
        diagnostics = diagnostics[:len(blanks)]

    data["diagnostic_blanks"] = diagnostics
    for i, diag in enumerate(diagnostics):
        if isinstance(diag, str):
            diag = {"correct_reason": diag}
            diagnostics[i] = diag
        elif not isinstance(diag, dict):
            diag = {"correct_reason": str(diag)}
            diagnostics[i] = diag
        expected_idx = i + 1
        diag["blank_index"] = expected_idx
        diag["correct_answer"] = blanks[i]["correct"]
        if not diag.get("correct_reason", "").strip():
            for alt in ["reason", "explanation", "why_correct"]:
                if diag.get(alt, "").strip():
                    diag["correct_reason"] = diag[alt]
                    break
    return data


def parse_phase2_response(raw: str, blanks: list[dict]) -> dict | None:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return None

    if isinstance(data, list):
        synthesized = data
        data = {"synthesized_blanks": synthesized}
    elif isinstance(data, dict):
        synthesized = _find_list_key(data, "synthesized_blanks", ["blanks", "synthesized", "explanations", "items", "questions"])
    else:
        return None

    if not isinstance(synthesized, list) or len(synthesized) < len(blanks):
        return None

    if len(synthesized) > len(blanks):
        synthesized = synthesized[:len(blanks)]

    data["synthesized_blanks"] = synthesized
    for i, synth in enumerate(synthesized):
        if isinstance(synth, str):
            synth = {"student_explanation": synth}
            synthesized[i] = synth
        elif not isinstance(synth, dict):
            synth = {"student_explanation": str(synth)}
            synthesized[i] = synth
        synth["blank_index"] = i + 1
        if not synth.get("student_explanation", "").strip():
            for alt in ["explanation", "detailed_explanation"]:
                if synth.get(alt, "").strip():
                    synth["student_explanation"] = synth[alt]
                    break
    return data


def parse_phase3_response(raw: str, blanks: list[dict]) -> dict | None:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return None

    if isinstance(data, list):
        reviewed = data
        data = {"reviewed_blanks": reviewed}
    elif isinstance(data, dict):
        reviewed = _find_list_key(data, "reviewed_blanks", ["blanks", "reviewed", "review", "items", "questions"])
    else:
        return None

    if not isinstance(reviewed, list) or len(reviewed) < len(blanks):
        return None

    if len(reviewed) > len(blanks):
        reviewed = reviewed[:len(blanks)]

    data["reviewed_blanks"] = reviewed
    for i, rev in enumerate(reviewed):
        if isinstance(rev, str):
            rev = {"final_explanation": rev}
            reviewed[i] = rev
        elif not isinstance(rev, dict):
            rev = {"final_explanation": str(rev)}
            reviewed[i] = rev
        rev["blank_index"] = i + 1
        if not rev.get("final_explanation", "").strip():
            for alt in ["explanation", "student_explanation"]:
                if rev.get(alt, "").strip():
                    rev["final_explanation"] = rev[alt]
                    break
    return data


def load_sidecar(path: str) -> dict[int, dict]:
    records = {}
    if not os.path.exists(path):
        return records
    with open(path, "r", encoding="utf-8") as f:
        for line_num, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                rid = int(rec.get("id", 0))
                if rid > 0:
                    records[rid] = rec
            except (json.JSONDecodeError, ValueError):
                continue
    return records


def save_sidecar_atomic(path: str, records: dict[int, dict]):
    tmp_path = path + ".tmp"
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            for rid in sorted(records.keys()):
                f.write(json.dumps(records[rid], ensure_ascii=False) + "\n")
        
        for attempt in range(1, 6):
            try:
                os.replace(tmp_path, path)
                break
            except PermissionError:
                if attempt == 5:
                    raise
                time.sleep(0.5 * attempt)
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        raise


def load_workbook_data(workbook_path: str) -> dict[int, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=False)
    header_row = next(rows_iter)
    hmap = {}
    for i, cell in enumerate(header_row):
        if cell.value is not None:
            hmap[str(cell.value).strip()] = i

    id_idx = hmap.get("ID", hmap.get("id"))
    answer_idx = hmap.get("ANSWER", hmap.get("Answer"))
    full_text_idx = hmap.get("Full Text", hmap.get("fullText"))

    if id_idx is None or answer_idx is None or full_text_idx is None:
        raise ValueError(f"Missing required columns in workbook. Headers: {list(hmap.keys())}")

    questions = {}
    for row in rows_iter:
        cells = list(row)
        if id_idx >= len(cells):
            continue
        raw_id = cells[id_idx].value
        if raw_id is None:
            continue
        try:
            qid = int(raw_id)
        except (ValueError, TypeError):
            continue

        answer_text = str(cells[answer_idx].value or "") if answer_idx < len(cells) else ""
        full_text = str(cells[full_text_idx].value or "") if full_text_idx < len(cells) else ""

        if not answer_text.strip() or not full_text.strip():
            continue

        blanks = extract_blanks(answer_text)
        if not blanks:
            continue

        questions[qid] = {
            "id": qid,
            "answer_text": answer_text,
            "full_text": full_text,
            "blanks": blanks,
        }

    wb.close()
    return questions


def load_existing_explanations(jsonl_path: str) -> dict[int, dict]:
    records = {}
    if not os.path.exists(jsonl_path):
        return records
    with open(jsonl_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                rid = int(rec.get("id", 0))
                if rid > 0:
                    records[rid] = rec
            except (json.JSONDecodeError, ValueError):
                continue
    return records


def process_question(
    qid: int,
    question: dict,
    existing_explanation: dict | None,
) -> dict | None:
    blanks = question["blanks"]
    answer_text = question["answer_text"]
    full_text = question["full_text"]

    logging.info(f"--- Question {qid}: {len(blanks)} blanks ---")

    logging.info("  Phase 1 [DR] deepseek-r1:14b ...")
    p1_prompt = build_phase1_prompt(answer_text, full_text, blanks, existing_explanation)
    t1 = time.time()
    p1_raw = query_ollama(MODELS["dr"], p1_prompt, temperature=0.1, timeout=300, max_retries=4)
    p1_data = parse_phase1_response(p1_raw, blanks) if p1_raw else None

    if p1_data is None and existing_explanation is not None:
        logging.info("  Phase 1 retry without existing explanation context...")
        p1_prompt = build_phase1_prompt(answer_text, full_text, blanks, None)
        p1_raw = query_ollama(MODELS["dr"], p1_prompt, temperature=0.1, timeout=300, max_retries=4)
        p1_data = parse_phase1_response(p1_raw, blanks) if p1_raw else None

    t1_elapsed = time.time() - t1

    if p1_data is None:
        logging.error(f"  Phase 1 FAILED for Question {qid}: No valid response from deepseek-r1:14b")
        return None
    logging.info(f"  Phase 1 OK ({t1_elapsed:.1f}s)")

    logging.info("  Phase 2 [QW] qwen3:14b ...")
    p2_prompt = build_phase2_prompt(answer_text, full_text, blanks, p1_data)
    t2 = time.time()
    p2_raw = query_ollama(MODELS["qw"], p2_prompt, temperature=0.2, timeout=300, max_retries=4)
    p2_data = parse_phase2_response(p2_raw, blanks) if p2_raw else None
    t2_elapsed = time.time() - t2

    if p2_data is None:
        logging.error(f"  Phase 2 FAILED for Question {qid}: No valid response from qwen3:14b")
        return None
    logging.info(f"  Phase 2 OK ({t2_elapsed:.1f}s)")

    logging.info("  Phase 3 [GM] gemma4:latest ...")
    p3_prompt = build_phase3_prompt(answer_text, full_text, blanks, p2_data, existing_explanation)
    t3 = time.time()
    p3_raw = query_ollama(MODELS["gm"], p3_prompt, temperature=0.1, timeout=300, max_retries=4)
    p3_data = parse_phase3_response(p3_raw, blanks) if p3_raw else None
    t3_elapsed = time.time() - t3

    if p3_data is None:
        logging.error(f"  Phase 3 FAILED for Question {qid}: No valid response from gemma4:latest")
        return None
    logging.info(f"  Phase 3 OK ({t3_elapsed:.1f}s)")

    total_time = t1_elapsed + t2_elapsed + t3_elapsed

    reviewed_blanks = p3_data["reviewed_blanks"]
    diagnostic_blanks = p1_data["diagnostic_blanks"]
    synthesized_blanks = p2_data["synthesized_blanks"]

    revision_blanks = []
    overall_confidence_scores = []

    for i, rev in enumerate(reviewed_blanks):
        diag = diagnostic_blanks[i] if i < len(diagnostic_blanks) else {}
        synth = synthesized_blanks[i] if i < len(synthesized_blanks) else {}
        b_info = blanks[i]

        conf, flags = compute_objective_confidence(
            blank_index=i + 1,
            expected_correct=b_info["correct"],
            all_options=b_info["options"],
            revised_blank=rev,
        )
        overall_confidence_scores.append(conf)

        revision_blanks.append({
            "blank_index": rev.get("blank_index", i + 1),
            "correct_answer": rev.get("correct_answer", b_info["correct"]),
            "grammar_tag": rev.get("grammar_tag", synth.get("grammar_tag", "Grammar")),
            "dr_reasoning": diag.get("correct_reason", ""),
            "dr_distractor_analysis": diag.get("distractor_analysis", []),
            "dr_existing_errors": diag.get("existing_errors", ""),
            "qw_formatted": synth.get("student_explanation", ""),
            "qw_gaps_flagged": synth.get("gaps_flagged", ""),
            "gm_notes": rev.get("review_notes", ""),
            "gm_disagreements": rev.get("disagreements", ""),
            "final_explanation": rev.get("final_explanation", ""),
            "concise_explanation": rev.get("concise_explanation", ""),
            "confidence": conf,
            "confidence_flags": flags,
            "cohesion_tie": diag.get("cohesion_tie", ""),
        })

    errors_found = sum(1 for b in revision_blanks if b["dr_existing_errors"].strip())
    summary = f"Processed {len(blanks)} blanks cleanly."
    if errors_found:
        summary += f" DR identified issues in {errors_found} existing Gemma explanations."

    record = {
        "id": qid,
        "original_status": existing_explanation.get("status", "") if existing_explanation else "New",
        "revision_summary": summary,
        "blanks": revision_blanks,
        "phase_models": [MODELS["dr"], MODELS["qw"], MODELS["gm"]],
        "phase_times": {
            "dr_seconds": round(t1_elapsed, 1),
            "qw_seconds": round(t2_elapsed, 1),
            "gm_seconds": round(t3_elapsed, 1),
            "total_seconds": round(total_time, 1),
        },
        "raw_artifacts": {
            "dr_raw": p1_raw,
            "qw_raw": p2_raw,
            "gm_raw": p3_raw,
        },
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }

    logging.info(f"  [OK] Question {qid} complete in {total_time:.1f}s - {summary}")
    return record


def parse_args():
    parser = argparse.ArgumentParser(
        description="3-Model RFIB Explanation Revision Pipeline (DR -> QW -> GM)"
    )
    parser.add_argument("--input", default=INPUT_WORKBOOK)
    parser.add_argument("--existing", default=INPUT_JSONL)
    parser.add_argument("--out", default=OUTPUT_JSONL)
    parser.add_argument("--test", type=int, help="Use N curated test IDs")
    parser.add_argument("--ids", type=str, help="Comma-separated IDs")
    parser.add_argument("--start", type=int, help="Zero-based start index")
    parser.add_argument("--limit", type=int, help="Max questions to process")
    parser.add_argument("--no-resume", action="store_true")
    parser.add_argument("--reset-sidecar", action="store_true")
    parser.add_argument("--save-every", type=int, default=10)
    parser.add_argument("--translate", action="store_true",
                        help="Enrichment-only mode: add Vietnamese translations to existing sidecar records")
    return parser.parse_args()


def main():
    args = parse_args()

    if args.reset_sidecar and os.path.exists(args.out):
        os.remove(args.out)

    sidecar_records = load_sidecar(args.out)

    if args.translate:
        logging.info("=== Phase 4: Vietnamese Translation Enrichment ===")
        if not os.path.exists(args.out) or not sidecar_records:
            logging.error(f"Sidecar file '{args.out}' does not exist or is empty. Cannot run translation mode.")
            sys.exit(2)

        if args.ids is not None:
            id_list = [int(x.strip()) for x in args.ids.split(",") if x.strip()]
            missing_ids = [qid for qid in id_list if qid not in sidecar_records]
            if missing_ids:
                logging.error(f"Requested translation IDs missing from sidecar '{args.out}': {missing_ids}")
                sys.exit(2)
            target_ids = id_list
        elif args.test is not None:
            target_ids = [qid for qid in CURATED_QUALITY_GATE_IDS[:args.test] if qid in sidecar_records]
        elif args.start is not None:
            all_sidecar_ids = sorted(sidecar_records.keys())
            end = args.start + (args.limit or len(all_sidecar_ids))
            target_ids = all_sidecar_ids[args.start:end]
        else:
            all_sidecar_ids = sorted(sidecar_records.keys())
            target_ids = all_sidecar_ids[:(args.limit or len(all_sidecar_ids))]

        translate_count = 0
        failed_records = []
        failed_blanks_details = []

        for qid in target_ids:
            rec = sidecar_records[qid]
            logging.info(f"--- Translating Question {qid}: {len(rec.get('blanks', []))} blanks ---")

            rec, ok_idx, fail_idx = translate_blanks(rec, no_resume=args.no_resume)
            sidecar_records[qid] = rec

            if fail_idx:
                failed_records.append(qid)
                failed_blanks_details.append((qid, fail_idx))
                logging.error(f"Question {qid} failed translation on blank indexes: {fail_idx}")

            if ok_idx:
                translate_count += len(ok_idx)

            if (target_ids.index(qid) + 1) % max(1, args.save_every) == 0:
                save_sidecar_atomic(args.out, sidecar_records)

        save_sidecar_atomic(args.out, sidecar_records)
        logging.info(f"Translation phase finished. Blanks translated: {translate_count}, Failed records: {len(failed_records)}")

        if failed_records:
            failed_str = ",".join(map(str, failed_records))
            logging.error(f"Translation batch failed for {len(failed_records)} records: {failed_blanks_details}")
            logging.info(f"Rerun command to retry failed records:\n  python scripts/revise_rfib_explanations.py --translate --out \"{args.out}\" --ids {failed_str}")
            sys.exit(1)
        sys.exit(0)

    # --- Phase 1–3 Execution ---
    questions = load_workbook_data(args.input)
    existing = load_existing_explanations(args.existing)
    all_ids = sorted(questions.keys())

    if args.test is not None:
        selected_ids = CURATED_QUALITY_GATE_IDS[:args.test]
        selected_ids = [qid for qid in selected_ids if qid in questions]
    elif args.ids is not None:
        selected_ids = [int(x.strip()) for x in args.ids.split(",") if x.strip()]
    elif args.start is not None:
        end = args.start + (args.limit or len(all_ids))
        selected_ids = all_ids[args.start:end]
    else:
        selected_ids = all_ids[: (args.limit or len(all_ids))]

    succeeded = 0
    failed_ids = []

    for qid in selected_ids:
        if not args.no_resume and qid in sidecar_records:
            continue

        q_data = questions.get(qid)
        if not q_data:
            continue

        e_data = existing.get(qid)
        res = process_question(qid, q_data, e_data)
        if res:
            sidecar_records[qid] = res
            succeeded += 1
        else:
            failed_ids.append(qid)

        if succeeded % args.save_every == 0:
            save_sidecar_atomic(args.out, sidecar_records)

    save_sidecar_atomic(args.out, sidecar_records)

    logging.info(f"Done. Succeeded: {succeeded}, Failed: {len(failed_ids)}")
    if failed_ids:
        logging.error(f"Batch completed with {len(failed_ids)} failed IDs: {failed_ids}")
        sys.exit(1)


if __name__ == "__main__":
    main()