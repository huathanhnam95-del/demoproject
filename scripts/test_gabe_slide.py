import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_test_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Editorial Palette
    C_BG = RGBColor(250, 248, 245)
    C_CARD = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 221, 214)
    C_DARK = RGBColor(26, 29, 32)
    C_MUTED = RGBColor(95, 100, 110)
    C_TERRA = RGBColor(168, 67, 43)
    C_TERRA_LIGHT = RGBColor(253, 242, 240)
    C_SLATE = RGBColor(30, 58, 95)
    C_SLATE_LIGHT = RGBColor(238, 242, 246)
    C_GREEN = RGBColor(22, 101, 52)
    C_GREEN_LIGHT = RGBColor(236, 253, 245)
    C_BANNER_BG = RGBColor(243, 239, 233)

    plate_gabe = r"C:\Cursor AI\assets\speaker_portraits\plates\gabe_newell_plate.png"

    slide = prs.slides.add_slide(blank_layout)

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()

    # Header
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.25))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = "FOUNDATIONAL PHILOSOPHY • THE ECONOMIC & CREATIVE LOGIC"
    p.font.name = "Segoe UI"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TERRA
    p.space_after = Pt(2)

    p = tf.add_paragraph()
    p.text = "Gabe Newell's Ideology: Why Hierarchies Destroy Creative Value"
    p.font.name = "Georgia"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p.space_after = Pt(3)

    p = tf.add_paragraph()
    p.text = "From Ronald Coase's transaction costs to human leverage: why Valve eliminated bosses to unleash potential"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.color.rgb = C_MUTED

    # Left Column: Founder Profile & Hero Quote Card
    left_x = Inches(0.8)
    left_w = Inches(3.60)
    left_y = Inches(1.72)
    left_h = Inches(4.35)

    c_founder = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, left_y, left_w, left_h)
    c_founder.fill.solid()
    c_founder.fill.fore_color.rgb = C_CARD
    c_founder.line.color.rgb = C_TERRA
    c_founder.line.width = Pt(1.5)

    # Pill at top of left card
    pill_w = left_w - Inches(0.36)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.18), left_y + Inches(0.14), pill_w, Inches(0.28))
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_TERRA_LIGHT
    pill.line.color.rgb = C_TERRA
    pill.line.width = Pt(0.75)
    ptf = pill.text_frame
    ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
    p = ptf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "✦ THE FOUNDER'S VISION & LOGIC"
    p.font.name = "Segoe UI"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = C_TERRA

    # Portrait Plate (4:5 ratio)
    pw = Inches(1.76)
    ph = Inches(2.20)
    px = left_x + (left_w - pw) / 2
    py = left_y + Inches(0.48)
    if os.path.exists(plate_gabe):
        slide.shapes.add_picture(plate_gabe, px, py, pw, ph)
        pfr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, py, pw, ph)
        pfr.fill.background()
        pfr.line.color.rgb = C_CARD_BORDER
        pfr.line.width = Pt(1)

    # Name & Subtitle
    tb_name = slide.shapes.add_textbox(left_x + Inches(0.18), py + ph + Inches(0.04), left_w - Inches(0.36), Inches(0.42))
    ntf = tb_name.text_frame
    ntf.word_wrap = True
    ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
    p = ntf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Gabe Newell"
    p.font.name = "Georgia"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p.space_after = Pt(1)

    p = ntf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Co-Founder & CEO • Valve Corporation"
    p.font.name = "Segoe UI Semibold"
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TERRA

    # Hero Quote Box at bottom of left column
    qb_y = py + ph + Inches(0.50)
    qb_h = left_y + left_h - qb_y - Inches(0.14)
    qb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.16), qb_y, left_w - Inches(0.32), qb_h)
    qb.fill.solid()
    qb.fill.fore_color.rgb = C_BANNER_BG
    qb.line.color.rgb = C_CARD_BORDER
    qb.line.width = Pt(0.75)
    qtf = qb.text_frame
    qtf.word_wrap = True
    qtf.margin_left = qtf.margin_right = Inches(0.14)
    qtf.margin_top = Inches(0.07)
    qtf.margin_bottom = Inches(0.05)

    p = qtf.paragraphs[0]
    p.text = "“If you hire smart, creative people and then tell them what to do, you destroy most of their value. You want people who figure out what is valuable and just go do it.”"
    p.font.name = "Georgia"
    p.font.size = Pt(8.25)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p.line_spacing = 1.15
    p.space_after = Pt(2)

    p = qtf.add_paragraph()
    p.text = "— Gabe Newell (University of Texas & Washington Post)"
    p.font.name = "Segoe UI Semibold"
    p.font.size = Pt(7.5)
    p.font.color.rgb = C_TERRA

    # Right Column: 3 Strategic Cards
    right_x = left_x + left_w + Inches(0.24)
    right_w = Inches(11.733) - left_w - Inches(0.24)
    card_h = Inches(1.37)
    gap = Inches(0.12)

    cards_data = [
        (
            "01 • TRANSACTION COSTS",
            Inches(2.15),
            C_SLATE,
            C_SLATE_LIGHT,
            "Why Traditional Hierarchies Break in Creative Work",
            [
                ("The Industrial Legacy: ", "Ronald Coase showed firms exist to cut transaction costs. Hierarchies worked for factories, but in creative education, corporate managers actually introduce massive friction, red tape, and delay."),
                ("Zero Internal Friction: ", "By eliminating middle management, Valve operates as an agile internal network. Ideas move at the speed of direct conversation without waiting weeks for supervisor approvals.")
            ]
        ),
        (
            "02 • CREATIVE LEVERAGE",
            Inches(1.95),
            C_TERRA,
            C_TERRA_LIGHT,
            "Autonomy Multiplies Creative Leverage Over Direct Orders",
            [
                ("The Management Bottleneck: ", "When supervisors dictate daily tasks, educators can only produce what one manager imagines. Dictating tasks to masters caps their creative output at average execution."),
                ("Unlocking Full Leverage: ", "When educators own their curriculum directly, their impact multiplies. They spot subtle student confusions, craft engaging pronunciation games, and take personal pride in every outcome.")
            ]
        ),
        (
            "03 • LEARNER FIRST",
            Inches(1.65),
            C_GREEN,
            C_GREEN_LIGHT,
            "Eliminating Office Politics to Focus 100% on the Student",
            [
                ("Ending 'Managing Up': ", "In corporate ladders, people waste half their mental energy managing impressions and pleasing bosses. Without managers, internal politics and defensive turf wars evaporate."),
                ("One True North Star: ", "100% of energy is directed outward: “Does this genuinely help students learn?” When the learner is the only judge, quality remains uncompromising and transparent.")
            ]
        )
    ]

    for i, (badge_txt, pill_w_r, col, bg_col, title, bullets) in enumerate(cards_data):
        cy = left_y + i * (card_h + gap)
        cbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, cy, right_w, card_h)
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = C_CARD
        cbox.line.color.rgb = col
        cbox.line.width = Pt(1.5)

        # Header Pill
        pill_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x + Inches(0.16), cy + Inches(0.12), pill_w_r, Inches(0.26))
        pill_r.fill.solid()
        pill_r.fill.fore_color.rgb = bg_col
        pill_r.line.color.rgb = col
        pill_r.line.width = Pt(0.75)
        prtf = pill_r.text_frame
        prtf.margin_left = prtf.margin_right = prtf.margin_top = prtf.margin_bottom = 0
        p = prtf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = badge_txt
        p.font.name = "Segoe UI"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = col

        # Card Title next to pill - generous width
        title_x = right_x + Inches(0.16) + pill_w_r + Inches(0.14)
        title_w = right_w - Inches(0.30) - pill_w_r - Inches(0.14)
        tb_t = slide.shapes.add_textbox(title_x, cy + Inches(0.10), title_w, Inches(0.28))
        ttf = tb_t.text_frame
        ttf.word_wrap = True
        ttf.margin_left = ttf.margin_right = ttf.margin_top = ttf.margin_bottom = 0
        p = ttf.paragraphs[0]
        p.text = title
        p.font.name = "Georgia"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_DARK

        # Bullets below
        tb_b = slide.shapes.add_textbox(right_x + Inches(0.16), cy + Inches(0.44), right_w - Inches(0.32), card_h - Inches(0.48))
        btf = tb_b.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0

        for j, (b_lead, b_body) in enumerate(bullets):
            p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
            p.text = "• "
            p.font.name = "Segoe UI"
            p.font.size = Pt(9)
            p.font.color.rgb = col
            r1 = p.add_run()
            r1.text = b_lead
            r1.font.name = "Segoe UI Semibold"
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = C_DARK
            r2 = p.add_run()
            r2.text = b_body
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(9)
            r2.font.color.rgb = C_MUTED
            p.line_spacing = 1.14
            p.space_after = Pt(2.5)

    # Bottom Unifying Bridge Banner
    banner_y = Inches(6.20)
    ban_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), banner_y, Inches(11.733), Inches(0.90))
    ban_box.fill.solid()
    ban_box.fill.fore_color.rgb = C_BANNER_BG
    ban_box.line.color.rgb = C_CARD_BORDER
    ban_box.line.width = Pt(1)

    ban_tf = ban_box.text_frame
    ban_tf.word_wrap = True
    ban_tf.margin_left = ban_tf.margin_right = Inches(0.25)
    ban_tf.margin_top = Inches(0.08)

    p = ban_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "💡 HOW GABE'S FOUNDING LOGIC EMPOWERS OUR TEACHING TEAM"
    p.font.name = "Segoe UI"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TERRA
    p.space_after = Pt(2)

    p = ban_tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "“We aren't making software, but educational design is deeply creative knowledge work. You don't need an administrator's permission slip to make a lesson better. When we trust each other as equals, our students get the very best curriculum.”"
    p.font.name = "Georgia"
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = C_DARK

    test_out = r"C:\Cursor AI\test_gabe_slide.pptx"
    prs.save(test_out)

    # Export to PNG
    import win32com.client
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = ppt.Presentations.Open(os.path.abspath(test_out), WithWindow=False)
        out_img = os.path.abspath(r"C:\Cursor AI\exported_slides\test_gabe_slide.png")
        pres.Slides[0].Export(out_img, "PNG", 1920, 1080)
        pres.Close()
    finally:
        ppt.Quit()

if __name__ == "__main__":
    build_test_slide()
