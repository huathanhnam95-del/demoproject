"""
Automated Artifact Validator for RFIB 3-Model Revision Sidecar
Supports:
- Selective ID validation (--ids 1,2,9)
- Full workbook validation (--full)
- Vietnamese translation validation (--check-vi)
"""

import sys
import os
import json
import argparse
import re

DEFAULT_SIDECAR = r"public\database\RFIB\RFIB_3model_revision.jsonl"
DEFAULT_INPUT = r"public\database\RFIB\RFIB Final ver.xlsx"
DEFAULT_CURATED_IDS = [1, 2, 9, 13, 375, 398, 597, 677, 944, 1315]

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def load_workbook_data(workbook_path: str) -> dict[int, dict]:
    """Helper to load workbook questions using openpyxl."""
    import openpyxl

    if not os.path.exists(workbook_path):
        return {}

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=False)
    header_row = next(rows_iter)
    hmap = {}
    for i, cell in enumerate(header_row):
        if cell.value is not None:
            hmap[str(cell.value).strip()] = i

    id_idx = hmap.get("ID", hmap.get("id"))
    answer_idx = hmap.get("ANSWER", hmap.get("Answer"))
    full_text_idx = hmap.get("Full Text", hmap.get("fullText"))

    if id_idx is None or answer_idx is None or full_text_idx is None:
        return {}

    blank_re = re.compile(r"__([^_]+?)__")
    questions = {}

    for row in rows_iter:
        cells = list(row)
        if id_idx >= len(cells):
            continue
        raw_id = cells[id_idx].value
        if raw_id is None:
            continue
        try:
            qid = int(raw_id)
        except (ValueError, TypeError):
            continue

        answer_text = str(cells[answer_idx].value or "") if answer_idx < len(cells) else ""
        full_text = str(cells[full_text_idx].value or "") if full_text_idx < len(cells) else ""

        if not answer_text.strip():
            continue

        matches = blank_re.findall(answer_text)
        blanks = []
        for i, payload in enumerate(matches, start=1):
            options = [s.strip() for s in payload.split("/")]
            blanks.append({"index": i, "correct": options[0], "options": options})

        if not blanks:
            continue

        questions[qid] = {
            "id": qid,
            "answer_text": answer_text,
            "full_text": full_text,
            "blanks": blanks,
        }
    return questions


def validate_rfib_sidecar(
    sidecar_path: str = DEFAULT_SIDECAR,
    input_path: str = DEFAULT_INPUT,
    ids: list[int] | None = None,
    full: bool = False,
    check_vi: bool = False,
) -> tuple[bool, list[str]]:
    if not os.path.exists(sidecar_path):
        return False, [f"Sidecar file not found at '{sidecar_path}'"]

    records = []
    errors = []

    with open(sidecar_path, "r", encoding="utf-8") as f:
        for line_num, line in enumerate(f, start=1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                errors.append(f"Line {line_num}: Malformed JSON in sidecar: {e}")

    if errors:
        return False, errors

    sidecar_by_id = {}
    for r in records:
        qid = r.get("id")
        if qid in sidecar_by_id:
            errors.append(f"Duplicate record ID {qid} found in sidecar!")
        else:
            sidecar_by_id[qid] = r

    wb_questions = load_workbook_data(input_path)

    if ids is not None:
        target_ids = sorted(set(ids))
    elif full and wb_questions:
        target_ids = sorted(wb_questions.keys())
    else:
        target_ids = DEFAULT_CURATED_IDS

    total_blanks = 0
    conf_counts = {"high": 0, "medium": 0, "low": 0}
    truncated_raw_records = []

    for qid in target_ids:
        if qid not in sidecar_by_id:
            errors.append(f"Missing required Question ID {qid} in sidecar")
            continue

        r = sidecar_by_id[qid]
        blanks = r.get("blanks", [])
        total_blanks += len(blanks)

        if wb_questions and qid in wb_questions:
            expected_b_count = len(wb_questions[qid]["blanks"])
            if len(blanks) != expected_b_count:
                errors.append(f"ID {qid}: Blank count mismatch (expected {expected_b_count}, got {len(blanks)})")

        for b in blanks:
            b_idx = b.get("blank_index")
            c = b.get("confidence", "low")
            conf_counts[c] = conf_counts.get(c, 0) + 1

            if not b.get("final_explanation", "").strip():
                errors.append(f"ID {qid} Blank {b_idx}: Missing or empty final_explanation")

            if check_vi:
                vi_exp = b.get("vi_explanation", "")
                vi_raw = b.get("vi_raw", "")
                vi_model = b.get("vi_model", "")
                vi_ts = b.get("vi_timestamp", "")

                if not vi_exp or not isinstance(vi_exp, str):
                    errors.append(f"ID {qid} Blank {b_idx}: Missing or empty vi_explanation")
                elif "Tóm lại" not in vi_exp:
                    errors.append(f"ID {qid} Blank {b_idx}: vi_explanation missing required 'Tóm lại' marker")

                if not vi_raw or not isinstance(vi_raw, str):
                    errors.append(f"ID {qid} Blank {b_idx}: Missing or empty vi_raw")

                if not vi_model or not isinstance(vi_model, str) or "qwen3" not in vi_model.lower():
                    errors.append(f"ID {qid} Blank {b_idx}: Missing or invalid vi_model ('{vi_model}')")

                if not vi_ts or not isinstance(vi_ts, str):
                    errors.append(f"ID {qid} Blank {b_idx}: Missing or empty vi_timestamp")

        raws = r.get("raw_artifacts", {})
        for key in ["dr_raw", "qw_raw", "gm_raw"]:
            val = raws.get(key, "")
            if len(val) == 1500:
                truncated_raw_records.append((qid, key))

        models = r.get("phase_models", [])
        if len(models) != 3 or "deepseek" not in models[0].lower() or "qwen" not in models[1].lower() or "gemma" not in models[2].lower():
            errors.append(f"ID {qid}: Incorrect phase_models {models}")

    if truncated_raw_records:
        errors.append(f"Truncated 1,500-char raw artifacts detected in records: {truncated_raw_records}")

    print("============================================================")
    print("RFIB SIDECAR VALIDATION REPORT")
    print("============================================================")
    print(f"Target Records Checked: {len(target_ids)}")
    print(f"Total Blanks Checked:   {total_blanks}")
    print(f"Confidences:            High: {conf_counts.get('high', 0)}, Medium: {conf_counts.get('medium', 0)}, Low: {conf_counts.get('low', 0)}")
    print(f"Vietnamese Checked:     {'YES' if check_vi else 'NO'}")
    print("============================================================")

    if errors:
        print("VALIDATION FAILED WITH ERRORS:")
        for err in errors:
            print(f" - {err}")
        return False, errors
    else:
        print("SUCCESS: Sidecar fully passes all quality gate assertions!")
        return True, []


def parse_args():
    parser = argparse.ArgumentParser(description="Validate RFIB Sidecar Artifacts")
    parser.add_argument("--sidecar", default=DEFAULT_SIDECAR, help="Path to sidecar JSONL")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="Path to workbook XLSX")
    parser.add_argument("--ids", type=str, help="Comma-separated target IDs to validate")
    parser.add_argument("--full", action="store_true", help="Validate all workbook records")
    parser.add_argument("--check-vi", action="store_true", help="Validate Vietnamese translation fields")
    return parser.parse_args()


def main():
    args = parse_args()
    id_list = [int(x.strip()) for x in args.ids.split(",") if x.strip()] if args.ids else None

    success, errors = validate_rfib_sidecar(
        sidecar_path=args.sidecar,
        input_path=args.input,
        ids=id_list,
        full=args.full,
        check_vi=args.check_vi
    )
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
