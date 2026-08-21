from __future__ import annotations

import json
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


LEVELS = ("a2_b1", "b2", "c1")


@dataclass(frozen=True)
class QuestionSources:
    questions: dict[str, dict[str, Any]]
    collocations: set[str]
    invalid_questions: dict[str, list[str]]
    source_hashes: dict[str, str]


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _prompt_key(value: Any) -> str:
    text = _clean(value)
    return (text.replace("“", '"').replace("”", '"').replace("’", "'")
            .replace("–", "-").replace("—", "-").lower())


def _task_type(value: Any) -> str:
    text = _clean(value).lower()
    if "agree or disagree" in text:
        return "agree_disagree"
    if "advantages outweigh" in text or "positive or negative" in text:
        return "advantages_disadvantages"
    if "both sides" in text:
        return "discuss_both_views"
    if "problems" in text and "solutions" in text:
        return "problems_solutions"
    if "responsib" in text:
        return "responsibility"
    if any(term in text for term in ("choose", "which", "rather than", "better", "prefer")):
        return "choose_between"
    return "other"


def _read_workbook_questions(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["453 Prompts"]
    output: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        question_id = str(int(row[0]))
        output[question_id] = {
            "prompt": _clean(row[1]),
            "promptType": _task_type(row[2]),
            "verifiedPrimaryTopic": _clean(row[3]) or None,
            "verifiedSecondaryTopic1": _clean(row[4]) or None,
            "verifiedSecondaryTopic2": _clean(row[5]) or None,
        }
    return output


def _strip_pos(value: Any) -> str:
    return re.sub(r"\s*\([^)]*\)\s*", "", _clean(value)).strip(" ,;:")


def load_collocation_allowlist(root: Path) -> set[str]:
    path = root / "public" / "database" / "The_Academic_Collocation_List.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Academic Collocation List"]
    allowlist: set[str] = set()
    component_one = ""
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if row[2]:
            component_one = _strip_pos(row[2])
        component_two = _strip_pos(row[3]) if len(row) > 3 else ""
        if component_one and component_two:
            allowlist.add(f"{component_one} {component_two}".lower())
    return allowlist


def load_question_sources(root: Path) -> QuestionSources:
    source_path = root / "public" / "database" / "Write Essay" / "essay-questions-with-vocab.json"
    data = json.loads(source_path.read_text(encoding="utf-8-sig"))
    workbook_path = root / "public" / "database" / "Write Essay" / "PTE_453_Cleaned_Prompts_Topics_Task_Types.xlsx"
    collocation_path = root / "public" / "database" / "The_Academic_Collocation_List.xlsx"
    workbook_questions = _read_workbook_questions(workbook_path)
    questions: dict[str, dict[str, Any]] = {}
    invalid_questions: dict[str, list[str]] = {}
    for item in data:
        question_id = str(item.get("id") or "").strip()
        if not question_id or question_id not in workbook_questions:
            raise ValueError(f"question {question_id!r} is missing from the authoritative workbook")
        workbook_item = workbook_questions[question_id]
        if _prompt_key(item.get("prompt")) != _prompt_key(workbook_item["prompt"]):
            raise ValueError(f"question {question_id} prompt differs from authoritative workbook")
        sample_responses = item.get("sampleResponses")
        if not isinstance(sample_responses, dict) or not isinstance(sample_responses.get("levels"), dict):
            raise ValueError(f"question {question_id} has no levelled sample responses")
        for level_id in LEVELS:
            level = sample_responses["levels"].get(level_id)
            variants = [v for v in (level or {}).get("variants", []) if isinstance(v, dict)]
            if not variants:
                raise ValueError(f"question {question_id} has no {level_id} sample variant metadata")
            approved = [v for v in variants if v.get("qa", {}).get("status") == "approved"]
            # Some legacy records contain complete samples without the later QA envelope.
            # Preserve those as explicit fallback sources; the generator/audit records the
            # fallback and never labels it as an approved sample.
            sample_responses["levels"][level_id]["_supportApprovedVariantIds"] = [
                str(v.get("id") or "") for v in approved
            ]
            sample_responses["levels"][level_id]["_supportEssayVariantIds"] = [
                str(v.get("id") or "") for v in variants if v.get("essay")
            ]
        source_issues = []
        if not workbook_item["prompt"]:
            source_issues.append("missing_prompt")
            invalid_questions[question_id] = source_issues
        questions[question_id] = {
            **workbook_item,
            "id": question_id,
            "title": _clean(item.get("title")),
            "targetVocabulary": item.get("targetVocabulary") or {},
            "sampleResponses": sample_responses,
            "auditStatus": item.get("auditStatus"),
            "auditConsensus": item.get("auditConsensus"),
            "sourceIssues": source_issues,
        }
    if set(questions) != set(workbook_questions):
        raise ValueError("runtime question data and workbook question IDs do not reconcile")
    return QuestionSources(
        questions=questions,
        collocations=load_collocation_allowlist(root),
        invalid_questions=invalid_questions,
        source_hashes={
            "questionWorkbookSha256": _sha256_file(workbook_path),
            "questionJsonSha256": _sha256_file(source_path),
            "collocationWorkbookSha256": _sha256_file(collocation_path),
        },
    )
