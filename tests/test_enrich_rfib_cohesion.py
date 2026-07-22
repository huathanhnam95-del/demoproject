import unittest
from unittest.mock import patch, MagicMock
import sys
import os
import openpyxl

# Temporarily add path to import the script
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import scripts.enrich_rfib_cohesion as script

class TestRFIBCohesionCore(unittest.TestCase):

    def test_blank_extraction(self):
        # Blank extraction returns sequential indexes and first option as correct
        answer = "This is a __proposal/reform/change__ and another __all/all of which__ text."
        blanks = script.extract_blanks(answer)
        self.assertEqual(len(blanks), 2)
        self.assertEqual(blanks[0], {"index": 1, "correct": "proposal", "options": ["proposal", "reform", "change"]})
        self.assertEqual(blanks[1], {"index": 2, "correct": "all", "options": ["all", "all of which"]})

    def test_cohesion_boolean_normalization(self):
        # Boolean False/True normalization
        self.assertTrue(script.is_cohesion_false(False))
        self.assertTrue(script.is_cohesion_false("false"))
        self.assertTrue(script.is_cohesion_false("  FALSE  "))
        self.assertFalse(script.is_cohesion_false(True))
        self.assertFalse(script.is_cohesion_false("true"))
        
        self.assertEqual(script.cohesion_label(False), "FALSE")
        self.assertEqual(script.cohesion_label(True), "TRUE")

    def test_prompt_content(self):
        # Prompt includes Column H and the four-tier rubric
        blanks = [{"index": 1, "correct": "proposal", "options": ["proposal", "reform"]}]
        prompt = script.build_prompt("__proposal/reform__", "Full solved text", "Existing details", blanks, "FALSE")
        self.assertIn("FALSE", prompt)
        self.assertIn("Existing details", prompt)
        self.assertIn("Correct", prompt)
        self.assertIn("Partially Correct", prompt)
        self.assertIn("Incorrect", prompt)
        self.assertIn("No Cohesion", prompt)

    def test_response_validation_valid(self):
        blanks = [
            {"index": 1, "correct": "proposal", "options": ["proposal", "reform"]},
            {"index": 2, "correct": "all", "options": ["all", "all of which"]}
        ]
        valid_response = {
            "cohesion_verification_status": "Correct",
            "cohesion_verification_notes": "All clean.",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "proposal",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good noun."
                },
                {
                    "blank_index": 2,
                    "correct_answer": "all",
                    "selection_reason": "pronoun",
                    "cohesion_tie": "ref",
                    "detailed_student_explanation": "Good pronoun."
                }
            ]
        }
        errors = script.validate_response(valid_response, blanks)
        self.assertEqual(errors, [])

    def test_response_validation_invalid_status(self):
        blanks = [{"index": 1, "correct": "proposal", "options": ["proposal", "reform"]}]
        response = {
            "cohesion_verification_status": "InvalidStatus",
            "cohesion_verification_notes": "Note",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "proposal",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good."
                }
            ]
        }
        errors = script.validate_response(response, blanks)
        self.assertTrue(any("Invalid status" in e for e in errors))

    def test_response_validation_count_mismatch(self):
        blanks = [
            {"index": 1, "correct": "proposal", "options": ["proposal", "reform"]},
            {"index": 2, "correct": "all", "options": ["all", "all of which"]}
        ]
        response = {
            "cohesion_verification_status": "Correct",
            "cohesion_verification_notes": "Note",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "proposal",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good."
                }
            ]
        }
        errors = script.validate_response(response, blanks)
        self.assertTrue(any("Explanation count mismatch" in e for e in errors))

    def test_response_validation_duplicate_indexes(self):
        blanks = [
            {"index": 1, "correct": "proposal", "options": ["proposal", "reform"]},
            {"index": 2, "correct": "all", "options": ["all", "all of which"]}
        ]
        response = {
            "cohesion_verification_status": "Correct",
            "cohesion_verification_notes": "Note",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "proposal",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good."
                },
                {
                    "blank_index": 1,
                    "correct_answer": "all",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good."
                }
            ]
        }
        errors = script.validate_response(response, blanks)
        self.assertTrue(any("blank_index" in e or "Duplicate" in e for e in errors))

    def test_response_validation_wrong_answer(self):
        blanks = [{"index": 1, "correct": "proposal", "options": ["proposal", "reform"]}]
        response = {
            "cohesion_verification_status": "Correct",
            "cohesion_verification_notes": "Note",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "reform", # Should be proposal
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "Good."
                }
            ]
        }
        errors = script.validate_response(response, blanks)
        self.assertTrue(any("correct_answer" in e for e in errors))

    def test_response_validation_empty_explanation(self):
        blanks = [{"index": 1, "correct": "proposal", "options": ["proposal", "reform"]}]
        response = {
            "cohesion_verification_status": "Correct",
            "cohesion_verification_notes": "Note",
            "detailed_cohesion_explanations": [
                {
                    "blank_index": 1,
                    "correct_answer": "proposal",
                    "selection_reason": "noun",
                    "cohesion_tie": "",
                    "detailed_student_explanation": "  " # Empty
                }
            ]
        }
        errors = script.validate_response(response, blanks)
        self.assertTrue(any("detailed_student_explanation is empty" in e for e in errors))


class TestRFIBCliAndSelection(unittest.TestCase):

    def test_save_every_zero_rejected(self):
        # --save-every 0 is rejected with exit code 2
        with self.assertRaises(SystemExit) as cm:
            script.parse_arguments(["--save-every", "0"])
        self.assertEqual(cm.exception.code, 2)

    def test_mutually_exclusive_selectors(self):
        # --test, --ids, and --start are mutually exclusive
        with self.assertRaises(SystemExit) as cm:
            script.parse_arguments(["--test", "10", "--ids", "13"])
        self.assertEqual(cm.exception.code, 2)

        with self.assertRaises(SystemExit) as cm:
            script.parse_arguments(["--test", "10", "--start", "5"])
        self.assertEqual(cm.exception.code, 2)

        with self.assertRaises(SystemExit) as cm:
            script.parse_arguments(["--ids", "13", "--start", "5"])
        self.assertEqual(cm.exception.code, 2)

    @patch('openpyxl.load_workbook')
    def test_empty_selector_exits_2(self, mock_load):
        # Mock a workbook that has IDs, but none match a specific requested ID selector
        mock_wb = MagicMock()
        mock_ws = MagicMock()
        mock_load.return_value = mock_wb
        mock_wb.active = mock_ws
        
        # Row 1 headers, Row 2 dummy data with ID=100
        mock_ws.max_row = 2
        mock_ws.max_column = 2
        
        def cell_val(row, column):
            cell = MagicMock()
            if row == 1:
                cell.value = "ID" if column == 1 else "TITLE"
            else:
                cell.value = 100 if column == 1 else "Title 100"
            return cell
            
        mock_ws.cell.side_effect = cell_val

        # Call script main or target function that resolves candidate rows
        with self.assertRaises(SystemExit) as cm:
            script.resolve_candidate_rows(mock_ws, {"ID": 1}, ids_str="13,14", start=0, limit=0, test=0)
        self.assertEqual(cm.exception.code, 2)

    @patch('openpyxl.load_workbook')
    def test_curated_sample_builder(self, mock_load):
        # Verify curated sample selector builds exactly [13, 14, 9, 15, 3, 4, 1, 2, 5, 6]
        # We need to mock the worksheet cells for IDs 1..20
        # and check the output IDs
        mock_wb = MagicMock()
        mock_ws = MagicMock()
        mock_load.return_value = mock_wb
        mock_wb.active = mock_ws
        
        mock_ws.max_row = 21
        mock_ws.max_column = 6
        
        hmap = {
            "ID": 1,
            "TITLE": 2,
            "ANSWER": 3,
            "Full Text": 4,
            "Cohesion Feature": 5,
            "Cohesion Feature Details": 6
        }

        # ID values correspond to row index - 1 (row 2 -> ID 1, row 14 -> ID 13)
        # We make ID 9 and 15 have Cohesion Feature = False
        # We make IDs 3, 4, 5, 6 have 5+ blanks in ANSWER
        def cell_val(row, column):
            cell = MagicMock()
            if row == 1:
                # Return header names
                for k, v in hmap.items():
                    if v == column:
                        cell.value = k
                        return cell
            
            rid = row - 1
            if column == hmap["ID"]:
                cell.value = rid
            elif column == hmap["Cohesion Feature"]:
                cell.value = False if rid in [9, 15] else True
            elif column == hmap["ANSWER"]:
                if rid in [3, 4, 5, 6]:
                    cell.value = "__a/b__ __c/d__ __e/f__ __g/h__ __i/j__"
                else:
                    cell.value = "__a/b__"
            else:
                cell.value = "dummy"
            return cell

        mock_ws.cell.side_effect = cell_val

        candidate_rows = script._build_curated_sample(mock_ws, hmap, 10)
        # Extract the resolved IDs
        resolved_ids = []
        for r in candidate_rows:
            resolved_ids.append(mock_ws.cell(row=r, column=hmap["ID"]).value)
            
        self.assertEqual(resolved_ids, [13, 14, 9, 15, 3, 4, 1, 2, 5, 6])


class TestOllamaRetries(unittest.TestCase):

    @patch('scripts.enrich_rfib_cohesion.requests.post')
    def test_retry_exhaustion_makes_three_calls(self, mock_post):
        mock_post.side_effect = Exception("Ollama connection failed")
        
        with patch('time.sleep'): # skip delay
            response = script.query_gemma("test prompt", max_retries=3)
            
        self.assertIsNone(response)
        self.assertEqual(mock_post.call_count, 3)


class TestCheckpointAndRecovery(unittest.TestCase):

    @patch('builtins.open')
    @patch('os.path.exists')
    def test_duplicate_sidecar_ids_rejected(self, mock_exists, mock_open):
        mock_exists.return_value = True
        mock_open.return_value.__enter__.return_value = [
            '{"id": 13, "status": "Correct"}\n',
            '{"id": 13, "status": "Incorrect"}\n' # Duplicate ID
        ]
        with self.assertRaises(ValueError) as cm:
            script.load_sidecar_to_dict("dummy_sidecar.jsonl")
        self.assertIn("Duplicate ID 13 detected in sidecar", str(cm.exception))

    @patch('os.replace')
    @patch('builtins.open')
    def test_atomic_sidecar_rewrite(self, mock_open, mock_replace):
        records = {
            13: {"id": 13, "status": "Correct"},
            14: {"id": 14, "status": "Incorrect"}
        }
        script.save_sidecar_atomic("dummy_sidecar.jsonl", records)
        mock_replace.assert_called_once_with("dummy_sidecar.jsonl.tmp", "dummy_sidecar.jsonl")

    @patch('scripts.enrich_rfib_cohesion.save_workbook_safe')
    @patch('scripts.enrich_rfib_cohesion.save_sidecar_atomic')
    def test_checkpoint_order(self, mock_save_sidecar, mock_save_wb):
        # Verify sidecar is saved first, workbook second
        call_order = []
        mock_save_sidecar.side_effect = lambda *a, **kw: call_order.append("sidecar")
        mock_save_wb.side_effect = lambda *a, **kw: call_order.append("workbook")

        mock_wb = MagicMock()
        script.checkpoint(mock_wb, "out.xlsx", "sidecar.jsonl", {})
        self.assertEqual(call_order, ["sidecar", "workbook"])

    def test_reconciliation_excel_ahead_and_mismatches(self):
        # Mock workbook sheet cell access
        mock_ws = MagicMock()
        col_detailed = 1
        col_verification = 2
        
        # Scenario 1: Valid matching sidecar, missing Excel -> render from sidecar
        mock_ws.cell(row=5, column=col_detailed).value = ""
        mock_ws.cell(row=5, column=col_verification).value = ""
        
        sidecar_records = {
            13: {
                "id": 13,
                "status": "Correct",
                "notes": "No issues",
                "explanations": [{"blank_index": 1, "correct_answer": "a", "detailed_student_explanation": "exp"}]
            }
        }
        
        # Test sidecar-ahead (Excel missing)
        action = script.reconcile_row(mock_ws, 5, 13, col_detailed, col_verification, sidecar_records)
        self.assertEqual(action, "render")
        
        # Scenario 2: Valid Excel cells, missing/invalid sidecar -> reprocess from Ollama
        mock_ws.cell(row=6, column=col_detailed).value = "Blank 1 ('a'):\n  Why this answer: x\n  Student explanation: exp"
        mock_ws.cell(row=6, column=col_verification).value = "[Correct] Note"
        
        action = script.reconcile_row(mock_ws, 6, 14, col_detailed, col_verification, {})
        self.assertEqual(action, "reprocess")

        # Scenario 3: Valid matching both -> skip
        mock_ws.cell(row=7, column=col_detailed).value = "Blank 1 ('a'):\n  Why this answer: x\n  Student explanation: exp"
        mock_ws.cell(row=7, column=col_verification).value = "[Correct] Note"
        sidecar_records[14] = {
            "id": 14,
            "status": "Correct",
            "notes": "Note",
            "explanations": [{"blank_index": 1, "correct_answer": "a", "detailed_student_explanation": "exp"}]
        }
        action = script.reconcile_row(mock_ws, 7, 14, col_detailed, col_verification, sidecar_records)
        self.assertEqual(action, "skip")

    @patch('shutil.copy2')
    @patch('os.path.exists')
    @patch('os.remove')
    def test_reset_sidecar_logs_and_clears(self, mock_remove, mock_exists, mock_copy):
        mock_exists.return_value = True
        with patch('builtins.open') as mock_open:
            mock_open.return_value.__enter__.return_value = ['{"id":1}\n']
            script.reset_sidecar_if_requested("sidecar.jsonl", reset=True)
            mock_remove.assert_called_once_with("sidecar.jsonl")


class TestValidationAndPreservation(unittest.TestCase):

    def test_preflight_check_unsupported(self):
        # Create a mock sheet that has conditional formatting
        mock_ws = MagicMock()
        mock_ws.conditional_formatting = ["dummy"]
        mock_ws.data_validations = []
        mock_ws.tables = []
        mock_ws._images = []
        mock_ws._charts = []
        mock_wb = MagicMock()
        mock_wb.defined_names = []
        mock_wb._external_links = []
        
        with self.assertRaises(ValueError) as cm:
            script.run_preflight_check(mock_wb, mock_ws)
        self.assertIn("Unsupported advanced features detected", str(cm.exception))

    def test_preflight_check_supported(self):
        mock_ws = MagicMock()
        mock_ws.conditional_formatting = []
        mock_ws.data_validations = []
        mock_ws.tables = []
        mock_ws._images = []
        mock_ws._charts = []
        mock_wb = MagicMock()
        mock_wb.defined_names = []
        mock_wb._external_links = []
        
        # Should not raise any error
        script.run_preflight_check(mock_wb, mock_ws)

    def test_semantic_comparison_mismatch(self):
        # Create mock cell styles and values that mismatch
        cell_a = MagicMock(value="A", font=MagicMock(name="Arial"), fill=MagicMock(fill_type="solid"), border=MagicMock(), alignment=MagicMock(), number_format="General", protection=MagicMock())
        cell_b = MagicMock(value="B", font=MagicMock(name="Arial"), fill=MagicMock(fill_type="solid"), border=MagicMock(), alignment=MagicMock(), number_format="General", protection=MagicMock())
        
        # Test values mismatch in column 1 (ID)
        self.assertFalse(script.compare_cell_style_and_value(cell_a, cell_b, col_idx=1))
        # Appended columns (e.g. index 13 Detailed Cohesion Explanation) are allowed to differ
        self.assertTrue(script.compare_cell_style_and_value(cell_a, cell_b, col_idx=13))

    @patch('scripts.enrich_rfib_cohesion.load_sidecar_to_dict')
    @patch('openpyxl.load_workbook')
    def test_validation_logic(self, mock_load, mock_load_sidecar):
        mock_wb = MagicMock()
        mock_ws = MagicMock()
        mock_load.return_value = mock_wb
        mock_wb.active = mock_ws
        
        mock_ws.max_row = 3
        mock_ws.max_column = 14
        
        hmap = {
            "ID": 1,
            "TITLE": 2,
            "ANSWER": 3,
            "Full Text": 4,
            "Cohesion Feature": 5,
            "Cohesion Feature Details": 6,
            "Detailed Cohesion Explanation": 13,
            "Cohesion Verification": 14
        }
        
        def cell_val(row, column):
            cell = MagicMock()
            if row == 1:
                for k, v in hmap.items():
                    if v == column:
                        cell.value = k
                        return cell
            
            rid = row - 1
            if column == hmap["ID"]:
                cell.value = rid
            elif column == hmap["ANSWER"]:
                cell.value = "__ans1/ans2__"
            elif column == hmap["Detailed Cohesion Explanation"]:
                cell.value = "Blank 1 ('ans1'): Explanation text"
            elif column == hmap["Cohesion Verification"]:
                cell.value = "[Correct] Notes here"
            else:
                cell.value = "dummy"
            return cell
            
        mock_ws.cell.side_effect = cell_val

        # Sidecar records with mismatched status in ID 2
        sidecar_records = {
            1: {
                "id": 1,
                "status": "Correct",
                "notes": "Notes here",
                "explanations": [{"blank_index": 1, "correct_answer": "ans1", "detailed_student_explanation": "Explanation text"}]
            },
            2: {
                "id": 2,
                "status": "Incorrect", # Mismatched with Excel [Correct] status
                "notes": "Notes here",
                "explanations": [{"blank_index": 1, "correct_answer": "ans1", "detailed_student_explanation": "Explanation text"}]
            }
        }
        mock_load_sidecar.return_value = sidecar_records

        args = MagicMock(input="dummy.xlsx", sidecar="dummy.jsonl", ids="", validate_only=True, baseline="")
        
        with self.assertRaises(SystemExit) as cm:
            script.run_validate_only(args)
        self.assertEqual(cm.exception.code, 1)


if __name__ == '__main__':
    unittest.main()
