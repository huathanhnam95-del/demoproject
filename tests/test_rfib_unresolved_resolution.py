import hashlib
import copy
from collections import Counter
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.patch_rfib_unresolved import _validate_aggregate_shape, _validate_fresh_audit, build_aggregate_report


ROOT = Path(__file__).resolve().parents[1]
RFIB_DIR = ROOT / "public" / "database" / "RFIB"

EXPECTED = {
    (146, 3): {
        "answer": "open",
        "grammar_tag": "Collocation (Verb + Predicate Adjective)",
        "final_explanation": (
            "The correct answer is 'open' because it pairs with the verb 'leaves' to form the "
            "fixed collocation 'leave open the possibility,' which means to allow for a chance, "
            "alternative interpretation, or unresolved outcome. In this sentence, the object noun "
            "clause ('the possibility that they are not...') is placed after the adjective 'open' "
            "for syntactic balance.\n\n"
            "Why the distractors are incorrect:\n"
            "- 'go': This is a base-form verb and cannot function as a predicate adjective complement "
            "after 'leaves' (one cannot say *leaves go the possibility* in standard English).\n"
            "- 'covered': While 'covered' is an adjective/participle, the phrase *'leaves covered the "
            "possibility'* is not a recognized English collocation and creates a semantic contradiction "
            "by implying concealment rather than availability.\n"
            "- 'undoubted': An adjective meaning 'certain' or 'indisputable.' Using it here would "
            "contradict the entire premise of the sentence, which emphasizes maintaining ambiguity and "
            "non-confinement rather than certainty."
        ),
        "concise_explanation": (
            "'Leaves open the possibility' is a fixed collocation meaning to allow for an alternative "
            "chance or interpretation. The distractors fail grammatically ('go' is a verb) or "
            "semantically ('covered' and 'undoubted' contradict the idea of keeping an option available)."
        ),
    },
    (520, 4): {
        "answer": "tend",
        "grammar_tag": "Grammar/Usage (Catenative Verb of Tendency / Habit)",
        "final_explanation": (
            "The correct answer is 'tend' because it functions as a lexical verb expressing a general "
            "inclination, habit, or typical characteristic when paired with a to-infinitive ('tend to be'). "
            "In this context, 'tend to be higher-priced' accurately states that B2B transactions are "
            "typically or usually more expensive than consumer sales.\n\n"
            "Why the distractors are incorrect:\n"
            "- 'extend': Means to lengthen, prolong, or stretch out physically or temporally. It cannot "
            "describe a general characteristic of pricing (e.g., sales do not 'extend to be expensive').\n"
            "- 'contend': Means to assert or argue a position firmly ('contend that') or to struggle "
            "against an opponent ('contend with'). It does not express an intrinsic tendency of goods or services.\n"
            "- 'pretend': Means to behave deceptively or simulate an imaginary state. It is logically and "
            "contextually incompatible with describing factual corporate pricing structures."
        ),
        "concise_explanation": (
            "'Tend to be' is a standard construction expressing a general habit or typical state (B2B sales "
            "are usually higher-priced). The rhyming distractors mean to lengthen ('extend'), argue/compete "
            "('contend'), or fake ('pretend'), which are logically impossible here."
        ),
    },
}

EXPECTED_OTHER_REVISION_DIGEST = "68bdc38c19041ec2c2ad8be29c8d0d5b9cdfcc3f43bb6aa70aecb26f618dbbf7"
EXPECTED_OTHER_AUDITED_DIGEST = "38727d01ea2ae3754abcdfd9863b6bf584298f16bc26550f369c633a72202264"
EXPECTED_WORKBOOK_DIGEST = "417a24c142af91c4b42a469a91a905197b46dd3e330210468fef0211e3c2733b"


def load_jsonl(path: Path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def records_by_id(path: Path):
    return {int(record["id"]): record for record in load_jsonl(path)}


def target_blank(record, blank_index):
    return next(blank for blank in record["blanks"] if int(blank["blank_index"]) == blank_index)


def other_records_digest(path: Path):
    records = [record for record in load_jsonl(path) if int(record["id"]) not in {146, 520}]
    payload = "\n".join(
        json.dumps(record, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        for record in records
    ) + "\n"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def workbook_preservation_digest(workbook_path: Path):
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    target_cells = {(147, 11), (147, 13), (480, 11), (480, 13)}
    cells = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if (cell.row, cell.column) in target_cells:
                    continue
                cells.append(
                    (
                        worksheet.title,
                        cell.coordinate,
                        cell.value,
                        cell.data_type,
                        cell.style_id,
                        cell.number_format,
                        cell.alignment.horizontal,
                        cell.alignment.vertical,
                        cell.alignment.wrap_text,
                        cell.protection.locked,
                        cell.protection.hidden,
                    )
                )
    metadata = [
        (
            worksheet.title,
            worksheet.max_row,
            worksheet.max_column,
            worksheet.freeze_panes,
            worksheet.auto_filter.ref,
            tuple(sorted(str(value) for value in worksheet.merged_cells.ranges)),
            worksheet.sheet_view.showGridLines,
        )
        for worksheet in workbook.worksheets
    ]
    payload = {
        "cells": cells,
        "meta": metadata,
        "sheets": workbook.sheetnames,
        "defined_names": sorted(str(value) for value in workbook.defined_names),
    }
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=False, default=str, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


class RFIBUnresolvedResolutionTest(unittest.TestCase):
    def test_both_corrections_and_artifact_preservation(self):
        revision_records = records_by_id(RFIB_DIR / "RFIB_3model_revision.jsonl")
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")

        for key, expected in EXPECTED.items():
            qid, blank_index = key
            revision_blank = target_blank(revision_records[qid], blank_index)
            audited_blank = target_blank(audited_records[qid], blank_index)
            for blank in (revision_blank, audited_blank):
                self.assertEqual(blank["correct_answer"], expected["answer"])
                self.assertEqual(blank["grammar_tag"], expected["grammar_tag"])
                self.assertEqual(blank["final_explanation"], expected["final_explanation"])
                self.assertEqual(blank["concise_explanation"], expected["concise_explanation"])

            self.assertIn(
                audited_blank.get("audit_verdict"),
                {"PASS", "FIXED_CONSENSUS", "FIXED_MAJORITY"},
            )
            self.assertIn(
                audited_blank.get("audit_consensus_type"),
                {"ROUND1_APPROVAL", "MUTUAL_CONSENT", "MAJORITY_CONSENT"},
            )

            summary = audited_records[qid]["audit_summary"]
            self.assertEqual(summary["unresolved_blanks"], 0)
            self.assertEqual(summary["total_blanks"], len(audited_records[qid]["blanks"]))
            detail = next(item for item in summary["details"] if item["blank_index"] == blank_index)
            self.assertEqual(detail["final_verdict"], audited_blank["audit_verdict"])
            self.assertEqual(detail["final_grammar_tag"], expected["grammar_tag"])
            self.assertEqual(detail["final_explanation"], expected["final_explanation"])

        report = json.loads((RFIB_DIR / "RFIB_audit_debate_report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["total_target_questions"], 221)
        self.assertEqual(report["processed_questions"], 221)
        status_counts = Counter(record["audit_status"] for record in audited_records.values())
        self.assertTrue(set(status_counts).issubset({"PASS", "REVISED_WITH_CONSENSUS", "UNRESOLVED"}))
        self.assertEqual(report["round1_pass_count"], status_counts["PASS"])
        self.assertEqual(report["debate_revised_count"], status_counts["REVISED_WITH_CONSENSUS"])
        self.assertEqual(report["unresolved_count"], status_counts["UNRESOLVED"])
        self.assertEqual(report["unresolved_count"], 0)
        self.assertEqual(
            report["round1_pass_count"] + report["debate_revised_count"] + report["unresolved_count"],
            report["processed_questions"],
        )
        for transcript in report["debate_transcripts"]:
            self.assertIn(transcript["question_id"], audited_records)
            detail = next(
                detail
                for detail in audited_records[transcript["question_id"]]["audit_summary"]["details"]
                if detail["blank_index"] == transcript["blank_index"]
            )
            for field in ("round_1_votes", "debate_rounds", "consensus_type", "final_explanation", "debate_history"):
                self.assertEqual(transcript[field], detail[field])
            self.assertTrue(transcript["debate_rounds"] > 0)
            self.assertEqual(len(transcript["debate_history"]), transcript["debate_rounds"])

        workbook_path = RFIB_DIR / "RFIB Final ver.xlsx"
        workbook = load_workbook(workbook_path, data_only=False, read_only=False)
        worksheet = workbook["Sheet1"]
        for row, key in ((147, (146, 3)), (480, (520, 4))):
            expected = EXPECTED[key]
            analysis = json.loads(worksheet.cell(row, 11).value)
            workbook_blank = target_blank({"blanks": analysis}, key[1])
            self.assertEqual(workbook_blank["grammar_tag"], expected["grammar_tag"])
            self.assertEqual(workbook_blank["concise_explanation"], expected["concise_explanation"])
            self.assertIn(expected["final_explanation"], worksheet.cell(row, 13).value)
            self.assertIn(expected["concise_explanation"], worksheet.cell(row, 13).value)

        self.assertEqual(
            other_records_digest(RFIB_DIR / "RFIB_3model_revision.jsonl"),
            EXPECTED_OTHER_REVISION_DIGEST,
        )
        self.assertEqual(
            other_records_digest(RFIB_DIR / "RFIB_audited_sample20.jsonl"),
            EXPECTED_OTHER_AUDITED_DIGEST,
        )
        self.assertEqual(workbook_preservation_digest(workbook_path), EXPECTED_WORKBOOK_DIGEST)

    def test_fresh_validator_rejects_wrong_question_status(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        fresh_records = {qid: copy.deepcopy(audited_records[qid]) for qid in (146, 520)}
        fresh_records[520]["audit_status"] = "PASS"
        report = json.loads((RFIB_DIR / "RFIB_audit_debate_report.json").read_text(encoding="utf-8"))
        report.update(
            {
                "total_target_questions": 2,
                "processed_questions": 2,
                "round1_pass_count": 2,
                "debate_revised_count": 0,
                "unresolved_count": 0,
            }
        )
        with self.assertRaises(ValueError):
            _validate_fresh_audit(fresh_records, report)

    def test_fresh_validator_rejects_missing_target_vote_evidence(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        fresh_records = {qid: copy.deepcopy(audited_records[qid]) for qid in (146, 520)}
        target_detail = next(
            detail
            for detail in fresh_records[146]["audit_summary"]["details"]
            if detail["blank_index"] == 3
        )
        del target_detail["round_1_votes"]
        report = json.loads((RFIB_DIR / "RFIB_audit_debate_report.json").read_text(encoding="utf-8"))
        report.update(
            {
                "total_target_questions": 2,
                "processed_questions": 2,
                "round1_pass_count": 1,
                "debate_revised_count": 1,
                "unresolved_count": 0,
            }
        )
        with self.assertRaises(ValueError):
            _validate_fresh_audit(fresh_records, report)

    def test_aggregate_shape_rejects_target_only_report(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        report = json.loads((RFIB_DIR / "RFIB_audit_debate_report.json").read_text(encoding="utf-8"))
        report["total_target_questions"] = 2
        report["processed_questions"] = 2
        with self.assertRaises(ValueError):
            _validate_aggregate_shape(report, audited_records)

    def test_fresh_validator_rejects_absent_target_detail(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        fresh_records = {qid: copy.deepcopy(audited_records[qid]) for qid in (146, 520)}
        fresh_records[146]["audit_summary"]["details"] = [
            detail
            for detail in fresh_records[146]["audit_summary"]["details"]
            if detail["blank_index"] != 3
        ]
        report = json.loads((RFIB_DIR / "RFIB_audit_debate_report.json").read_text(encoding="utf-8"))
        report.update(
            {
                "total_target_questions": 2,
                "processed_questions": 2,
                "round1_pass_count": 1,
                "debate_revised_count": 1,
                "unresolved_count": 0,
            }
        )
        with self.assertRaises(ValueError):
            _validate_fresh_audit(fresh_records, report)

    def test_aggregate_probe_derives_all_pass_counts_instead_of_stale_constants(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        all_pass_records = copy.deepcopy(list(audited_records.values()))
        for record in all_pass_records:
            record["audit_status"] = "PASS"
        report = build_aggregate_report(all_pass_records)
        self.assertEqual(report["total_target_questions"], 221)
        self.assertEqual(report["processed_questions"], 221)
        self.assertEqual(report["round1_pass_count"], 221)
        self.assertEqual(report["debate_revised_count"], 0)
        self.assertEqual(report["unresolved_count"], 0)

    def test_aggregate_rejects_unknown_record_status(self):
        audited_records = records_by_id(RFIB_DIR / "RFIB_audited_sample20.jsonl")
        invalid_records = copy.deepcopy(list(audited_records.values()))
        invalid_records[0]["audit_status"] = "FABRICATED"
        with self.assertRaises(ValueError):
            build_aggregate_report(invalid_records)


if __name__ == "__main__":
    unittest.main()
